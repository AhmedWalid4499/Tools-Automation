"""Order tracker: the orders read in the Overview as a tracker sheet to share with management or customers.

Two views:  "management" - everything (people, value, lead time, latest GOLD note)
            "customer"   - only what a customer should see, in plain words
Comments / next steps typed in the app are kept per order (tracker_comments.json) and go into the sheet.
"""
import json
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import gold_overview as go

ORANGE, INK, MUTED, LINE = "FF7900", "1F2328", "6B7280", "E5E7EB"
FONT = "Segoe UI"
VIEWS = {"management": "Management", "customer": "Customer"}

# Delivery status in plain words (the RAG of the target date)
STATUS_WORDS = {"Green": "On track", "Amber": "Due soon", "Red": "Past target date", "No date": "Date to be confirmed"}
LEAD_WORDS = {"Green": "Within 120 days", "Amber": "Close to 120 days", "Red": "Over 120 days", "No date": "-"}


def _site(o):
    return ", ".join(x for x in (o.get("city") or "", (o.get("country") or "").title()) if x)


def _next(o):
    if not o.get("next_milestone"):
        return ""
    when = go.fmt(o.get("next_milestone_date"))
    return o["next_milestone"].replace("...", "") + (f" ({when})" if when else "")


def _note(o):
    if not o.get("latest_note"):
        return ""
    return f"{o.get('latest_note_date', '')} {o.get('latest_note_by', '')}: {o['latest_note']}".strip()


# (key, title, width in the sheet, views it is in, value(o, comment))
COLUMNS = [
    ("order", "GOLD order", 13, "mc", lambda o, c: o["order"]),
    ("customer", "Customer", 24, "m", lambda o, c: o.get("customer", "")),
    ("project", "Project", 24, "mc", lambda o, c: o.get("project", "")),
    ("site", "Site", 22, "mc", lambda o, c: _site(o)),
    ("service", "Service", 26, "mc", lambda o, c: o.get("product", "")),
    ("order_status", "Order status", 12, "mc", lambda o, c: o.get("order_status", "")),
    ("target", "Target delivery date", 14, "mc", lambda o, c: o.get("key_date")),
    ("status", "Delivery status", 16, "mc", lambda o, c: STATUS_WORDS.get(o.get("rag"), o.get("rag", ""))),
    ("days", "Days left", 9, "m", lambda o, c: o.get("days_left")),
    ("revised", "Revised", 10, "mc", lambda o, c: "Yes" if o.get("revised_tdd") else "No"),
    ("next", "Next milestone", 30, "mc", lambda o, c: _next(o)),
    ("stm", "STM responsible", 20, "m", lambda o, c: o.get("stm", "")),
    ("lead", "120-day lead time", 15, "m", lambda o, c: LEAD_WORDS.get(o.get("lt_rag"), "")),
    ("value", "Order value", 15, "m", lambda o, c: go.money(o.get("total_value"), o.get("currency"))
     if o.get("total_value") is not None else ""),
    ("note", "Latest GOLD note", 45, "m", lambda o, c: _note(o)),
    ("comment", "Comments / next steps", 45, "mc", lambda o, c: c or ""),
]


def columns(view):
    flag = "c" if view == "customer" else "m"
    return [c for c in COLUMNS if flag in c[3]]


def rows(overviews, view, comments):
    cols = columns(view)
    return [[fn(o, comments.get(o["order"], "")) for _, _, _, _, fn in cols] for o in overviews]


# ---------------------------------------------------------------- comments kept per order
def load_comments(path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save_comments(path, comments):
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps({k: v for k, v in comments.items() if v}, indent=1, ensure_ascii=False),
                   encoding="utf-8")
    tmp.replace(path)


# ---------------------------------------------------------------- the Excel tracker
RAG_FILLS = {"On track": ("C6EFCE", "006100"), "Due soon": ("FFEB9C", "9C5700"),
             "Past target date": ("FFC7CE", "9C0006"), "Date to be confirmed": ("EDEDED", "595959"),
             "Within 120 days": ("C6EFCE", "006100"), "Close to 120 days": ("FFEB9C", "9C5700"),
             "Over 120 days": ("FFC7CE", "9C0006")}


def export(path, overviews, view, comments, title="Order delivery tracker", prepared_for=""):
    cols = columns(view)
    data = rows(overviews, view, comments)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    ws.sheet_view.showGridLines = False
    n = len(cols)
    thin = Side(style="thin", color=LINE)

    # title block
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=18, bold=True, color=INK)
    customers = sorted({o.get("customer", "") for o in overviews if o.get("customer")})
    sub = [f"{len(overviews)} order(s)"]
    if prepared_for:
        sub.append(f"prepared for {prepared_for}")
    elif view == "customer" and len(customers) == 1:
        sub.append(customers[0])
    sub.append(f"status as of {datetime.now():%d-%b-%Y %H:%M}")
    ws["A2"] = "  ·  ".join(sub)
    ws["A2"].font = Font(name=FONT, size=10, color=MUTED)
    ws.row_dimensions[1].height = 30

    # summary: delivery status counts as coloured boxes
    counts = {w: sum(1 for o in overviews if STATUS_WORDS.get(o.get("rag")) == w) for w in STATUS_WORDS.values()}
    col = 1
    for word, count in counts.items():
        if not count and word == "Date to be confirmed":
            continue
        fill, fnt = RAG_FILLS[word]
        a, b = ws.cell(4, col, count), ws.cell(5, col, word)
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        for r in (4, 5):
            for c in (col, col + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=fill)
        a.font = Font(name=FONT, size=16, bold=True, color=fnt)
        b.font = Font(name=FONT, size=9, bold=True, color=fnt)
        a.alignment = b.alignment = Alignment(horizontal="center", vertical="center")
        col += 3
    if view == "management":
        totals = {}
        for o in overviews:
            if o.get("total_value") is not None:
                totals[o.get("currency", "")] = totals.get(o.get("currency", ""), 0) + o["total_value"]
        if totals:
            c = ws.cell(4, col + 1, "Total order value")
            c.font = Font(name=FONT, size=9, bold=True, color=MUTED)
            v = ws.cell(5, col + 1, "   ".join(go.money(t, cur) for cur, t in totals.items()))
            v.font = Font(name=FONT, size=12, bold=True, color=ORANGE)
    ws.row_dimensions[4].height = 26

    # the tracker table
    top = 7
    for j, (key, head, width, _, _) in enumerate(cols, 1):
        c = ws.cell(top, j, head)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ORANGE)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[top].height = 30
    wrap = {"project", "service", "next", "note", "comment", "customer", "site"}
    for i, values in enumerate(data, top + 1):
        for j, ((key, _, _, _, _), v) in enumerate(zip(cols, values), 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=FONT, size=10, color=INK, bold=key == "order")
            c.alignment = Alignment(vertical="top", wrap_text=key in wrap)
            c.border = Border(bottom=thin)
            if isinstance(v, (date, datetime)):
                c.number_format = "dd-mmm-yyyy"
            if key in ("status", "lead") and v in RAG_FILLS:
                fill, fnt = RAG_FILLS[v]
                c.fill = PatternFill("solid", fgColor=fill)
                c.font = Font(name=FONT, size=10, bold=True, color=fnt)
                c.alignment = Alignment(horizontal="center", vertical="top")
    last = top + max(len(data), 1)
    if data:
        t = Table(displayName="Tracker", ref=f"A{top}:{get_column_letter(n)}{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(t)
    ws.freeze_panes = ws.cell(top + 1, 2)
    note = ("Delivery status compares the target delivery date (revised when there is one) with today: "
            f"Due soon = {go.AMBER_DAYS} days or less.")
    ws.cell(last + 2, 1, note).font = Font(name=FONT, size=8, italic=True, color=MUTED)

    # print: landscape, one page wide, title row repeated
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{top}:{top}"
    ws.oddFooter.center.text = "&P / &N"
    ws.oddHeader.left.text = title

    # one block per order (easy to read or forward one order)
    ds = wb.create_sheet("Per order")
    ds.sheet_view.showGridLines = False
    ds.column_dimensions["A"].width = 26
    ds.column_dimensions["B"].width = 80
    r = 1
    for o, values in zip(overviews, data):
        ds.cell(r, 1, f"{o['order']}  ·  {o.get('customer', '')}" if view == "management" else o["order"]).font = \
            Font(name=FONT, size=13, bold=True, color=ORANGE)
        r += 1
        for (key, head, _, _, _), v in zip(cols, values):
            if key == "order" or v in (None, ""):
                continue
            ds.cell(r, 1, head).font = Font(name=FONT, size=10, bold=True, color=MUTED)
            c = ds.cell(r, 2, v)
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(v, (date, datetime)):
                c.number_format = "dd-mmm-yyyy"
            if key in ("status", "lead") and v in RAG_FILLS:
                fill, fnt = RAG_FILLS[v]
                c.fill = PatternFill("solid", fgColor=fill)
                c.font = Font(name=FONT, size=10, bold=True, color=fnt)
            r += 1
        r += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path

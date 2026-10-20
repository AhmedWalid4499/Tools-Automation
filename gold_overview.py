"""Order overview: collect the key facts of a GOLD order (read-only) and export them to Excel.

Sources: order header (customer, status, type) · Information (project, scope, local site,
dates, references, account manager) · Billing > Site (city, state, postcode, country, site ids,
product) · Milestones (initial / revised target delivery date, next open milestone) · Teams (STM
and other owners) · Notes (latest note).

Delivery status (RAG) against the Revised Target Delivery Date, or the Initial one if there
is no revision: Red = date passed, Amber = 14 days or less left, Green = more than 14 days.
"""
import re
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import gold_extract as gx
import gold_milestones as gm
import open_order as oo

AMBER_DAYS = 14
LEAD_TIME_DAYS = 120  # an order should be delivered within 120 days of the customer signature
TOTAL_RE = re.compile(r"Total Charges in (\w+):\s*([-\d,.]+)\s+([-\d,.]+)\s+([-\d,.]+)\s+([-\d,.]+)")
RAG_COLORS = {  # Excel fill, font
    "Red": ("FFC7CE", "9C0006"), "Amber": ("FFEB9C", "9C5700"), "Green": ("C6EFCE", "006100"),
    "No date": ("EDEDED", "595959"),
}


def _d(text):
    try:
        return gm.parse_date(text) if text else None
    except ValueError:
        return None


def _field(fields, *names, section=None):
    """First value whose label matches one of names (case-insensitive), optionally in a section."""
    for name in names:
        for sec, label, value in fields:
            if label.strip().lower() == name.lower() and (section is None or sec.lower().startswith(section.lower())):
                if value and value.lower() not in ("n/a", "none selected"):
                    return value.strip()
    return ""


def rag(target, today=None):
    """(RAG, days_left, text) for a target date."""
    if not target:
        return "No date", None, "No target date"
    today = today or date.today()
    days = (target - today).days
    if days < 0:
        return "Red", days, f"Elapsed {-days} day{'s' if days != -1 else ''} ago"
    if days == 0:
        return "Amber", 0, "Due today"
    return ("Amber" if days <= AMBER_DAYS else "Green"), days, f"{days} day{'s' if days != 1 else ''} left"


def _num(text):
    try:
        return float(text.replace(",", ""))
    except (AttributeError, ValueError):
        return None


def lead_time(sign_date, key_date):
    """120-day lead-time check: (RAG, deadline, margin_days, text). margin = days the planned
    delivery date is before (+) or after (-) the deadline."""
    if not sign_date:
        return "No date", None, None, "No customer signature date"
    deadline = sign_date + timedelta(days=LEAD_TIME_DAYS)
    if not key_date:
        return "No date", deadline, None, "No target delivery date"
    margin = (deadline - key_date).days
    if margin < 0:
        return "Red", deadline, margin, f"Planned {-margin} day{'s' if margin != -1 else ''} after the 120-day deadline"
    if margin <= AMBER_DAYS:
        return "Amber", deadline, margin, f"Planned only {margin} day{'s' if margin != 1 else ''} before the deadline"
    return "Green", deadline, margin, f"Planned {margin} days before the deadline"


def money(value, currency=""):
    if value is None:
        return ""
    return f"{currency} {value:,.2f}".strip()


def collect(driver, wait, order, log=print):
    """Read an order and return its overview dict (nothing is changed in GOLD)."""
    oo.open_orders_tab(driver, wait, discard_changes=True)
    oo.untick_my_orders(driver)
    row = oo.search_order(driver, wait, order)
    # Search-result row: Order, Type, Customer, BAN, Orange Site Ref, Site Id, Cust. Site Ref,
    # Status, Installed Offer (clean values; the Billing page hyphenates long words).
    cells = [c.text.strip() for c in row.find_elements("tag name", "td")]
    cells = cells[cells.index(order):] if order in cells else []
    oo.open_result(driver, wait, row, order)

    info = gx.extract_tab(driver)  # the order opens on Information
    banner = info["banner"]
    f = info["fields"]
    oo.open_tab(driver, wait, "Billing")
    bill = gx.extract_tab(driver)["fields"]
    oo.open_tab(driver, wait, "Milestones")
    ms_rows = []
    for g in gx.extract_tab(driver)["grids"]:
        if "Title" in g["columns"] and "Target Date" in g["columns"]:
            ms_rows += [dict(zip(g["columns"], r)) for r in g["rows"]]
    oo.open_tab(driver, wait, "Teams")
    teams = oo.read_teams(driver)
    oo.open_tab(driver, wait, "Pricing")
    m = TOTAL_RE.search(oo.content_text(driver))
    currency, mrc_existing, mrc_new, mrc_change, one_off = (
        (m.group(1), *(_num(m.group(i)) for i in range(2, 6))) if m else ("", None, None, None, None))
    oo.open_tab(driver, wait, "Notes")
    notes = []
    for g in gx.extract_tab(driver)["grids"]:
        if "Date Created" in g["columns"]:
            notes += [dict(zip(g["columns"], r)) for r in g["rows"]]

    # Header: "Customer: X Site: Y Order Status: Z | Change order # : 10-..."
    m = re.search(r"Customer:\s*(.*?)\s+Site:\s*(.*?)\s+Order Status:\s*(.*?)\s*(\||$)", banner)
    customer, site_label, status = (m.group(1), m.group(2), m.group(3)) if m else ("", "", "")
    m = re.search(r"\|\s*(\w+) order", banner)
    order_type = m.group(1).title() if m else ""

    # Milestones
    def ms(title):
        return next((r for r in ms_rows if r["Title"].startswith(title)), {})
    initial = _d(ms("Initial Target Delivery Date").get("Target Date")) or \
        _d(_field(f, "Initial Target Delivery Date"))
    rtdds = [r for r in ms_rows if r["Title"].startswith("Revised Target Date")]
    current = ([r for r in rtdds if not r.get("Completed Date")] or rtdds or [{}])[-1]
    revised = _d(current.get("Target Date"))
    revised_type = ("Customer reason" if "Customer" in current.get("Title", "") else
                    "Orange Business Services reason" if current else "")
    key_date = revised or initial
    status_rag, days_left, rag_text = rag(key_date)
    skip = ("Initial Target Delivery Date", "Revised Target Date", "Order Raised Date")
    nxt = next((r for r in ms_rows if not r.get("Completed Date") and r.get("Target Date")
                and not r["Title"].startswith(skip)), {})
    raised = _d(ms("Order Raised Date").get("Completed Date"))

    # Site / address
    city = _field(bill, "City", section="Site")
    state = _field(bill, "County/State", section="Site")
    postcode = _field(bill, "Postal/Zipcode", section="Site")
    country = _field(bill, "Country", section="Site")
    local_site = _field(f, "Local Site")
    address = ", ".join(x for x in (local_site, state, country.title() if country else "") if x)

    term = _num(re.sub(r"[^\d.]", "", _field(f, "Order Term (months)")) or "x")
    total_value = (mrc_new * term + (one_off or 0)) if (mrc_new is not None and term) else one_off
    signature = _d(_field(f, "Customer signature date"))
    lt_rag, lt_deadline, lt_margin, lt_text = lead_time(signature, key_date)

    stm = teams.get("STM", {})
    latest = notes[0] if notes else {}
    return {
        "order": order,
        "customer": customer,
        "project": _field(f, "Project", section="Product/Service") or _field(f, "Project"),
        "scope": _field(f, "Additional Comments"),
        "stm": stm.get("primary") or "Not assigned",
        "country": country,
        "city": city,
        "address": address,
        "initial_tdd": initial,
        "revised_tdd": revised,
        "revised_type": revised_type,
        "key_date": key_date,
        "rag": status_rag,
        "days_left": days_left,
        "rag_text": rag_text,
        # order value (Pricing totals; total = new monthly x term + one-off)
        "currency": currency,
        "mrc_existing": mrc_existing,
        "mrc_new": mrc_new,
        "mrc_change": mrc_change,
        "one_off": one_off,
        "term_months": int(term) if term else None,
        "total_value": total_value,
        # 120-day lead time from the customer signature (contract sign) date
        "lt_rag": lt_rag,
        "lt_deadline": lt_deadline,
        "lt_margin": lt_margin,
        "lt_text": lt_text,
        # extras
        "order_status": status,
        "order_type": order_type,
        "product": _field(bill, "Product/Service") or "",
        "site_label": site_label,
        "site_id": cells[5] if len(cells) > 5 else _field(bill, "Site ID", section="Site"),
        "orange_site_ref": cells[4] if len(cells) > 4 else _field(bill, "Orange Site Ref", section="Site"),
        "ban": cells[3] if len(cells) > 3 else "",
        "postcode": postcode,
        "state": state,
        "requested_date": _d(_field(f, "Requested Delivery Date")),
        "signature_date": signature,
        "order_raised": raised,
        "order_age_days": (date.today() - raised).days if raised else None,
        "delivery_owner": (teams.get("Delivery Order Owner") or {}).get("primary") or "",
        "order_entry": (teams.get("Order Entry") or {}).get("primary") or "",
        "tech_support": (teams.get("Technical Support") or {}).get("primary") or "",
        "account_manager": _field(f, "Account Manager"),
        "external_ref": _field(f, "External Reference Number"),
        "previous_order": _field(f, "Previous GOLD Order"),
        "next_milestone": nxt.get("Title", ""),
        "next_milestone_owner": nxt.get("Owner", ""),
        "next_milestone_date": _d(nxt.get("Target Date")),
        "rtdd_history": [(r["Title"].replace("...", ""), r.get("Target Date", ""), r.get("Completed Date", ""))
                         for r in rtdds],
        "latest_note": latest.get("Full note") or latest.get("Note", ""),
        "latest_note_date": latest.get("Last Modification Date", ""),
        "latest_note_by": latest.get("Created By", ""),
        "collected": datetime.now(),
    }


# What the app card and the Excel show, in order: (section, label, key)
LAYOUT = [
    ("Order value", "Order total value", None),
    ("Order value", "New monthly charge", None),
    ("Order value", "Monthly change", None),
    ("Order value", "One-off charge", None),
    ("Order value", "Order term", None),
    ("Order", "Customer", "customer"),
    ("Order", "Project", "project"),
    ("Order", "Order type / status", None),
    ("Order", "Product / service", "product"),
    ("Order", "Scope", "scope"),
    ("Delivery", "Initial Target Delivery Date", "initial_tdd"),
    ("Delivery", "Revised Target Delivery Date", "revised_tdd"),
    ("Delivery", "Revision reason type", "revised_type"),
    ("Delivery", "Delivery status", "rag_text"),
    ("Delivery", "Customer requested date", "requested_date"),
    ("Delivery", "Contract (customer) signature date", "signature_date"),
    ("Delivery", "120-day lead-time deadline", "lt_deadline"),
    ("Delivery", "Lead-time status", "lt_text"),
    ("Delivery", "Next open milestone", None),
    ("Site", "Country", "country"),
    ("Site", "City", "city"),
    ("Site", "Address", "address"),
    ("Site", "Site ID", "site_id"),
    ("Site", "Orange site ref", "orange_site_ref"),
    ("People", "STM responsible", "stm"),
    ("People", "Delivery Order Owner", "delivery_owner"),
    ("People", "Technical Support", "tech_support"),
    ("People", "Order Entry", "order_entry"),
    ("People", "Account Manager", "account_manager"),
    ("References", "External reference", "external_ref"),
    ("References", "Previous GOLD order", "previous_order"),
    ("References", "Order raised", None),
    ("Latest note", "Latest note", None),
]


def fmt(v):
    if isinstance(v, (date, datetime)):
        return gm.gold_date(v if isinstance(v, date) else v.date())
    return "" if v is None else str(v)


def display_value(o, label, key):
    if key:
        return fmt(o.get(key))
    if label == "Order total value":
        return money(o.get("total_value"), o.get("currency"))
    if label == "New monthly charge":
        return money(o.get("mrc_new"), o.get("currency"))
    if label == "Monthly change":
        return money(o.get("mrc_change"), o.get("currency"))
    if label == "One-off charge":
        return money(o.get("one_off"), o.get("currency"))
    if label == "Order term":
        return f"{o['term_months']} months" if o.get("term_months") else ""
    if label == "Order type / status":
        return " / ".join(x for x in (o["order_type"], o["order_status"]) if x)
    if label == "Next open milestone":
        if not o["next_milestone"]:
            return ""
        return f"{o['next_milestone']} - {fmt(o['next_milestone_date'])} ({o['next_milestone_owner']})"
    if label == "Order raised":
        return f"{fmt(o['order_raised'])} ({o['order_age_days']} days ago)" if o["order_raised"] else ""
    if label == "Latest note":
        if not o["latest_note"]:
            return ""
        return f"{o['latest_note_date']} - {o['latest_note_by']}: {o['latest_note']}"
    return ""


# ---------------------------------------------------------------- Excel for management

def export(overviews, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    orange = "FF7900"
    ws["A1"] = "GOLD order overview"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (f"Generated {datetime.now():%d-%b-%Y %H:%M}  ·  {len(overviews)} order(s)  ·  "
                f"Delivery status vs Revised Target Delivery Date (Initial if not revised): Red = passed, "
                f"Amber = {AMBER_DAYS} days or less, Green = more  ·  Lead time: planned delivery vs "
                f"contract signature + {LEAD_TIME_DAYS} days: Red = after, Amber = within {AMBER_DAYS} days "
                f"before, Green = earlier  ·  Order total value = new monthly charge x term + one-off")
    ws["A2"].font = Font(italic=True, color="595959")
    for row_no, (label, key) in enumerate((("Delivery", "rag"), ("120-day lead time", "lt_rag")), 3):
        ws.cell(row=row_no, column=1, value=label).font = Font(bold=True)
        for col, k in enumerate(("Red", "Amber", "Green", "No date"), 2):
            n = sum(1 for o in overviews if o[key] == k)
            c = ws.cell(row=row_no, column=col, value=f"{k}: {n}")
            fill, fnt = RAG_COLORS[k]
            c.fill, c.font = PatternFill("solid", fgColor=fill), Font(bold=True, color=fnt)
            c.alignment = Alignment(horizontal="center")
    totals = {}
    for o in overviews:
        if o.get("total_value") is not None:
            totals[o["currency"]] = totals.get(o["currency"], 0) + o["total_value"]
    ws.cell(row=3, column=7, value="Total order value: " + ("   ".join(
        money(v, cur) for cur, v in totals.items()) or "-")).font = Font(bold=True, size=12, color=orange)

    headers = ["Gold order", "Status", "Days left", "Lead-time status", "Lead-time margin (days)",
               "Order total value", "Currency", "Key date", "Customer", "Project", "Scope", "STM responsible",
               "Country", "City", "Address", "Initial TDD", "Revised TDD", "Revision type",
               "Order type", "Order status", "Product / service", "Delivery Order Owner",
               "Customer requested date", "Next open milestone", "Next milestone date", "Site ID",
               "Orange site ref", "External reference", "New monthly charge", "One-off charge",
               "Order term (months)", "Contract signature date", "120-day deadline", "Latest note"]
    top = 6
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=top, column=i, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=orange)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D9D9D9")
    date_cols = {"Key date", "Initial TDD", "Revised TDD", "Customer requested date", "Next milestone date",
                 "Contract signature date", "120-day deadline"}
    money_cols = {"Order total value", "New monthly charge", "One-off charge"}
    for r, o in enumerate(overviews, top + 1):
        values = [o["order"], o["rag"], o["days_left"], o["lt_rag"], o["lt_margin"],
                  o["total_value"], o["currency"], o["key_date"], o["customer"], o["project"], o["scope"],
                  o["stm"], o["country"], o["city"], o["address"], o["initial_tdd"], o["revised_tdd"],
                  o["revised_type"], o["order_type"], o["order_status"], o["product"], o["delivery_owner"],
                  o["requested_date"], o["next_milestone"], o["next_milestone_date"], o["site_id"],
                  o["orange_site_ref"], o["external_ref"], o["mrc_new"], o["one_off"], o["term_months"],
                  o["signature_date"], o["lt_deadline"],
                  (f"{o['latest_note_date']}: {o['latest_note']}" if o["latest_note"] else "")]
        for i, v in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=headers[i - 1] in ("Scope", "Address", "Latest note"))
            c.border = Border(bottom=thin)
            if headers[i - 1] in date_cols and v:
                c.number_format = "dd-mmm-yyyy"
            if headers[i - 1] in money_cols and v is not None:
                c.number_format = "#,##0.00"
        for cols, key in (((2, 3), "rag"), ((4, 5), "lt_rag")):
            fill, fnt = RAG_COLORS[o[key]]
            for i in cols:
                c = ws.cell(row=r, column=i)
                c.fill, c.font = PatternFill("solid", fgColor=fill), Font(bold=True, color=fnt)
                c.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=6).font = Font(bold=True)
    last = top + max(len(overviews), 1)
    if overviews:
        t = Table(displayName="OrderOverview", ref=f"A{top}:{get_column_letter(len(headers))}{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(t)
    widths = {"Gold order": 13, "Status": 10, "Days left": 10, "Lead-time status": 11,
              "Lead-time margin (days)": 12, "Order total value": 16, "Currency": 9, "Key date": 13,
              "New monthly charge": 14, "One-off charge": 13, "Order term (months)": 10,
              "Contract signature date": 13, "120-day deadline": 13, "Customer": 22, "Project": 28,
              "Scope": 55, "STM responsible": 20, "Country": 14, "City": 16, "Address": 34, "Initial TDD": 13,
              "Revised TDD": 13, "Revision type": 18, "Order type": 11, "Order status": 12,
              "Product / service": 22, "Delivery Order Owner": 20, "Customer requested date": 14,
              "Next open milestone": 24, "Next milestone date": 13, "Site ID": 14, "Orange site ref": 24,
              "External reference": 22, "Latest note": 60}
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 15)
    ws.freeze_panes = ws.cell(row=top + 1, column=2)
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[top].height = 32

    # One detail block per order (easy to read / forward)
    ds = wb.create_sheet("Order details")
    ds.column_dimensions["A"].width = 16
    ds.column_dimensions["B"].width = 30
    ds.column_dimensions["C"].width = 90
    ds.sheet_view.showGridLines = False
    row = 1
    for o in overviews:
        c = ds.cell(row=row, column=1, value=f"{o['order']}  -  {o['customer']}")
        c.font = Font(bold=True, size=13, color=orange)
        fill, fnt = RAG_COLORS[o["rag"]]
        s = ds.cell(row=row, column=3, value=f"Delivery {o['rag']}: {o['rag_text']}")
        s.fill, s.font = PatternFill("solid", fgColor=fill), Font(bold=True, color=fnt)
        row += 1
        lfill, lfnt = RAG_COLORS[o["lt_rag"]]
        s = ds.cell(row=row, column=3, value=f"Lead time {o['lt_rag']}: {o['lt_text']}")
        s.fill, s.font = PatternFill("solid", fgColor=lfill), Font(bold=True, color=lfnt)
        row += 1
        prev = None
        for section, label, key in LAYOUT:
            ds.cell(row=row, column=1, value=section if section != prev else "").font = Font(bold=True, color="595959")
            prev = section
            ds.cell(row=row, column=2, value=label).font = Font(bold=True)
            v = ds.cell(row=row, column=3, value=display_value(o, label, key))
            v.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        if o["rtdd_history"]:
            ds.cell(row=row, column=2, value="RTDD history").font = Font(bold=True)
            ds.cell(row=row, column=3, value="; ".join(
                f"{t} -> {d}" + (f" (closed {c})" if c else " (current)") for t, d, c in o["rtdd_history"]))
            ds.cell(row=row, column=3).alignment = Alignment(wrap_text=True)
            row += 1
        row += 2
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path

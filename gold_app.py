"""GOLD Express: a small desktop app for GOLD orders.

Run:  pythonw gold_app.py   (or double-click "GOLD Express.bat"; packaged as GOLD Express.exe)

- Order explorer: type an order number -> Open; click a tab on the left to see it.
  Screenshot -> <Tab>_<n>.png; Excel -> <Order>.xlsx (one sheet per tab);
  Documents: download; Teams: browse / find a person.
- Revised Target Date: add the RTDD milestone to many orders (Gold order | Date | Reason).
GOLD runs in Firefox in the background (Settings > Show the Firefox window to see it).
"""
import json
import os
import queue
import re
import sys
import threading
import time
import traceback
from datetime import date, datetime
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from PIL import ImageTk
from selenium.webdriver.support.ui import WebDriverWait

import gold_ai as gai
import gold_assignments as gas
import gold_extract as gx
import gold_markdown as gmd
import gold_milestones as gm
import gold_myorders as gmo
import gold_notify as gnotify
import gold_names as gn
import gold_orgchart as goc
import gold_overview as go
import gold_teams as gt
import gold_tracker as gtr
import gold_ui as gui
import gold_update as gu
from gold_grid import EditGrid
import open_order as oo
from open_gold import HELPER_PROFILES_DIR, Wait, APP_NAME, BUNDLE_DIR, GOLD_URL, HERE, start_firefox, wait_for_gold
from version import AUTHOR, __version__

SETTINGS_FILE = HERE / "gold_app_settings.json"
AI_KEY_CACHE = HERE / "ai_key.bin"  # the team's AI key, encrypted for this Windows user
OC_CACHE = HERE / "orgchart_cache.json"  # Org Chart lookups of GOLD team members (kept 7 days)
MY_SNAPSHOT = HERE / "my_orders_snapshot.json"  # last My orders read, to notice changes
CHATS_DIR = HERE / "chats"  # saved Ask GOLD chats
TRACKER_COMMENTS = HERE / "tracker_comments.json"  # Overview > Tracker: comments / next steps per order
OA_FILE = HERE / "apm_open_orders.xlsx"  # the latest APM Open Orders sheet (replaced on every refresh)

# Palette (warm light background, white rounded panels, orange accent) - see gold_ui
ORANGE = gui.ORANGE
ORANGE_DARK = gui.ORANGE_DARK
ORANGE_SOFT = gui.ORANGE_SOFT
INK = gui.INK
MUTED = gui.MUTED
FAINT = gui.FAINT
BG = gui.BG
CANVAS = BG
CARD = gui.CARD
LINE = gui.LINE
HEADER = "#111111"
GREEN = gui.GREEN_FG
RED = gui.RED_FG
FONT = "Segoe UI"


def safe_name(s):
    return re.sub(r"[^\w\-]+", "_", s).strip("_") or "tab"


# ---------------------------------------------------------------- settings

def load_settings():
    s = {"screenshots": str(HERE / "output" / "screenshots"), "exports": str(HERE / "output"),
         "full_page": True, "auto_shot": False, "per_order_folder": True, "show_browser": False,
         "parallel_sessions": True, "ai_key_file": "",
         # Windows notifications
         "notify_on": True, "notify_orders": True, "notify_every": "30", "notify_digest": True,
         "notify_updates": True, "notify_tasks": True, "notified_version": "", "digest_date": "",
         "theme": "light", "save_chats": True, "bot_speed": "fast"}
    try:
        saved = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        saved = {}
    # Older versions kept folders per project: reuse the last project's folders.
    old = saved.pop("projects", {}).get(saved.pop("project", ""), {})
    s.update({k: v for k, v in old.items() if v})
    s.update(saved)
    return s


def save_settings(s):
    try:
        SETTINGS_FILE.write_text(json.dumps(s, indent=2), encoding="utf-8")
    except OSError:
        pass


# ---------------------------------------------------------------- connection

GOLD_HOST = "gold.sso.infra.ftgroup"
NET_ERROR_TEXT = ("unable to connect", "having trouble finding that site", "secure connection failed",
                  "the connection was reset", "the connection has timed out", "server not found",
                  "problem loading page", "network protocol error", "proxy server is refusing")


def network_ok(timeout=3):
    """Can we reach GOLD's server right now? (plain TCP check, no browser)"""
    import socket
    try:
        socket.create_connection((GOLD_HOST, 443), timeout=timeout).close()
        return True
    except OSError:
        return False


def looks_like_connection_error(err):
    text = str(err).lower()
    return any(k in text for k in ("neterror", "connection", "timed out", "timeout", "reset", "dnsnotfound",
                                   "reached error page", "no connection", "session", "browsing context",
                                   "not reachable", "failed to decode response")) or not network_ok()


# ---------------------------------------------------------------- browser worker

class Worker:
    """Owns the Firefox driver; runs GOLD jobs one at a time on its own thread."""

    def __init__(self, app):
        self.app = app
        self.jobs = queue.Queue()
        self.driver = None
        self.wait = None
        self.restart = False  # set when the "show Firefox" setting changes
        self.helpers = []     # extra hidden GOLD sessions for reading many orders side by side
        threading.Thread(target=self._loop, daemon=True).start()

    def submit(self, label, fn, done=None, retry=False):
        """retry=True for read-only jobs: if they fail because of the connection, reconnect to
        GOLD (fresh session) and run them once more. Jobs that save in GOLD are never retried."""
        self.app.set_busy(label)
        self.jobs.put((label, fn, done, retry))

    def _loop(self):
        while True:
            label, fn, done, retry = self.jobs.get()
            if fn is None:
                break
            started = time.time()
            try:
                try:
                    result = fn()
                except Exception as first:
                    if not (retry and self.session_problem(first)):
                        raise
                    self.app.log(f"Connection problem ({str(first).splitlines()[0][:80]}) - reconnecting to GOLD...")
                    self.wait_for_network()
                    self.restart = True
                    result = fn()
                self.app.ui(lambda r=result: (self.app.set_idle(), done and done(r)))
                self.app.ui(lambda label=label, t=time.time() - started: self.app._task_done(label, t, True))
            except Exception as e:
                tb = traceback.format_exc()
                self.app.ui(lambda e=e, tb=tb, label=label: self.app.job_failed(label, e, tb))
                self.app.ui(lambda label=label, t=time.time() - started: self.app._task_done(label, t, False))

    def session_problem(self, err):
        """Did this failure come from a lost connection / expired GOLD session?"""
        return looks_like_connection_error(err) or (self.driver is not None and not self._healthy())

    def wait_for_network(self, timeout=120):
        """Wait (up to timeout seconds) for GOLD's server to be reachable again."""
        deadline = time.time() + timeout
        while not network_ok():
            if time.time() > deadline:
                raise RuntimeError("No connection to GOLD. Check your network / VPN - GOLD Express "
                                   "reconnects by itself when the connection is back.")
            time.sleep(5)

    def _healthy(self, driver=None):
        """Is the open Firefox alive, on GOLD and not stuck on an error page?"""
        driver = driver or self.driver
        try:
            url = driver.current_url
            if url.startswith("about:") and "blank" not in url:
                return False  # about:neterror, about:certerror ...
            if GOLD_HOST not in url:
                return False  # e.g. sign-in page after the session expired
            body = driver.execute_script("return (document.body && document.body.innerText || '')"
                                         ".slice(0, 3000).toLowerCase()")
            return not any(k in body for k in NET_ERROR_TEXT)
        except Exception:
            return False

    def _quit_driver(self):
        if self.driver is not None:
            try:
                self.driver.quit()
            except Exception:
                pass
        self.driver = None

    def ensure_browser(self):
        """Make sure there is a working GOLD session: reuse it, repair it (reload GOLD), or start
        a new one. Called before every GOLD action, so a dropped connection heals by itself."""
        if self.restart:
            self.restart = False
            self._quit_driver()
        if self.driver is not None:
            if self._healthy():
                return
            self.app.log("GOLD session lost - reconnecting...")
            self.wait_for_network()
            try:  # first try to repair the session in place
                self.driver.get(GOLD_URL)
                if wait_for_gold(self.driver, timeout=60):
                    self.app.log("Reconnected to GOLD.")
                    return
            except Exception:
                pass
            self._quit_driver()  # then start a fresh one
        if not network_ok():
            self.wait_for_network(timeout=20)
        show = bool(self.app.settings.get("show_browser"))
        self.app.log("Signing in to GOLD" + ("..." if show else " in the background..."))
        self.driver = start_firefox(headless=not show)
        self.wait = Wait(self.driver, 60)
        self.driver.get(GOLD_URL)
        if wait_for_gold(self.driver, timeout=300 if show else 90):
            return
        if not show:
            # Background sign-in didn't finish: GOLD probably wants a manual login. Show it.
            self.app.log("GOLD needs you to sign in - opening the Firefox window...")
            self._quit_driver()
            self.driver = start_firefox(headless=False)
            self.wait = Wait(self.driver, 60)
            self.driver.get(GOLD_URL)
            if wait_for_gold(self.driver):
                return
        raise RuntimeError("GOLD did not load. Sign in in the Firefox window, then try again.")

    def helper_drivers(self, count):
        """Up to `count` extra hidden GOLD sessions (read-only work), each with its own Firefox
        profile. Healthy ones are reused; missing ones are started side by side (~20 s, once)."""
        alive = []
        for d, i in self.helpers:
            if len(alive) < count and self._healthy(d):
                alive.append((d, i))
            else:
                try:
                    d.quit()
                except Exception:
                    pass
        self.helpers = alive
        missing = count - len(alive)
        if missing > 0:
            used = {i for _, i in alive}
            free = [k for k in range(1, count + len(used) + 2) if k not in used][:missing]
            self.app.log(f"Opening {missing} extra GOLD session(s) to read orders side by side...")
            started, lock = [], threading.Lock()

            def start(i):
                try:
                    d = start_firefox(headless=True, profile_dir=HELPER_PROFILES_DIR / f"session{i}")
                    d.get(GOLD_URL)
                    if wait_for_gold(d, timeout=90):
                        with lock:
                            started.append((d, i))
                    else:
                        d.quit()
                except Exception as e:  # e.g. profile in use: just read with fewer sessions
                    self.app.log(f"  extra session {i} not available ({str(e).splitlines()[0][:70]})")
            threads = [threading.Thread(target=start, args=(i,)) for i in free]
            for t in threads:
                t.start()
            for t in threads:
                t.join()
            self.helpers += started
        return [d for d, _ in self.helpers]

    def drop_helper(self, driver):
        self.helpers = [(d, i) for d, i in self.helpers if d is not driver]
        try:
            driver.quit()
        except Exception:
            pass

    def close(self):
        self.jobs.put(("quit", None, None, False))
        for d, _ in list(self.helpers):
            self.drop_helper(d)
        self._quit_driver()


# ---------------------------------------------------------------- the window

class App(tk.Tk):
    def __init__(self):
        gnotify.set_app_id()  # before the window exists: taskbar + notifications as GOLD Express
        super().__init__()
        self.title(APP_NAME)
        for icon in (BUNDLE_DIR / "gold_express.ico", Path(__file__).parent / "gold_express.ico"):
            if icon.exists():
                self.iconbitmap(default=str(icon))
                break
        self.geometry("1240x800")
        self.minsize(1040, 660)
        self.configure(background=CANVAS)
        self.settings = load_settings()
        self.dpm_names = gn.load()
        self.ui_queue = queue.Queue()
        self.worker = Worker(self)
        self.order = None          # currently open order number
        self.current_tab = None    # name of the GOLD tab shown
        self.tab_titles = []       # side-menu titles of the open order
        self.docs = []             # Documents tab rows
        self.teams = {}            # Teams tab data
        self.buttons = []          # disabled while a job runs
        self.busy = False
        # Settings shown in the Settings window
        self.shot_dir = tk.StringVar(value=self.settings["screenshots"])
        self.export_dir = tk.StringVar(value=self.settings["exports"])
        self.full_page = tk.BooleanVar(value=self.settings["full_page"])
        self.auto_shot = tk.BooleanVar(value=self.settings["auto_shot"])
        self.per_order = tk.BooleanVar(value=self.settings["per_order_folder"])
        self.show_browser = tk.BooleanVar(value=self.settings["show_browser"])
        self.parallel = tk.BooleanVar(value=self.settings["parallel_sessions"])
        self.ai_key_file = tk.StringVar(value=self.settings["ai_key_file"])
        self.notify_on = tk.BooleanVar(value=self.settings["notify_on"])
        self.notify_orders = tk.BooleanVar(value=self.settings["notify_orders"])
        self.notify_every = tk.StringVar(value=str(self.settings["notify_every"]))
        self.notify_digest = tk.BooleanVar(value=self.settings["notify_digest"])
        self.notify_updates = tk.BooleanVar(value=self.settings["notify_updates"])
        self.notify_tasks = tk.BooleanVar(value=self.settings["notify_tasks"])
        self.chat = None           # Ask GOLD conversation (gold_ai.Chat)

        self._style()
        self._build()
        self.report_callback_exception = self._callback_error
        self.protocol("WM_DELETE_WINDOW", self._request_close)
        gui.smooth_window_changes(self)  # no half-drawn frames when maximizing / restoring
        self.after(100, self._pump)
        self.online = True
        threading.Thread(target=self._watch_connection, daemon=True).start()
        self.new_release = None
        self.after(4000, lambda: self._check_updates(quiet=True))  # quiet check at start-up
        # Sign in to GOLD right away, in the background, so the first action doesn't wait for it.
        self.after(300, self._prewarm)
        # Background check of My orders for notifications (only while the user is not working in the app)
        self._last_input = time.time()
        for ev in ("<Any-KeyPress>", "<Any-Button>"):
            self.bind_all(ev, lambda e: setattr(self, "_last_input", time.time()), add="+")
        self.after(60_000, self._my_watch)

    def _prewarm(self):
        if self.busy or self.worker.driver is not None:
            return
        def warm():
            w = self.worker
            w.ensure_browser()
            try:  # also load the order search page now, so the first order opens straight away
                oo.open_orders_tab(w.driver, w.wait, discard_changes=True)
                oo.untick_my_orders(w.driver)
            except Exception:
                pass
            return True
        self.worker.submit("Connecting to GOLD", warm, lambda _: self.log("Connected to GOLD - ready."))

    # ----- thread-safe UI helpers
    def ui(self, fn):
        self.ui_queue.put(fn)

    def _pump(self):
        """Run what the worker asked the UI to do. One failing step must not stop the others (or the
        app would stay 'busy' for good), so each runs on its own."""
        try:
            while True:
                fn = self.ui_queue.get_nowait()
                try:
                    fn()
                except Exception:
                    self._callback_error(*sys.exc_info())
        except queue.Empty:
            pass
        finally:
            self.after(100, self._pump)

    def log(self, msg):
        def add():
            stamp = datetime.now().strftime("%H:%M")
            self.log_box.configure(state="normal")
            self.log_box.insert("end", f"{stamp}  {msg}\n")
            self.log_box.see("end")
            self.log_box.configure(state="disabled")
            self.last_msg.configure(text=msg if len(msg) < 140 else msg[:137] + "...")
        self.ui(add)

    # status pill colours on the light background, by the colour callers give for the state
    PILL = {"#7BD88F": (gui.GREEN_FG, gui.GREEN_BG), "#FF8A80": (gui.RED_FG, gui.RED_BG)}

    def _set_pill(self, text, fg, bg):
        fg, bg = self.PILL.get(fg, (ORANGE_DARK, ORANGE_SOFT))
        self.status.set(f"●  {text}", fg, bg)

    def set_busy(self, label):
        self.busy = True
        self.busy_label = label
        self._set_pill(f"{label}...", ORANGE, "#2A2A2A")
        for b in self.buttons:
            b.state(["disabled"])
        self.configure(cursor="watch")

    # ----- updates (GitHub Releases)
    def _check_updates(self, quiet=True):
        """quiet: background check at start-up - only lights the button up when there is an update."""
        if self.new_release and not quiet:
            return self._offer_update(self.new_release)

        def work():
            try:
                rel, newer = gu.check()
            except Exception as e:
                if not quiet:
                    self.ui(lambda e=e: self._message(
                        "Updates", "Could not check for updates",
                        f"{str(e).splitlines()[0][:160]}\n\nYou can also download new versions from "
                        f"{gu.RELEASES_PAGE}", icon=""))
                return
            if newer:
                self.new_release = rel
                self.ui(lambda: self.update_btn.configure(text=f"↑  Update to {rel['version']}",
                                                          style="Accent.TButton"))
                self.log(f"A new version of {APP_NAME} is available: {rel['version']}")
                if self.settings.get("notified_version") != rel["version"]:  # once per version
                    self.settings["notified_version"] = rel["version"]
                    self.ui(lambda: (save_settings(self.settings), self._notify(
                        f"GOLD Express {rel['version']} is available",
                        "Click Update in GOLD Express's side menu to install it.", "updates")))
                if not quiet:
                    self.ui(lambda: self._offer_update(rel))
            elif not quiet:
                self.ui(lambda: self._message(
                    "Updates", "You're up to date",
                    f"GOLD Express {__version__} is the newest version "
                    f"(newest on GitHub: {rel['version'] if rel else '?'}).", icon=""))
        threading.Thread(target=work, daemon=True).start()

    def _offer_update(self, rel):
        if getattr(self, "_updating", False):
            if getattr(self, "_update_win", None) and self._update_win.winfo_exists():
                self._update_win.lift()  # the update window is already open
            return
        if self.busy:  # e.g. still connecting to GOLD: offer the update as soon as that has finished
            if not getattr(self, "_update_waiting", False):
                self._update_waiting = True
                self.update_btn.configure(text="↑  Update when ready...")
                self.log("The update will start as soon as the current action has finished.")
                self._update_when_idle(rel)
            return
        if not gu.can_self_update():
            self._message("Update available", f"Version {rel['version']} is available.",
                          "This copy runs from the source code - update it with 'git pull'.\n\n"
                          f"{gu.RELEASES_PAGE}", icon="")
            return
        self._updating = True
        self._update_dialog(rel)

    def _update_when_idle(self, rel):
        if self.busy:
            self.after(500, self._update_when_idle, rel)
            return
        self._update_waiting = False
        self.update_btn.configure(text=f"↑  Update to {rel['version']}")
        self._offer_update(rel)

    # ----- app-styled small windows
    def _dialog(self, title, width=560):
        """A small window in the app's style: a rounded white tile on the warm background.
        Returns (window, body frame)."""
        win = tk.Toplevel(self, background=BG)
        win.title(f"{APP_NAME} - {title}")
        win.transient(self)
        win.resizable(False, False)
        tile = gui.Tile(win, pad=14, radius=20)
        tile.pack(fill="both", expand=True, padx=10, pady=10)
        tile.body.configure(width=int(width * self.winfo_fpixels("1i") / 96))
        return win, tile.body

    def _place_dialog(self, win):
        win.update_idletasks()
        x = self.winfo_rootx() + (self.winfo_width() - win.winfo_reqwidth()) // 2
        y = self.winfo_rooty() + max(40, (self.winfo_height() - win.winfo_reqheight()) // 3)
        win.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        win.lift()
        win.focus_force()

    def _dialog_head(self, body, title, sub, icon="", badge=None):
        head = tk.Frame(body, background=CARD)
        head.pack(fill="x")
        logo = self._asset_image("logo_48.png")
        if logo:
            lab = tk.Label(head, image=logo, bg=CARD)
            lab.image = logo
        else:
            lab = tk.Label(head, text=icon, bg=CARD, fg=ORANGE, font=(gui.ICONS, 22))
        lab.pack(side="left", padx=(0, 14))
        words = tk.Frame(head, background=CARD)
        words.pack(side="left", fill="x", expand=True)
        tk.Label(words, text=title, bg=CARD, fg=INK, font=(FONT, 14, "bold"), anchor="w").pack(fill="x")
        tk.Label(words, text=sub, bg=CARD, fg=MUTED, font=(FONT, 9), anchor="w", justify="left",
                 wraplength=420).pack(fill="x")
        if badge:
            gui.Pill(head, badge, ORANGE_DARK, ORANGE_SOFT, bg=CARD).pack(side="right", anchor="n")
        return head

    def _message(self, title, heading, text, icon="", button="OK"):
        """An app-styled message with one button (replaces plain Windows message boxes)."""
        win, body = self._dialog(title, width=460)
        self._dialog_head(body, heading, text, icon)
        row = tk.Frame(body, background=CARD)
        row.pack(fill="x", pady=(16, 0))
        ttk.Button(row, text=button, style="Accent.TButton", command=win.destroy).pack(side="right")
        win.bind("<Escape>", lambda e: win.destroy())
        win.bind("<Return>", lambda e: win.destroy())
        self._place_dialog(win)
        return win

    # ----- light / dark mode
    def _set_theme(self, name):
        """Save the theme; the colours are set when the app starts, so offer to restart now."""
        if name == gui.THEME:
            return
        self.settings["theme"] = name
        self._store_settings()
        win, body = self._dialog("Appearance", width=440)
        self._dialog_head(body, f"Switch to {name} mode?",
                          "GOLD Express restarts to change its colours (a few seconds). Nothing is lost - "
                          "your settings and saved chats stay.", "")
        row = tk.Frame(body, background=CARD)
        row.pack(fill="x", pady=(16, 0))
        ttk.Button(row, text="Restart now", style="Accent.TButton",
                   command=lambda: (win.destroy(), self._restart())).pack(side="right")
        ttk.Button(row, text="Later", command=win.destroy).pack(side="right", padx=(0, 8))
        self._place_dialog(win)

    def _when_idle(self, what, then):
        """Run then() now, or - while a task is running (a bulk update may be saving in GOLD) - once it has
        finished, so an order is never left half saved."""
        if not self.busy:
            then()
            return
        self.log(f"{what} when the current task has finished...")
        self._set_pill(f"{what} after this task...", ORANGE, "#2A2A2A")

        def wait():
            if self.busy:
                self.after(500, wait)
            else:
                then()
        self.after(500, wait)

    def _request_close(self):
        """The window's X: while a task runs, ask first, then stop a bulk run after its current order."""
        if self.busy:
            if not messagebox.askyesno(
                    APP_NAME, f"GOLD Express is still working:\n{getattr(self, 'busy_label', 'a task')}.\n\n"
                              "Close anyway? A bulk update stops after the order it is saving now, then "
                              "GOLD Express closes.", icon="warning"):
                return
            self.bulk_stop_flag.set()
        self._when_idle("Closing", self._on_close)

    def _restart(self):
        """Start a new GOLD Express that waits for this one to close, then close (after any running task)."""
        if self.busy:
            self._when_idle("Restarting", self._restart)
            return
        import subprocess
        import sys as _sys
        args = [_sys.executable] + ([] if getattr(_sys, "frozen", False) else [str(Path(__file__).resolve())])
        subprocess.Popen(args + ["--after-pid", str(os.getpid())], close_fds=True,
                         creationflags=0x00000008 | 0x00000200)  # DETACHED_PROCESS | CREATE_NEW_PROCESS_GROUP
        self._on_close()

    # ----- the update window
    def _update_dialog(self, rel):
        win, body = self._dialog("Update", width=560)
        self._update_win = win
        self._dialog_head(body, "A new version is ready",
                          f"GOLD Express {rel['version']}  ·  you have {__version__}", "",
                          badge=f"v{rel['version']}")
        tk.Label(body, text="WHAT'S NEW", bg=CARD, fg=FAINT, font=(FONT, 8, "bold"), anchor="w").pack(
            fill="x", pady=(16, 4))
        notes = tk.Text(body, height=9, wrap="word", relief="flat", background=gui.FIELD, foreground=INK,
                        padx=12, pady=8, highlightthickness=1, highlightbackground=LINE, borderwidth=0,
                        font=(FONT, 9), cursor="arrow")
        notes.pack(fill="both", expand=True)
        gmd.setup(notes, font=FONT, ink=INK, muted=MUTED, accent=ORANGE, soft=ORANGE_SOFT, line=LINE, size=9)
        gmd.render(notes, rel.get("notes") or "- Improvements and fixes.", 500)
        notes.configure(state="disabled")
        tk.Label(body, text="GOLD Express closes, updates itself and opens again (about a minute). "
                            "Your settings and files are kept.", bg=CARD, fg=MUTED, font=(FONT, 8),
                 anchor="w", justify="left", wraplength=520).pack(fill="x", pady=(10, 0))

        foot = tk.Frame(body, background=CARD)
        foot.pack(fill="x", pady=(14, 0))
        status = tk.Label(foot, text="", bg=CARD, fg=MUTED, font=(FONT, 9), anchor="w")
        bar = gui.Progress(foot, height=10)
        later = ttk.Button(foot, text="Later", style="TButton")
        go_btn = ttk.Button(foot, text="↑  Update now", style="Accent.TButton")
        go_btn.pack(side="right")
        later.pack(side="right", padx=(0, 8))
        state = {"downloading": False}

        def close():
            if state["downloading"]:
                return  # the download finishes, then the app restarts
            self._updating = False
            win.destroy()

        def progress(done, total):
            frac = done / total if total else 0
            text = f"Downloading...  {done * 100 // total}%" if total else f"Downloading...  {done // 1_000_000} MB"
            self.ui(lambda: (bar.set(frac), status.configure(text=text)))

        def failed(err):
            state["downloading"] = False
            self.set_idle()
            bar.pack_forget()
            status.configure(text=f"The download failed: {str(err).splitlines()[0][:90]}", fg=RED)
            go_btn.configure(text="Try again")
            go_btn.state(["!disabled"])
            later.configure(text="Close")
            later.state(["!disabled"])

        def start():
            if self.busy:  # the installer closes the app: never in the middle of a task
                status.configure(text="Waiting for the current task to finish...", fg=MUTED)
                status.pack(side="left")
                go_btn.state(["disabled"])
                self._when_idle("Updating", start)
                return
            state["downloading"] = True
            go_btn.state(["disabled"])
            later.state(["disabled"])
            status.configure(text="Downloading...", fg=MUTED)
            status.pack(side="left")
            bar.pack(side="left", fill="x", expand=True, padx=(12, 16))
            self.set_busy("Downloading update")

            def work():
                try:
                    setup = gu.download(rel, progress)
                except Exception as e:
                    self.ui(lambda e=e: failed(e))
                    return
                self.log(f"Installing {rel['name']}...")
                self.ui(lambda: (bar.set(1), status.configure(text="Installing - GOLD Express will reopen..."),
                                 self.after(600, lambda: self._install_update(setup))))
            threading.Thread(target=work, daemon=True).start()

        go_btn.configure(command=start)
        later.configure(command=close)
        win.protocol("WM_DELETE_WINDOW", close)
        win.bind("<Escape>", lambda e: close())
        self._place_dialog(win)

    def _install_update(self, setup):
        try:
            gu.run_installer(setup)
        except Exception as e:
            self._updating = False
            self.set_idle()
            if getattr(self, "_update_win", None) and self._update_win.winfo_exists():
                self._update_win.destroy()
            self._message("Update", "The update could not start", str(e), icon="")
            return
        self._on_close()  # the setup replaces the app's files and starts it again

    def _watch_connection(self, every=20):
        """Background thread: notice when the network drops and comes back. When it is back,
        replace the (now dead) GOLD session with a fresh one so the app just keeps working."""
        while True:
            time.sleep(every)
            ok = network_ok()
            if ok == self.online:
                continue
            self.online = ok
            if not ok:
                self.log("Connection to GOLD lost - waiting for the network to come back...")
                self.ui(lambda: None if self.busy else self._set_pill("Offline - waiting for network", "#FF8A80",
                                                                      "#2A2A2A"))
                continue
            self.log("Connection is back.")
            if self.worker.driver is None:
                self.ui(lambda: None if self.busy else self.set_idle())
                continue
            self.worker.restart = True  # the old session is gone: start a fresh one

            def reconnect():
                if self.busy:
                    return  # the running job reconnects by itself (ensure_browser)
                self.worker.submit("Reconnecting to GOLD", lambda: (self.worker.ensure_browser(), "ok")[1],
                                   lambda _: self.log("New GOLD session ready."))
            self.ui(reconnect)

    def set_idle(self):
        self.busy = False
        for b in ("rtdd_stop_btn", "assign_stop_btn"):
            if hasattr(self, b):
                getattr(self, b).state(["disabled"])
        if getattr(self, "online", True):
            self._set_pill("Ready", "#7BD88F", "#2A2A2A")
        else:
            self._set_pill("Offline - waiting for network", "#FF8A80", "#2A2A2A")
        for b in self.buttons:
            b.state(["!disabled"])
        if hasattr(self, "my_empty_hint") and not self.my_time:  # e.g. reading My orders failed
            self.my_empty_hint.configure(text="Click below to read your orders from GOLD (My orders only).")
        self.configure(cursor="")
        self._refresh_action_state()

    def _callback_error(self, exc, err, tb):
        """An error in a button / window callback: the packaged app has no console, so log it and say so
        instead of silently doing nothing."""
        text = "".join(traceback.format_exception(exc, err, tb))
        try:
            (HERE / "gold_app_error.log").write_text(text, encoding="utf-8")
        except OSError:
            pass
        self.log(f"ERROR: {err}")
        messagebox.showerror(APP_NAME, f"Something went wrong:\n\n{err}\n\nDetails: {HERE / 'gold_app_error.log'}")

    # ----- Windows notifications
    QUIET_TASKS = ("Connecting to GOLD", "Checking my orders", "Reading my orders")
    TASK_NOTIFY_S = 20  # tell about tasks that took at least this long

    def _notify(self, title, text, kind, force=False):
        """Show a Windows notification if the user wants this kind (Settings > Notifications)."""
        wanted = {"orders": self.notify_orders, "digest": self.notify_digest, "updates": self.notify_updates,
                  "tasks": self.notify_tasks}.get(kind)
        if not force and (not self.notify_on.get() or (wanted is not None and not wanted.get())):
            return
        gnotify.notify(title, text, tag=kind)
        self.log(f"Notification: {title} - {text}")

    def _app_in_front(self):
        try:
            return self.focus_displayof() is not None
        except (KeyError, tk.TclError):
            return False

    def _task_done(self, label, seconds, ok):
        """A worker task has finished: notify when it was long and the user is in another window."""
        if seconds < self.TASK_NOTIFY_S or label.startswith(self.QUIET_TASKS) or self._app_in_front():
            return
        self._notify(f"{'Done' if ok else 'Failed'}: {label}",
                     "GOLD Express has finished." if ok else "Open GOLD Express to see what went wrong.", "tasks")

    def job_failed(self, label, err, tb):
        self.set_idle()
        self.log(f"ERROR while {label.lower()}: {err}")
        try:
            (HERE / "gold_app_error.log").write_text(tb, encoding="utf-8")
        except OSError:
            pass  # the log is locked (open, OneDrive...): still tell the user
        messagebox.showerror(APP_NAME, f"{label} failed:\n\n{err}")

    # ----- look
    def _style(self):
        st = ttk.Style(self)
        st.theme_use("clam")
        self.option_add("*Font", (FONT, 10))
        self.option_add("*TCombobox*Listbox.font", (FONT, 10))
        self.option_add("*TCombobox*Listbox.selectBackground", ORANGE)
        # drop-down lists and right-click menus follow the theme too
        for pattern, value in (("*TCombobox*Listbox.background", CARD), ("*TCombobox*Listbox.foreground", INK),
                               ("*TCombobox*Listbox.selectForeground", "white"), ("*Menu.background", CARD),
                               ("*Menu.foreground", INK), ("*Menu.activeBackground", ORANGE_SOFT),
                               ("*Menu.activeForeground", INK), ("*Text.insertBackground", INK)):
            self.option_add(pattern, value)
        st.configure(".", background=CARD, foreground=INK, font=(FONT, 10), bordercolor=LINE,
                     lightcolor=LINE, darkcolor=LINE, focuscolor=ORANGE)
        st.configure("TFrame", background=CARD)
        st.configure("Canvas.TFrame", background=CANVAS)
        st.configure("Card.TFrame", background=CARD, relief="solid", borderwidth=1)
        st.configure("TLabel", background=CARD, foreground=INK)
        st.configure("Card.TLabel", background=CARD)
        st.configure("Muted.TLabel", background=CARD, foreground=MUTED)
        st.configure("CardMuted.TLabel", background=CARD, foreground=MUTED)
        st.configure("Head.TLabel", background=CARD, font=(FONT, 11, "bold"))
        st.configure("CardHead.TLabel", background=CARD, font=(FONT, 11, "bold"))
        st.configure("Big.TLabel", background=CARD, font=(FONT, 13, "bold"))
        # Buttons
        st.configure("TButton", background="#F6F7F9", foreground=INK, borderwidth=1, relief="solid",
                     padding=(12, 6), bordercolor="#CFD3D9", lightcolor="#F6F7F9", darkcolor="#F6F7F9")
        st.map("TButton", background=[("disabled", "#F6F7F9"), ("pressed", "#E4E6EA"), ("active", "#ECEEF1")],
               bordercolor=[("active", ORANGE), ("focus", "#CFD3D9")],
               lightcolor=[("active", "#ECEEF1")], darkcolor=[("active", "#ECEEF1")],
               foreground=[("disabled", "#A0A4AB")])
        st.configure("Accent.TButton", background=ORANGE, foreground="white", bordercolor=ORANGE,
                     lightcolor=ORANGE, darkcolor=ORANGE, font=(FONT, 10, "bold"), padding=(16, 6))
        st.map("Accent.TButton", background=[("disabled", "#F7C49A"), ("pressed", ORANGE_DARK),
                                             ("active", ORANGE_DARK)],
               bordercolor=[("disabled", "#F7C49A"), ("active", ORANGE_DARK)],
               lightcolor=[("disabled", "#F7C49A"), ("active", ORANGE_DARK)],
               darkcolor=[("disabled", "#F7C49A"), ("active", ORANGE_DARK)],
               foreground=[("disabled", "white")])
        st.configure("Header.TButton", background=HEADER, foreground="white", bordercolor="#3A3A3A",
                     lightcolor=HEADER, darkcolor=HEADER, padding=(12, 5))
        st.map("Header.TButton", background=[("active", "#2A2A2A"), ("pressed", "#000000")],
               lightcolor=[("active", "#2A2A2A")], darkcolor=[("active", "#2A2A2A")],
               bordercolor=[("active", ORANGE)])
        # Segmented buttons: radio buttons drawn without the round indicator.
        st.layout("Seg.TRadiobutton", [("Radiobutton.border", {"sticky": "nswe", "children": [
            ("Radiobutton.padding", {"sticky": "nswe", "children": [
                ("Radiobutton.label", {"sticky": "nswe"})]})]})])
        st.configure("Seg.TRadiobutton", background="#F6F7F9", foreground=INK, padding=(16, 7),
                     borderwidth=1, relief="solid", bordercolor="#CFD3D9", font=(FONT, 10, "bold"))
        st.map("Seg.TRadiobutton", background=[("selected", ORANGE), ("active", ORANGE_SOFT)],
               foreground=[("selected", "white")])
        st.configure("TCheckbutton", background=CARD)
        st.configure("Card.TCheckbutton", background=CARD)
        for cb in ("TCheckbutton", "Card.TCheckbutton"):
            st.configure(cb, indicatorbackground=CARD, indicatorforeground="white",
                         indicatormargin=(0, 0, 8, 0), padding=(0, 3))
            st.map(cb, indicatorbackground=[("selected", ORANGE), ("active", ORANGE_SOFT)],
                   indicatorforeground=[("selected", "white")])
        # Inputs
        st.configure("TEntry", fieldbackground=CARD, padding=6, bordercolor=LINE)
        st.map("TEntry", bordercolor=[("focus", ORANGE)], lightcolor=[("focus", ORANGE)])
        st.configure("TCombobox", fieldbackground=CARD, padding=4, arrowsize=14)
        st.map("TCombobox", fieldbackground=[("readonly", CARD)], bordercolor=[("focus", ORANGE)])
        # Tabs
        st.configure("TNotebook", background=CANVAS, borderwidth=0, tabmargins=(0, 6, 0, 0))
        st.configure("TNotebook.Tab", background=CANVAS, foreground=MUTED, padding=(14, 8),
                     font=(FONT, 10, "bold"), borderwidth=0)
        st.map("TNotebook.Tab", background=[("selected", CARD)], foreground=[("selected", ORANGE)],
               expand=[("selected", (0, 0, 0, 0))])
        # Tables
        st.configure("Treeview", background=CARD, fieldbackground=CARD, foreground=INK, rowheight=28,
                     borderwidth=0, font=(FONT, 10))
        st.map("Treeview", background=[("selected", ORANGE_SOFT)], foreground=[("selected", INK)])
        st.configure("Treeview.Heading", background="#F7F7F8", foreground=MUTED, relief="flat",
                     font=(FONT, 9, "bold"), padding=(8, 6), borderwidth=0)
        st.map("Treeview.Heading", background=[("active", "#EFEFF1")])
        st.configure("Grid.Treeview", rowheight=28)
        st.configure("TPanedwindow", background=CARD)
        st.configure("Sash", sashthickness=6, gripcount=0, background=CARD)
        st.configure("Vertical.TScrollbar", background="#E5E7EB", troughcolor=CARD, borderwidth=0,
                     arrowsize=12)
        st.configure("Horizontal.TScrollbar", background="#E5E7EB", troughcolor=CARD, borderwidth=0,
                     arrowsize=12)
        # Softer tables, rounded buttons and fields
        st.configure("Treeview", rowheight=30, foreground=INK)
        st.configure("Treeview.Heading", background=gui.HEAD_BG, foreground=MUTED, padding=(8, 7))
        st.map("Treeview.Heading", background=[("active", gui.HEAD_ACTIVE)])
        gui.style_rounded(st, self)

    def _btn(self, parent, text, cmd, style=None, **grid):
        b = ttk.Button(parent, text=text, command=cmd, style=style or "TButton")
        self.buttons.append(b)
        return b

    def _card(self, parent, **pack):
        c = ttk.Frame(parent, style="Card.TFrame", padding=(14, 12))
        c.pack(**pack)
        return c

    # ----- layout
    def _build(self):
        # Left: side menu. Right: page title bar, the page in a rounded panel, status line.
        self.configure(background=BG)
        right = tk.Frame(self, background=BG)
        side = tk.Frame(self, background=CARD)
        side.pack(side="left", fill="y")
        right.pack(side="left", fill="both", expand=True)

        # Page title bar
        top = tk.Frame(right, background=BG, padx=24, pady=16)
        top.pack(fill="x")
        titles = tk.Frame(top, background=BG)
        titles.pack(side="left")
        self.page_title = tk.Label(titles, text="", bg=BG, fg=INK, font=("Segoe UI Light", 20), anchor="w")
        self.page_title.pack(anchor="w")
        self.page_sub = tk.Label(titles, text="", bg=BG, fg=MUTED, font=(FONT, 9), anchor="w")
        self.page_sub.pack(anchor="w")
        self.status = gui.Pill(top, "Ready")
        self.status.pack(side="right", pady=(4, 0))
        self._set_pill("Ready", "#7BD88F", "#2A2A2A")

        # Status line (bottom): last message + activity toggle; the log panel above it
        bottom = tk.Frame(right, background=BG)
        bottom.pack(side="bottom", fill="x", padx=24, pady=(0, 10))
        self.last_msg = tk.Label(bottom, text="Ready.", bg=BG, fg=MUTED, anchor="w", font=(FONT, 9))
        self.last_msg.pack(side="left", fill="x", expand=True)
        self.log_toggle = tk.Label(bottom, text="Show activity ▲", bg=BG, fg=ORANGE, cursor="hand2",
                                   font=(FONT, 9, "bold"))
        self.log_toggle.pack(side="right")
        self.log_toggle.bind("<Button-1>", lambda e: self._toggle_log())
        self.log_frame = gui.Tile(right, pad=4)
        self.log_box = tk.Text(self.log_frame.body, height=7, state="disabled", font=("Consolas", 9),
                               background=CARD, foreground=INK, relief="flat", padx=6, pady=2,
                               highlightthickness=0, borderwidth=0)
        self.log_box.pack(fill="both", expand=True)

        # The page: its parts sit in rounded tiles on the background
        content_host = tk.Frame(right, background=BG)
        content_host.pack(fill="both", expand=True, padx=(12, 18), pady=(0, 4))

        # Side menu: logo, pages by section, then updates / settings
        nav = self.notebook = gui.SideNav(side, content_host, on_change=self._page_changed)
        nav.pack(fill="y", expand=True)
        brand = tk.Frame(nav.top, background=CARD, padx=20, pady=20)
        brand.pack(fill="x")
        self._logo_img = self._asset_image("logo_36.png")
        if self._logo_img:
            tk.Label(brand, image=self._logo_img, bg=CARD).pack(side="left")
        words = tk.Frame(brand, background=CARD)
        words.pack(side="left", padx=(10, 0))
        name = tk.Frame(words, background=CARD)
        name.pack(anchor="w")
        tk.Label(name, text="GOLD", bg=CARD, fg=INK, font=(FONT, 13, "bold")).pack(side="left")
        tk.Label(name, text=" Express", bg=CARD, fg=ORANGE, font=(FONT, 13)).pack(side="left")
        tk.Label(words, text="Order Lifecycle Delivery", bg=CARD, fg=FAINT, font=(FONT, 8)).pack(anchor="w")

        def page(pad=True):
            return tk.Frame(content_host, background=BG)

        my, overview, explorer, ask = page(), page(), page(), page(False)
        rtdd, assign, oa = page(), page(), page()
        org, names = page(False), page()
        self.org_frame, self.explorer_frame, self.overview_frame = org, explorer, overview
        self.my_frame, self.oa_frame = my, oa
        nav.section("Orders")
        nav.add(my, "My Orders", "", "Your orders in GOLD (My orders only) at a glance")
        nav.add(oa, "Order Assignment", "",
                "Which DPM has which APM order - from the APM Open Orders sheet (refreshed daily)")
        nav.add(overview, "Overview", "", "Key facts and delivery status of your orders")
        nav.add(explorer, "Order explorer", "", "Open an order, browse its tabs, export and download")
        nav.add(ask, "GOLD bot", "", "The Friendly GOLD bot - ask about deadlines, pricing, sites and people")
        nav.section("Bulk updates")
        nav.add(rtdd, "Revised Target Date", "", "Add the Revised Target Date milestone to many orders")
        nav.add(assign, "Bulk assignment", "", "Make a DPM the primary STM of many orders")
        nav.section("People")
        nav.add(org, "Org Chart", "", "Find anyone at Orange Business and who they report to")
        nav.add(names, "DPM names", "", "The DPM list used by Order Assignment and Bulk assignment")

        self.update_btn = ttk.Button(nav.bottom, text="↻  Check for updates", style="Side.TButton",
                                     command=lambda: self._check_updates(quiet=False))
        self.update_btn.pack(fill="x", pady=(0, 2))
        ttk.Button(nav.bottom, text="⚙  Settings", style="Side.TButton",
                   command=self._open_settings).pack(fill="x")
        ttk.Button(nav.bottom, text="☼  Light mode" if gui.DARK else "☾  Dark mode", style="Side.TButton",
                   command=lambda: self._set_theme("light" if gui.DARK else "dark")).pack(fill="x", pady=(2, 0))
        tk.Label(nav.bottom, text=f"Version {__version__}  ·  © {AUTHOR}", bg=CARD, fg=FAINT,
                 font=(FONT, 8), anchor="w").pack(fill="x", padx=12, pady=(10, 0))
        self.bulk_stop_flag = threading.Event()
        self._build_my_tab(my)
        self._build_oa_tab(oa)
        self._build_overview_tab(overview)
        self._build_explorer(explorer)
        self._build_rtdd_tab(rtdd)
        self._build_assign_tab(assign)
        self._build_names_tab(names)
        self._build_org_tab(org)
        self._build_ask_tab(ask)
        self._refresh_action_state()

    def _page_changed(self, title, subtitle):
        self.page_title.configure(text=title)
        self.page_sub.configure(text=subtitle)
        if title == "Order Assignment" and hasattr(self, "oa_tree"):
            self.after(50, self._oa_shown)

    def _asset_path(self, name):
        for base in (BUNDLE_DIR / "assets", Path(__file__).parent / "assets"):
            if (base / name).exists():
                return base / name
        return None

    def _asset_image(self, name):
        for base in (BUNDLE_DIR / "assets", Path(__file__).parent / "assets"):
            if (base / name).exists():
                return tk.PhotoImage(file=str(base / name), master=self)
        return None

    def _toggle_log(self):
        if self.log_frame.winfo_ismapped():
            self.log_frame.pack_forget()
            self.log_toggle.configure(text="Show activity ▲")
        else:
            self.log_frame.pack(side="bottom", fill="x", padx=(12, 18))
            self.log_toggle.configure(text="Hide activity ▼")

    def _tile(self, parent, pad=10, **pack):
        """A rounded white tile; widgets go in its .body. Packed when pack options are given."""
        t = gui.Tile(parent, pad=pad)
        if pack:
            t.pack(**pack)
        return t

    def _build_explorer(self, f):
        # Order bar
        top = self._tile(f, pad=6, fill="x")
        bar = top.body
        tk.Label(bar, text="", bg=CARD, fg=ORANGE, font=(gui.ICONS, 13)).pack(side="left", padx=(4, 8))
        self.order_var = tk.StringVar()
        ent = ttk.Entry(bar, textvariable=self.order_var, width=16, font=(FONT, 10))
        ent.pack(side="left", padx=(0, 8))
        ent.bind("<Return>", lambda e: self.open_order())
        ent.focus_set()
        self._btn(bar, "Open order", self.open_order, "Accent.TButton").pack(side="left")
        self.order_label = tk.Label(bar, text="Type an order number, e.g. 10-1771012", bg=CARD, fg=MUTED,
                                    font=(FONT, 9))
        self.order_label.pack(side="left", padx=16)

        main = ttk.Panedwindow(f, orient="horizontal")
        main.pack(fill="both", expand=True)

        lt = self._tile(main)
        left = lt.body
        tk.Label(left, text="ORDER TABS", bg=CARD, fg=FAINT, font=(FONT, 8, "bold")).pack(anchor="w", pady=(0, 6))
        self.tab_list = tk.Listbox(left, width=20, activestyle="none", font=(FONT, 10),
                                   selectbackground=ORANGE_SOFT, selectforeground=ORANGE_DARK,
                                   background=CARD, foreground=INK, borderwidth=0,
                                   highlightthickness=0, relief="flat")
        self.tab_list.pack(fill="both", expand=True)
        self.tab_list.bind("<<ListboxSelect>>", lambda e: self._tab_clicked())
        main.add(lt, weight=0)

        rt = self._tile(main)
        right = rt.body
        main.add(rt, weight=1)
        head = tk.Frame(right, background=CARD)
        head.pack(fill="x")
        self.tab_title = ttk.Label(head, text="No order open", style="CardHead.TLabel",
                                   font=(FONT, 13, "bold"))
        self.b_all = self._btn(head, "⤓  All tabs", self.export_all)
        self.b_all.pack(side="right")
        self.b_xls = self._btn(head, "⤓  This tab", self.export_tab)
        self.b_xls.pack(side="right", padx=6)
        self.b_shot = self._btn(head, "Screenshot", self.screenshot)
        self.b_shot.pack(side="right")
        self.tab_title.pack(side="left")  # after the buttons: they keep their size

        self.content = tk.Frame(right, background=CARD, pady=8)
        self.content.pack(fill="both", expand=True)
        self._build_text_view()
        self._build_docs_view()
        self._build_teams_view()
        self._show(self.text_view)

    # ----- settings window
    def _open_settings(self):
        if getattr(self, "_settings_win", None) and self._settings_win.winfo_exists():
            self._settings_win.lift()
            return
        win = self._settings_win = tk.Toplevel(self, background=CARD)
        win.title(f"{APP_NAME} - Settings")
        win.transient(self)
        win.resizable(False, False)
        body = ttk.Frame(win, style="Card.TFrame", padding=(20, 16))
        body.pack(fill="both", expand=True)
        was_shown = self.show_browser.get()

        ttk.Label(body, text="Folders", style="CardHead.TLabel").grid(row=0, column=0, sticky="w",
                                                                      columnspan=4)
        for r, (label, var) in enumerate([("Screenshots", self.shot_dir),
                                          ("Excel & downloads", self.export_dir)], 1):
            ttk.Label(body, text=label, style="Card.TLabel").grid(row=r, column=0, sticky="w", pady=4)
            ttk.Entry(body, textvariable=var, width=62).grid(row=r, column=1, sticky="ew", padx=8)
            ttk.Button(body, text="Browse...", command=lambda v=var: self._browse(v)).grid(row=r, column=2)
            ttk.Button(body, text="Open", command=lambda v=var: self._open_folder(v.get())).grid(
                row=r, column=3, padx=(6, 0))

        ttk.Label(body, text="Screenshots", style="CardHead.TLabel").grid(
            row=3, column=0, sticky="w", columnspan=4, pady=(16, 2))
        for r, (text, var) in enumerate([
                ("Capture the whole page (not just the visible part)", self.full_page),
                ("Take a screenshot automatically when a tab opens", self.auto_shot),
                ("Put each order's screenshots in its own sub-folder", self.per_order)], 4):
            ttk.Checkbutton(body, text=text, variable=var, style="Card.TCheckbutton").grid(
                row=r, column=0, columnspan=4, sticky="w", pady=2)

        ttk.Label(body, text="Browser", style="CardHead.TLabel").grid(
            row=7, column=0, sticky="w", columnspan=4, pady=(16, 2))
        ttk.Checkbutton(body, text="Show the Firefox window (normally GOLD runs hidden in the background)",
                        variable=self.show_browser, style="Card.TCheckbutton").grid(
            row=8, column=0, columnspan=4, sticky="w", pady=2)
        ttk.Checkbutton(body, text="Overview: read 3+ orders side by side (up to 3 GOLD sessions - much faster)",
                        variable=self.parallel, style="Card.TCheckbutton").grid(
            row=9, column=0, columnspan=4, sticky="w", pady=2)

        ttk.Label(body, text="AI assistant (Ask GOLD)", style="CardHead.TLabel").grid(
            row=10, column=0, sticky="w", columnspan=4, pady=(16, 2))
        ttk.Label(body, text="Key file (optional)", style="Card.TLabel").grid(row=11, column=0, sticky="w", pady=4)
        ttk.Entry(body, textvariable=self.ai_key_file, width=62).grid(row=11, column=1, sticky="ew", padx=8)
        key_state = ttk.Label(body, style="CardMuted.TLabel", font=(FONT, 8))
        key_state.grid(row=12, column=1, columnspan=3, sticky="w", padx=8)

        def show_key_state(*_):
            if not key_state.winfo_exists():
                return
            chosen = self.ai_key_file.get().strip()
            key_state.configure(
                text="Leave empty: the team's AI key is fetched automatically from the team SharePoint "
                     "folder (nothing to set up)" if not chosen else
                f"Using {chosen}" if Path(chosen).is_file() else "File not found")
        ttk.Button(body, text="Browse...", command=lambda: (self._browse_key(), show_key_state())).grid(
            row=11, column=2)
        key_trace = self.ai_key_file.trace_add("write", show_key_state)
        show_key_state()
        old_key = self.ai_key_file.get()

        ttk.Label(body, text="Appearance", style="CardHead.TLabel").grid(
            row=19, column=0, sticky="w", columnspan=4, pady=(16, 2))
        look = tk.Frame(body, background=CARD)
        look.grid(row=20, column=0, columnspan=4, sticky="w", pady=2)
        theme = tk.StringVar(value=gui.THEME)
        for value, label in (("light", "☼  Light"), ("dark", "☾  Dark")):
            ttk.Radiobutton(look, text=label, value=value, variable=theme, style="SegSm.TRadiobutton",
                            command=lambda: (close(), self._set_theme(theme.get()))).pack(side="left", padx=(0, 4))
        tk.Label(look, text="GOLD Express restarts to switch", bg=CARD, fg=MUTED, font=(FONT, 8)).pack(
            side="left", padx=10)
        speed_row = tk.Frame(body, background=CARD)
        speed_row.grid(row=21, column=0, columnspan=4, sticky="w", pady=(8, 2))
        tk.Label(speed_row, text="GOLD bot answers", bg=CARD, fg=INK, font=(FONT, 9)).pack(side="left", padx=(0, 10))
        speed = tk.StringVar(value=self.settings.get("bot_speed", "fast"))

        def set_speed():
            self.settings["bot_speed"] = speed.get()
            save_settings(self.settings)
            self.chat = None  # the next question uses the new setting (the conversation is kept)
        for value, label in (("fast", "Fast"), ("thorough", "Thorough")):
            ttk.Radiobutton(speed_row, text=label, value=value, variable=speed, style="SegSm.TRadiobutton",
                            command=set_speed).pack(side="left", padx=(0, 4))
        tk.Label(speed_row, text="Fast: quick answers (recommended).  Thorough: thinks longer.", bg=CARD,
                 fg=MUTED, font=(FONT, 8)).pack(side="left", padx=10)

        ttk.Label(body, text="Notifications", style="CardHead.TLabel").grid(
            row=13, column=0, sticky="w", columnspan=4, pady=(16, 2))
        ttk.Checkbutton(body, text="Show Windows notifications", variable=self.notify_on,
                        style="Card.TCheckbutton").grid(row=14, column=0, columnspan=3, sticky="w", pady=2)
        ttk.Button(body, text="Send a test", command=lambda: self._notify(
            "GOLD Express notifications are on", "You will be told about your orders here.", "test",
            force=True)).grid(row=14, column=3, sticky="e")
        every = tk.Frame(body, background=CARD)
        every.grid(row=15, column=0, columnspan=4, sticky="w", pady=2, padx=(24, 0))
        ttk.Checkbutton(every, text="My orders: new orders and status changes - check every",
                        variable=self.notify_orders, style="Card.TCheckbutton").pack(side="left")
        ttk.Combobox(every, textvariable=self.notify_every, values=("15", "30", "60", "120"), width=4,
                     state="readonly").pack(side="left", padx=6)
        ttk.Label(every, text="minutes", style="Card.TLabel").pack(side="left")
        for r, (text, var) in enumerate([
                ("A daily reminder of late orders and orders due within 7 days", self.notify_digest),
                ("When a new version of GOLD Express is available", self.notify_updates),
                ("When a long task finishes while you are working in another window", self.notify_tasks)], 16):
            ttk.Checkbutton(body, text=text, variable=var, style="Card.TCheckbutton").grid(
                row=r, column=0, columnspan=4, sticky="w", pady=2, padx=(24, 0))

        def close():
            self.ai_key_file.trace_remove("write", key_trace)
            if self.ai_key_file.get() != old_key:
                self.chat = None  # the next question uses the new key file
            save_oa_link()
            self._store_settings()
            if self.show_browser.get() != was_shown and self.worker.driver is not None:
                self.worker.restart = True
                self.log("Firefox will " + ("show" if self.show_browser.get() else "run hidden") +
                         " from the next action.")
            win.destroy()

        ttk.Label(body, text="Order Assignment", style="CardHead.TLabel").grid(
            row=22, column=0, sticky="w", columnspan=4, pady=(16, 2))
        ttk.Label(body, text="Sheet link", style="Card.TLabel").grid(row=23, column=0, sticky="w", pady=4)
        oa_link = tk.StringVar(value=self._oa_link())
        ttk.Entry(body, textvariable=oa_link, width=62).grid(row=23, column=1, columnspan=2, sticky="ew", padx=8)
        ttk.Button(body, text="Default", command=lambda: oa_link.set(gas.SOURCE_LINK)).grid(row=23, column=3)
        ttk.Label(body, text="The APM Open Orders sheet on SharePoint (refreshed there every day). The orders are "
                             "always read from this link.", style="CardMuted.TLabel", font=(FONT, 8)).grid(
            row=24, column=1, columnspan=3, sticky="w", padx=8)

        def save_oa_link():
            link = " ".join(oa_link.get().split())
            if link and link != self._oa_link():
                self.settings["oa_url"] = link
                self.settings["oa_checked"] = 0  # read from the new link when the tab is opened
                self.log("Order Assignment will read the sheet from the new link.")
                if self.notebook.current() is self.oa_frame:
                    self.after(200, lambda: self._oa_load(force=True))

        btns = ttk.Frame(body, style="Card.TFrame")
        btns.grid(row=25, column=0, columnspan=4, sticky="ew", pady=(18, 0))
        ttk.Label(btns, text=f"{APP_NAME} {__version__}  ·  \u00a9 2026 {AUTHOR}  ·  data folder: {HERE}",
                  style="CardMuted.TLabel", font=(FONT, 8)).pack(side="left")
        ttk.Button(btns, text="Done", style="Accent.TButton", command=close).pack(side="right")
        win.protocol("WM_DELETE_WINDOW", close)
        win.bind("<Escape>", lambda e: close())
        win.update_idletasks()
        x = self.winfo_rootx() + (self.winfo_width() - win.winfo_width()) // 2
        y = self.winfo_rooty() + 80
        win.geometry(f"+{x}+{y}")
        win.grab_set()

    # ----- My Orders tab (dashboard of the orders GOLD lists under "My orders only")
    STATUS_COLOURS = {"Creation": "#9AA5B1", "Pricing": "#A78BFA", "Approval": "#F59E0B", "Release": "#38BDF8",
                      "Manage": "#22C55E", "Acceptance": "#14B8A6", "Billing": "#3B82F6",
                      "Awaiting Close": "#FB923C", "On-Hold": "#EF4444", "Closed": "#94A3B8"}
    CUSTOMER_COLOURS = ("#1F86E0", "#56A72C", "#FF7A00", "#8A3FB8", "#128577", "#DD3F58")
    MY_COLUMNS = (("order", "Order", 95), ("customer", "Customer name", 170), ("cust_site_ref", "Cust. site ref", 120),
                  ("status", "Order status", 105), ("target", "Target date", 95), ("service", "Service", 200),
                  ("type", "Type", 90), ("orange_site_ref", "Orange site ref", 240))
    MY_SHOWN = ("order", "customer", "cust_site_ref", "status", "target")  # the rest is in the Excel export

    def _build_my_tab(self, f):
        self.my_frame = f
        self.my_orders, self.my_time = [], None
        self.my_filter = {"scope": "open", "status": None, "customer": None, "when": None}

        # Colourful counters
        kpi = tk.Frame(f, background=BG)
        kpi.pack(fill="x", padx=6, pady=(2, 8))
        self.my_cards = {}
        for col, (key, caption, icon, palette, action) in enumerate((
                ("total", "My orders", "", "v_blue", lambda: self._my_set_filter(scope="all")),
                ("open", "In progress", "", "v_green", lambda: self._my_set_filter(scope="open")),
                ("due", "Due in 30 days", "", "v_orange", lambda: self._my_set_filter(when="due30")),
                ("late", "Past target date", "", "v_rose", lambda: self._my_set_filter(when="late")),
                ("customers", "Customers", "", "v_purple", lambda: self._my_set_filter(scope="all")))):
            card = gui.StatCard(kpi, caption, icon, palette, sub_lines=1, big_size=15, command=action)
            card.grid(row=0, column=col, sticky="ew", padx=(0 if col == 0 else 12, 0))
            kpi.columnconfigure(col, weight=1, uniform="mykpi")
            self.my_cards[key] = card

        # Below the counters: the orders (left, large) and a column of charts (right)
        lower = tk.Frame(f, background=BG)
        lower.pack(fill="both", expand=True)
        scale = self.winfo_fpixels("1i") / 96  # sizes below are for 100% screen scaling
        side = tk.Frame(lower, background=BG)
        side.grid(row=0, column=1, sticky="nsew")
        lower.rowconfigure(0, weight=1)
        lower.columnconfigure(0, weight=3, uniform="mylower")  # the orders get 3/5 of the width
        lower.columnconfigure(1, weight=2, uniform="mylower")

        def head(parent, title, icon):
            h = tk.Frame(parent, background=CARD)
            h.pack(fill="x", pady=(0, 6))
            tk.Label(h, text=icon, bg=CARD, fg=ORANGE, font=(gui.ICONS, 11)).pack(side="left", padx=(0, 8))
            tk.Label(h, text=title, bg=CARD, fg=INK, font=(FONT, 11, "bold")).pack(side="left")
            return h

        t1 = self._tile(side, fill="x")
        head(t1.body, "Orders by status", "")
        row = tk.Frame(t1.body, background=CARD)
        row.pack(fill="x")
        self.my_donut = gui.Donut(row)
        self.my_donut.pack(side="left", padx=(0, 14))
        self.my_legend = tk.Frame(row, background=CARD)
        self.my_legend.pack(side="left", fill="x", expand=True)

        t2 = self._tile(side, fill="both", expand=True)
        head(t2.body, "Orders per customer", "")
        self.my_bars = gui.Bars(t2.body, rows=5, on_click=lambda name: self._my_set_filter(customer=name))
        self.my_bars.pack(fill="x")

        # The orders
        tt = self._tile(lower)
        tt.grid(row=0, column=0, sticky="nsew")
        title = tk.Frame(tt.body, background=CARD)
        title.pack(fill="x", pady=(0, 6))
        tk.Label(title, text="", bg=CARD, fg=ORANGE, font=(gui.ICONS, 11)).pack(side="left", padx=(0, 8))
        tk.Label(title, text="Orders", bg=CARD, fg=INK, font=(FONT, 11, "bold")).pack(side="left")
        ttk.Button(title, text="⤓  Excel", style="Small.TButton", command=self._my_export).pack(side="right")
        self._btn(title, "↻  Refresh", lambda: self._my_load(refresh=True), "Small.TButton").pack(side="right",
                                                                                                padx=6)
        self.my_info = tk.Label(title, text="", bg=CARD, fg=MUTED, font=(FONT, 9))
        self.my_info.pack(side="right", padx=6)
        bar = tk.Frame(tt.body, background=CARD)
        bar.pack(fill="x", pady=(0, 8))
        tk.Label(bar, text="", bg=CARD, fg=ORANGE, font=(gui.ICONS, 12)).pack(side="left", padx=(2, 6))
        self.my_search = tk.StringVar()
        ent = ttk.Entry(bar, textvariable=self.my_search, width=16)
        ent.pack(side="left")
        ent.bind("<KeyRelease>", lambda e: self._my_fill())
        self.my_scope = tk.StringVar(value="open")
        seg = tk.Frame(bar, background=CARD)
        seg.pack(side="left", padx=10)
        for value, label in (("open", "Open"), ("closed", "Closed"), ("all", "All")):
            ttk.Radiobutton(seg, text=label, value=value, variable=self.my_scope, style="SegSm.TRadiobutton",
                            command=lambda: self._my_set_filter(scope=self.my_scope.get())).pack(side="left",
                                                                                                padx=(0, 4))
        self.my_chip = gui.Pill(bar, "", ORANGE_DARK, ORANGE_SOFT, bg=CARD,
                                command=lambda: self._my_set_filter(status=None, customer=None, when=None))

        tf = tk.Frame(tt.body, background=CARD)
        tf.pack(fill="both", expand=True)
        self.my_tree = ttk.Treeview(tf, columns=[c for c, _, _ in self.MY_COLUMNS], show="headings",
                                    selectmode="browse")
        self.my_tree_copy = gui.TableCopy(self.my_tree, self.log, menu=False)  # Excel-like copying
        for c, t, w in self.MY_COLUMNS:
            self.my_tree.heading(c, text=t, anchor="w", command=lambda c=c: self._my_sort(c))
            self.my_tree.column(c, width=int(w * scale), minwidth=int(70 * scale), anchor="w",
                                stretch=True)  # share the tile's width
        self.my_tree.configure(displaycolumns=self.MY_SHOWN)
        self.my_tree.tag_configure("Closed", foreground=MUTED)
        self.my_tree.tag_configure("late", foreground=RED)
        self.my_tree.tag_configure("odd", background=gui.ZEBRA)
        sb = ttk.Scrollbar(tf, orient="vertical", command=self.my_tree.yview)
        self.my_tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.my_tree.pack(side="left", fill="both", expand=True)
        # Until the user asks for them, the table shows a "Get my orders" button (nothing is read at start-up)
        self.my_empty = tk.Frame(tf, background=CARD)
        tk.Label(self.my_empty, text="", bg=CARD, fg=ORANGE, font=(gui.ICONS, 26)).pack(pady=(0, 6))
        tk.Label(self.my_empty, text="Your orders are not loaded yet", bg=CARD, fg=INK,
                 font=(FONT, 12, "bold")).pack()
        self.my_empty_hint = tk.Label(self.my_empty, text="Click below to read your orders from GOLD "
                                      "(My orders only).", bg=CARD, fg=MUTED, font=(FONT, 9))
        self.my_empty_hint.pack(pady=(2, 10))
        self.my_get = self._btn(self.my_empty, "↓  Get my orders", self._my_load, "Accent.TButton")
        self.my_get.pack()
        self.my_tree.bind("<Double-1>", lambda e: self._my_open())
        self.my_tree.bind("<Return>", lambda e: self._my_open())
        self.my_tree.bind("<Button-3>", self._my_menu)
        self.my_sort_key = ("target", False)
        self._my_render()

    # --- loading (only when the user asks: "Get my orders" / Refresh)
    def _my_load(self, refresh=False, background=False):
        if self.busy:
            return
        w = self.worker

        def job():
            w.ensure_browser()
            self.log("Reading your orders from GOLD (My orders only)...")
            return gmo.read_my_orders(w.driver, w.wait, self.log)

        self.my_info.configure(text="Reading GOLD...")
        if not background:
            self.my_empty_hint.configure(text="Reading your orders from GOLD - this takes a few seconds...")
        w.submit("Checking my orders" if background else "Reading my orders", job, self._my_loaded, retry=True)

    def _my_loaded(self, orders):
        self.my_orders, self.my_time = orders, time.time()
        self.log(f"My orders: {len(orders)} order(s) read from GOLD")
        self._my_render()
        self._my_changes(orders)

    # --- watching My orders for changes (Windows notifications)
    def _my_watch(self):
        """Every minute: when it is time (Settings > Notifications), re-read My orders in the background -
        only while GOLD Express is idle and the user is not typing / clicking in it."""
        self.after(60_000, self._my_watch)
        try:
            every = int(self.notify_every.get()) * 60
        except ValueError:
            every = 1800
        due = self.my_time and time.time() - self.my_time >= every  # only once the user has got them
        idle = time.time() - getattr(self, "_last_input", 0) >= 60 or not self._app_in_front()
        if (self.notify_on.get() and self.notify_orders.get() and due and idle and not self.busy
                and self.worker.driver is not None):
            self._my_load(background=True)

    def _my_changes(self, orders):
        """Compare with the last read (kept on disk, so changes while the app was closed count too) and
        notify new orders, status changes and moved target dates; once a day, a reminder."""
        try:
            old = json.loads(MY_SNAPSHOT.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            old = None
        now = {o["order"]: {"status": o.get("status", ""), "ittd": str(o.get("ittd_date") or ""),
                            "customer": o.get("customer", ""), "site": o.get("cust_site_ref", "")}
               for o in orders}
        try:
            MY_SNAPSHOT.write_text(json.dumps(now), encoding="utf-8")
        except OSError:
            pass
        if old is not None:  # nothing to compare with the very first time
            news = []
            for order, o in now.items():
                was = old.get(order)
                if was is None:
                    news.append((f"New order {order}", f"{o['customer']} · {o['site']} · {o['status']}"))
                elif was["status"] != o["status"]:
                    news.append((f"{order}: {was['status']} → {o['status']}", f"{o['customer']} · {o['site']}"))
                elif was["ittd"] != o["ittd"] and o["ittd"]:
                    d = datetime.fromisoformat(o["ittd"]).strftime("%d-%b-%Y")
                    news.append((f"{order}: target date now {d}", f"{o['customer']} · {o['site']}"))
            if len(news) > 3:
                self._notify(f"{len(news)} changes in My orders",
                             "; ".join(t for t, _ in news[:4]) + ("..." if len(news) > 4 else ""), "orders")
            else:
                for title, text in news:
                    self._notify(title, text, "orders")
        today = datetime.now().date()
        if self.settings.get("digest_date") != str(today):
            self.settings["digest_date"] = str(today)
            save_settings(self.settings)
            live = [o for o in orders if o.get("status") != "Closed" and o.get("ittd_date")]
            late = sum(1 for o in live if o["ittd_date"] < today)
            soon = sum(1 for o in live if 0 <= (o["ittd_date"] - today).days <= 7)
            if late or soon:
                self._notify("Your orders today",
                             f"{late} past their target date, {soon} due within 7 days. Open My Orders for "
                             "the list.", "digest")

    # --- drawing
    def _my_render(self):
        s = gmo.summary(self.my_orders)
        loaded = self.my_time is not None
        if loaded:
            self.my_empty.place_forget()
        else:
            self.my_empty.place(relx=0.5, rely=0.45, anchor="center")
        n = lambda v: str(v) if loaded else "-"
        self.my_cards["total"].set(n(s["total"]), "in My orders" if loaded else "not loaded yet")
        self.my_cards["open"].set(n(s["open"]), f"{s['closed']} closed" if loaded else "")
        self.my_cards["due"].set(n(s["due_30"]), "target in 30 days" if loaded else "")
        self.my_cards["late"].set(n(s["late"]), "still open" if loaded else "")
        self.my_cards["customers"].set(n(s["customers"]), (s["by_customer"][0][0] if s["by_customer"] else "")
                                       if loaded else "")
        # donut + legend
        items = [(k, v, self.STATUS_COLOURS.get(k, "#B0B7C0")) for k, v in s["by_status"]]
        self.my_donut.set(items, n(s["total"]), "orders")
        for w in self.my_legend.winfo_children():
            w.destroy()
        for k, v, colour in items:
            r = tk.Frame(self.my_legend, background=CARD, cursor="hand2")
            r.pack(fill="x", pady=1)
            dot = tk.Label(r, text="●", bg=CARD, fg=colour, font=(FONT, 11))
            dot.pack(side="left")
            pct = f"{v * 100 // max(s['total'], 1)}%"
            cnt = tk.Label(r, text=f"{v}  ·  {pct}", bg=CARD, fg=MUTED, font=(FONT, 9, "bold"))
            cnt.pack(side="right")
            name = tk.Label(r, text=k, bg=CARD, fg=INK, font=(FONT, 9), cursor="hand2", anchor="w")
            name.pack(side="left", padx=(4, 10), fill="x", expand=True)
            for wdg in (r, dot, name, cnt):
                wdg.bind("<Button-1>", lambda e, k=k: self._my_set_filter(status=k, scope="all"))
        # customers
        top = s["by_customer"][:5]
        self.my_bars.set([(name, v, self.CUSTOMER_COLOURS[i % len(self.CUSTOMER_COLOURS)])
                          for i, (name, v) in enumerate(top)])
        self._my_fill()

    def _my_set_filter(self, **kw):
        self.my_filter.update(kw)
        if "scope" in kw:
            self.my_scope.set(kw["scope"])
        if kw.get("when"):
            self.my_filter["scope"] = "open"
            self.my_scope.set("open")
        self._my_fill()

    def _my_rows(self):
        f, today = self.my_filter, datetime.now().date()
        q = self.my_search.get().strip().lower()
        out = []
        for o in self.my_orders:
            closed = o.get("status") == "Closed"
            if (f["scope"] == "open" and closed) or (f["scope"] == "closed" and not closed):
                continue
            if f["status"] and o.get("status") != f["status"]:
                continue
            if f["customer"] and o.get("customer") != f["customer"]:
                continue
            d = o.get("ittd_date")
            if f["when"] == "late" and not (d and d < today):
                continue
            if f["when"] == "due30" and not (d and 0 <= (d - today).days <= 30):
                continue
            if q and q not in " ".join(str(v) for v in o.values()).lower():
                continue
            out.append(o)
        key, rev = self.my_sort_key
        sort_val = (lambda o: o.get("ittd_date") or datetime.max.date()) if key == "target" else \
            (lambda o: str(o.get(key, "")).lower())
        return sorted(out, key=sort_val, reverse=rev)

    def _my_fill(self):
        rows = self._my_rows()
        today = datetime.now().date()
        self.my_tree.delete(*self.my_tree.get_children())
        for i, o in enumerate(rows):
            d = o.get("ittd_date")
            late = d and d < today and o.get("status") != "Closed"
            tags = [o.get("status", "")] + (["late"] if late else []) + (["odd"] if i % 2 else [])
            self.my_tree.insert("", "end", iid=o["order"], tags=tags, values=(
                o["order"], o.get("customer", ""), o.get("cust_site_ref", ""), "●  " + o.get("status", ""),
                f"{d:%d-%b-%Y}" if d else "", o.get("service", ""), o.get("type", ""), o.get("orange_site_ref", "")))
        f = self.my_filter
        parts = [p for p in (f["status"] and f"Status: {f['status']}", f["customer"] and f"Customer: {f['customer']}",
                             {"late": "Past target date", "due30": "Due in 30 days"}.get(f["when"])) if p]
        if parts:
            self.my_chip.set("  ·  ".join(parts) + "   ✕")
            self.my_chip.pack(side="left", padx=(4, 0))
        else:
            self.my_chip.pack_forget()
        when = datetime.fromtimestamp(self.my_time).strftime("%H:%M") if self.my_time else ""
        self.my_info.configure(text=f"{len(rows)} of {len(self.my_orders)}" + (f"  ·  {when}" if when else ""))

    def _my_sort(self, col):
        key, rev = self.my_sort_key
        self.my_sort_key = (col, not rev if key == col else False)
        self._my_fill()

    # --- actions
    def _my_selected(self):
        sel = self.my_tree.selection()
        return sel[0] if sel else None

    def _my_open(self, order=None):
        order = order or self._my_selected()
        if not order or self.busy:
            return
        self.notebook.select(self.explorer_frame)
        self.order_var.set(order)
        self.open_order()

    def _my_menu(self, e):
        iid = self.my_tree.identify_row(e.y)
        if not iid:
            return
        self.my_tree.selection_set(iid)
        m = tk.Menu(self, tearoff=0)
        m.add_command(label="Open in Order explorer", command=lambda: self._my_open(iid))
        m.add_command(label="Show in Overview", command=lambda: (self.notebook.select(self.overview_frame),
                                                                 self.ov_input.set(iid), self._ov_run()))
        m.add_command(label="Ask the GOLD bot about it", command=lambda: (self.notebook.select(self.ask_frame),
                                                                  self.order_var.set(iid), self._ask_tab_shown()))
        m.add_separator()
        cell = self.my_tree_copy._cell_at(e)
        if cell and cell[1] not in self.my_tree_copy.block()[1]:
            self.my_tree_copy.anchor = self.my_tree_copy.cur = cell
            self.my_tree_copy._draw()
        self.my_tree_copy.add_items(m, *(cell or (iid, None)))
        m.tk_popup(e.x_root, e.y_root)

    def _my_export(self):
        rows = self._my_rows()
        if not rows:
            return
        path = Path(self._ctx()["workbook"]).parent / f"My_orders_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
        cols = [("Order", "order"), ("Customer", "customer"), ("Cust. Site Ref", "cust_site_ref"),
                ("Order Status", "status"), ("Service", "service"), ("Target date (iTTD)", "ittd_date"),
                ("Order Type", "type"), ("Orange Site Ref", "orange_site_ref"), ("Site Id", "site_id"),
                ("BAN", "ban"), ("Sales Country", "sales_country"), ("Installed Offer Id", "installed_offer")]
        try:
            gx.write_sheet(path, "My orders", [f"Exported {datetime.now():%d-%b-%Y %H:%M}  ·  {len(rows)} orders"],
                           [("My orders", [c for c, _ in cols],
                             [[(o.get(k).strftime("%d-%b-%Y") if k == "ittd_date" and o.get(k) else o.get(k, ""))
                               for _, k in cols] for o in rows])])
        except PermissionError:
            messagebox.showerror(APP_NAME, f"{path.name} is open in Excel. Close it and try again.")
            return
        self.log(f"My orders exported: {path}")
        if messagebox.askyesno(APP_NAME, f"Saved to\n{path}\n\nOpen it now?"):
            os.startfile(path)

    # ----- Overview tab
    RAG_UI = {"Red": (gui.RED_FG, gui.RED_BG), "Amber": (gui.AMBER_FG, gui.AMBER_BG),
              "Green": (gui.GREEN_FG, gui.GREEN_BG), "No date": (gui.GREY_FG, gui.GREY_BG)}

    # ----- Order Assignment tab: which DPM has which open APM order (APM Open Orders sheet, refreshed daily)
    OA_COLUMNS = (("order", "Order", 95), ("status", "Status", 85), ("customer", "Customer", 190),
                  ("dpm", "STM (DPM)", 150), ("region", "Region", 95), ("target", "Target date", 95),
                  ("due_days", "Days to due", 80), ("site_ref", "Cust. site ref", 110), ("country", "Country", 100))
    OA_SHOWS = (("all", "All"), (gas.ASSIGNED, "Assigned"), (gas.UNASSIGNED, "Unassigned"),
                (gas.NOT_IN_LIST, "Not in DPM list"))

    def _build_oa_tab(self, f):
        self.oa_orders, self.oa_time, self.oa_loading = [], None, False
        self.oa_sort = ("target", False)

        kpi = tk.Frame(f, background=BG)
        kpi.pack(fill="x", padx=6, pady=(2, 8))
        self.oa_cards = {}
        for col, (key, caption, icon, palette, show) in enumerate((
                ("total", "Open APM orders", "", "v_blue", "all"),
                ("assigned", "Assigned to DPMs", "", "v_green", gas.ASSIGNED),
                ("unassigned", "Unassigned", "", "v_rose", gas.UNASSIGNED),
                ("not_in_list", "Not in DPM list", "", "v_orange", gas.NOT_IN_LIST),
                ("dpms", "DPMs with orders", "", "v_purple", gas.ASSIGNED))):
            card = gui.StatCard(kpi, caption, icon, palette, sub_lines=1, big_size=15,
                                command=lambda s=show: self._oa_filter(show=s, dpm=""))
            card.grid(row=0, column=col, sticky="ew", padx=(0 if col == 0 else 12, 0))
            kpi.columnconfigure(col, weight=1, uniform="oakpi")
            self.oa_cards[key] = card

        lower = tk.Frame(f, background=BG)
        lower.pack(fill="both", expand=True)
        lower.rowconfigure(0, weight=1)
        lower.columnconfigure(0, weight=3, uniform="oalower")
        lower.columnconfigure(1, weight=1, uniform="oalower")
        scale = self.winfo_fpixels("1i") / 96

        # The orders (left)
        tt = self._tile(lower)
        tt.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        head = self._tile_head(tt.body, "Orders", "")
        ttk.Button(head, text="⤓  Excel dashboard", style="Small.TButton", command=self._oa_export).pack(side="right")
        self._btn(head, "↻  Refresh", lambda: self._oa_load(force=True), "Small.TButton").pack(side="right", padx=6)
        self.oa_info = tk.Label(head, text="", bg=CARD, fg=MUTED, font=(FONT, 9))
        self.oa_info.pack(side="right", padx=6)

        bar = tk.Frame(tt.body, background=CARD)
        bar.pack(fill="x", pady=(0, 8))
        tk.Label(bar, text="", bg=CARD, fg=ORANGE, font=(gui.ICONS, 12)).pack(side="left", padx=(2, 6))
        self.oa_search = tk.StringVar()
        ent = ttk.Entry(bar, textvariable=self.oa_search, width=14)
        ent.pack(side="left")
        ent.bind("<KeyRelease>", lambda e: self._oa_fill())
        self.oa_show = tk.StringVar(value="all")
        seg = tk.Frame(bar, background=CARD)
        seg.pack(side="left", padx=10)
        for value, label in self.OA_SHOWS:
            ttk.Radiobutton(seg, text=label, value=value, variable=self.oa_show, style="SegSm.TRadiobutton",
                            command=lambda: self._oa_filter(show=self.oa_show.get())).pack(side="left", padx=(0, 4))
        filt = tk.Frame(tt.body, background=CARD)
        filt.pack(fill="x", pady=(0, 8))
        ttk.Label(filt, text="Region", style="Muted.TLabel").pack(side="left", padx=(2, 6))
        self.oa_region = tk.StringVar(value="All regions")
        self.oa_region_cb = ttk.Combobox(filt, textvariable=self.oa_region, state="readonly", width=16)
        self.oa_region_cb.pack(side="left")
        self.oa_region_cb.bind("<<ComboboxSelected>>", lambda e: self._oa_filter(region=self.oa_region.get(), dpm=""))
        ttk.Label(filt, text="STM (DPM)", style="Muted.TLabel").pack(side="left", padx=(14, 6))
        self.oa_dpm = tk.StringVar()
        self.oa_dpm_cb = ttk.Combobox(filt, textvariable=self.oa_dpm, width=26)
        self.oa_dpm_cb.pack(side="left")

        def dpm_typed(e):
            if e.keysym in ("Up", "Down", "Escape"):
                return
            if e.keysym == "Return":
                fits = [n for n in self._oa_people() if gn.matches(n, self.oa_dpm.get())]
                if len(fits) == 1:
                    self.oa_dpm.set(fits[0])
                self._oa_fill()
                return
            typed = self.oa_dpm.get().strip()
            self.oa_dpm_cb.configure(values=[n for n in self._oa_people() if gn.matches(n, typed)])
            self._oa_fill()
        self.oa_dpm_cb.bind("<KeyRelease>", dpm_typed)
        self.oa_dpm_cb.bind("<<ComboboxSelected>>", lambda e: self._oa_fill())
        self.oa_chip = gui.Pill(filt, "", ORANGE_DARK, ORANGE_SOFT, bg=CARD,
                                command=lambda: self._oa_filter(show="all", region="All regions", dpm=""))

        # Where the data comes from (click: open the sheet in the browser)
        self.oa_source = tk.Label(tt.body, text="", bg=CARD, fg=MUTED, font=(FONT, 8), anchor="w", cursor="hand2")
        self.oa_source.pack(fill="x", pady=(0, 6))
        self.oa_source.bind("<Button-1>", lambda e: os.startfile(self._oa_link()))

        tf = tk.Frame(tt.body, background=CARD)
        tf.pack(fill="both", expand=True)
        self.oa_tree = ttk.Treeview(tf, columns=[c for c, _, _ in self.OA_COLUMNS], show="headings",
                                    selectmode="extended")
        self.oa_tree_copy = gui.TableCopy(self.oa_tree, self.log, menu=False)
        for c, t, w in self.OA_COLUMNS:
            self.oa_tree.heading(c, text=t, anchor="w", command=lambda c=c: self._oa_sort_by(c))
            self.oa_tree.column(c, width=int(w * scale), minwidth=int(60 * scale), anchor="w", stretch=True)
        self.oa_tree.tag_configure(gas.UNASSIGNED, foreground=RED)
        self.oa_tree.tag_configure(gas.NOT_IN_LIST, foreground=ORANGE_DARK)
        self.oa_tree.tag_configure("odd", background=gui.ZEBRA)
        sb = ttk.Scrollbar(tf, orient="vertical", command=self.oa_tree.yview)
        self.oa_tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.oa_tree.pack(side="left", fill="both", expand=True)
        self.oa_tree.bind("<Double-1>", lambda e: self._oa_open())
        self.oa_tree.bind("<Button-3>", self._oa_menu)
        self.oa_empty = tk.Frame(tf, background=CARD)
        tk.Label(self.oa_empty, text="", bg=CARD, fg=ORANGE, font=(gui.ICONS, 26)).pack(pady=(0, 6))
        tk.Label(self.oa_empty, text="Open APM orders and their DPMs", bg=CARD, fg=INK,
                 font=(FONT, 12, "bold")).pack()
        self.oa_empty_hint = tk.Label(self.oa_empty, text="Read from the APM Open Orders sheet on SharePoint "
                                      "(refreshed every day).", bg=CARD, fg=MUTED, font=(FONT, 9))
        self.oa_empty_hint.pack(pady=(2, 10))
        self._btn(self.oa_empty, "↓  Get the latest orders", lambda: self._oa_load(force=True),
                  "Accent.TButton").pack()

        # Orders per person (right): click a name to see their orders
        pt = self._tile(lower)
        pt.grid(row=0, column=1, sticky="nsew")
        self._tile_head(pt.body, "Orders per STM", "")
        pf = tk.Frame(pt.body, background=CARD)
        pf.pack(fill="both", expand=True)
        self.oa_people = ttk.Treeview(pf, columns=("name", "where", "n"), show="headings", selectmode="browse")
        for c, t, w, anchor in (("name", "Name", 140, "w"), ("where", "Region", 90, "w"), ("n", "Orders", 55, "e")):
            self.oa_people.heading(c, text=t, anchor=anchor)
            self.oa_people.column(c, width=int(w * scale), anchor=anchor, stretch=c != "n")
        self.oa_people.tag_configure(gas.NOT_IN_LIST, foreground=ORANGE_DARK)
        self.oa_people.tag_configure("odd", background=gui.ZEBRA)
        psb = ttk.Scrollbar(pf, orient="vertical", command=self.oa_people.yview)
        self.oa_people.configure(yscrollcommand=psb.set)
        psb.pack(side="right", fill="y")
        self.oa_people.pack(side="left", fill="both", expand=True)
        self.oa_people.bind("<<TreeviewSelect>>", lambda e: self._oa_person_picked())
        self.oa_people_note = tk.Label(pt.body, text="", bg=CARD, fg=MUTED, font=(FONT, 8), anchor="w",
                                       justify="left", wraplength=int(260 * scale))
        self.oa_people_note.pack(fill="x", pady=(6, 0))
        self._oa_render()

    # --- loading: live from the sheet's SharePoint link. Each time the tab is opened (checked at most every
    # OA_CHECK_MIN minutes) SharePoint is asked for the sheet's version; a new version replaces the old data.
    OA_CHECK_MIN = 30

    def _oa_link(self):
        return (self.settings.get("oa_url") or gas.SOURCE_LINK).strip()

    def _oa_shown(self):
        if self.oa_loading:
            return
        checked = self.settings.get("oa_checked", 0)
        recent = OA_FILE.exists() and time.time() - checked < self.OA_CHECK_MIN * 60 \
            and self.settings.get("oa_source") == self._oa_link()
        if recent:
            if not self.oa_time:
                self._oa_read_file()  # just checked: the sheet we have is the latest
        elif self.busy:
            if not self.oa_time and OA_FILE.exists():
                self._oa_read_file()  # show the last sheet now; it is checked next time
        else:
            self._oa_load()

    def _oa_load(self, force=False):
        """Ask SharePoint for the sheet's current version and read it when it changed (Refresh: always)."""
        if self.busy or self.oa_loading:
            return
        link = self._oa_link()
        same_source = self.settings.get("oa_source") == link
        known = None if force or not same_source else self.settings.get("oa_version")
        self.oa_loading = True
        self.oa_info.configure(text="Checking SharePoint...")
        self.oa_empty_hint.configure(text="Reading the APM Open Orders sheet from SharePoint...")
        w = self.worker

        def job():
            w.ensure_browser()
            self.log("Order Assignment: checking the APM Open Orders sheet on SharePoint...")
            info = gas.fetch(w.driver, OA_FILE, link, known)
            orders = gas.read(OA_FILE) if info["changed"] or not self.oa_time else None
            return info, orders

        def done(res):
            info, orders = res
            self.oa_loading = False
            self.settings.update(oa_checked=time.time(), oa_version=info["version"], oa_source=link,
                                 oa_modified=info["modified"].timestamp() if info["modified"] else None)
            save_settings(self.settings)
            if orders is not None:
                self.log("Order Assignment: " + ("a new version of the sheet was read" if info["changed"]
                                                 else "the sheet is up to date"))
                self._oa_loaded(orders)
            else:
                self.log("Order Assignment: the sheet has not changed since it was read")
                self._oa_render()

        w.submit("Reading the APM Open Orders sheet", job, done, retry=True)
        self.after(1000, self._oa_watch_failure)

    def _oa_watch_failure(self):
        if self.oa_loading and not self.busy:  # the job ended with an error (already shown)
            self.oa_loading = False
            self.oa_empty_hint.configure(text="The sheet could not be read - check the Orange network and the "
                                              "link (Settings > Order Assignment), then click below.")
            if OA_FILE.exists() and not self.oa_time:
                self._oa_read_file()  # meanwhile show the last version read
            self._oa_render()
        elif self.oa_loading:
            self.after(1000, self._oa_watch_failure)

    def _oa_read_file(self):
        """Read the copy fetched earlier today (no network), in the background."""
        self.oa_loading = True
        self.oa_info.configure(text="Reading...")

        def work():
            try:
                orders = gas.read(OA_FILE)
            except Exception as e:
                self.ui(lambda e=e: (setattr(self, "oa_loading", False), self.log(f"Order Assignment: {e}"),
                                     self._oa_load(force=True)))
                return
            self.ui(lambda: (setattr(self, "oa_loading", False), self._oa_loaded(orders)))
        threading.Thread(target=work, daemon=True).start()

    def _oa_loaded(self, orders):
        self.oa_orders = gas.classify(orders, self.dpm_names)
        self.oa_time = self.settings.get("oa_modified") or OA_FILE.stat().st_mtime  # when SharePoint refreshed it
        s = gas.summary(self.oa_orders)
        self.log(f"Order Assignment: {s['total']} open orders - {s['assigned']} with DPMs, "
                 f"{s['unassigned']} unassigned, {s['not_in_list']} with people not in the DPM list")
        self._oa_render()

    def _oa_reclassify(self):
        """The DPM list changed: sort the orders again."""
        if self.oa_orders:
            gas.classify(self.oa_orders, self.dpm_names)
            self._oa_render()

    # --- filters and drawing
    def _oa_filter(self, show=None, region=None, dpm=None):
        if show is not None:
            self.oa_show.set(show)
        if region is not None:
            self.oa_region.set(region)
        if dpm is not None:
            self.oa_dpm.set(dpm)
        if self.oa_show.get() == gas.UNASSIGNED:
            self.oa_region.set("All regions")
            self.oa_dpm.set("")
        self._oa_fill()

    def _oa_regions(self):
        found = {o["region"] for o in self.oa_orders if o["category"] == gas.ASSIGNED}
        order = [r for r in gn.TEAMS if r in found] + sorted(found - set(gn.TEAMS))
        return ["All regions"] + order

    def _oa_people(self):
        region, show = self.oa_region.get(), self.oa_show.get()
        names = {o["dpm"] for o in self.oa_orders if o["dpm"]
                 and (show == "all" or o["category"] == show)
                 and (region == "All regions" or o["region"] == region)}
        return sorted(names, key=str.lower)

    def _oa_rows(self):
        show, region = self.oa_show.get(), self.oa_region.get()
        dpm, q = self.oa_dpm.get().strip(), self.oa_search.get().strip().lower()
        out = []
        for o in self.oa_orders:
            if show != "all" and o["category"] != show:
                continue
            if region != "All regions" and o["region"] != region:
                continue
            if dpm and o["dpm"] != dpm and not gn.matches(o["dpm"], dpm):
                continue
            if q and q not in " ".join(str(v) for v in o.values() if v).lower():
                continue
            out.append(o)
        key, rev = self.oa_sort
        if key in ("target", "due_days"):
            far = date.max if key == "target" else 10 ** 9
            val = lambda o: (o.get(key) if o.get(key) is not None and o.get(key) != "" else far)
        else:
            val = lambda o: str(o.get(key) or "").lower()
        try:
            return sorted(out, key=val, reverse=rev)
        except TypeError:
            return sorted(out, key=lambda o: str(val(o)), reverse=rev)

    def _oa_sort_by(self, col):
        key, rev = self.oa_sort
        self.oa_sort = (col, not rev if key == col else False)
        self._oa_fill()

    def _oa_render(self):
        loaded = bool(self.oa_time)
        if loaded:
            self.oa_empty.place_forget()
        else:
            self.oa_empty.place(relx=0.5, rely=0.45, anchor="center")
        s = gas.summary(self.oa_orders)
        n = lambda v: str(v) if loaded else "-"
        pct = lambda v: f"{v * 100 // max(s['total'], 1)}% of the orders" if loaded else ""
        when = datetime.fromtimestamp(self.oa_time).strftime("%d-%b %H:%M") if loaded else ""
        self.oa_cards["total"].set(n(s["total"]), f"sheet refreshed {when}" if loaded else "not loaded yet")
        checked = self.settings.get("oa_checked")
        self.oa_source.configure(
            text=f"Live from SharePoint: {gas.file_name(self._oa_link())}" +
                 (f"  ·  refreshed there {when}" if loaded else "") +
                 (f"  ·  checked {datetime.fromtimestamp(checked):%H:%M}" if loaded and checked else ""))
        self.oa_cards["assigned"].set(n(s["assigned"]), pct(s["assigned"]))
        self.oa_cards["unassigned"].set(n(s["unassigned"]), pct(s["unassigned"]))
        self.oa_cards["not_in_list"].set(n(s["not_in_list"]), f"{s['outsiders']} people" if loaded else "")
        self.oa_cards["dpms"].set(n(s["dpms"]), f"of {len(self.dpm_names)} in the DPM list" if loaded else "")
        self.oa_region_cb.configure(values=self._oa_regions())
        self._oa_fill()

    def _oa_fill(self):
        rows = self._oa_rows()
        t = self.oa_tree
        t.delete(*t.get_children())
        for i, o in enumerate(rows):
            tags = [o["category"]] + (["odd"] if i % 2 else [])
            dpm = o["dpm"] or "— unassigned —"
            region = o["region"] or ("not in list" if o["category"] == gas.NOT_IN_LIST else "")
            t.insert("", "end", iid=f"{o['order']}#{i}", tags=tags, values=(
                o["order"], o.get("status") or "", o.get("customer") or "", dpm, region,
                gas.fmt_date(o.get("target")), "" if o.get("due_days") is None else o["due_days"],
                o.get("site_ref") or "", (o.get("country") or "").title()))
        self.oa_dpm_cb.configure(values=self._oa_people())
        parts = [p for p in (dict(self.OA_SHOWS).get(self.oa_show.get()) if self.oa_show.get() != "all" else "",
                             self.oa_region.get() if self.oa_region.get() != "All regions" else "",
                             self.oa_dpm.get().strip()) if p]
        if parts:
            self.oa_chip.set("  ·  ".join(parts) + "   ✕")
            self.oa_chip.pack(side="left", padx=(12, 0))
        else:
            self.oa_chip.pack_forget()
        self.oa_info.configure(text=f"{len(rows)} of {len(self.oa_orders)}" if self.oa_time else "")
        self._oa_fill_people()

    def _oa_fill_people(self):
        region, show = self.oa_region.get(), self.oa_show.get()
        orders = [o for o in self.oa_orders if (region == "All regions" or o["region"] == region)
                  and (show == "all" or o["category"] == show)]
        p = self.oa_people
        p.delete(*p.get_children())
        people = gas.per_person(orders)
        for i, (name, where, count, _) in enumerate(people):
            tags = ([gas.NOT_IN_LIST] if where == gas.NOT_IN_LIST else []) + (["odd"] if i % 2 else [])
            p.insert("", "end", iid=name, values=(name, "not in list" if where == gas.NOT_IN_LIST else where, count),
                     tags=tags)
        un = sum(1 for o in orders if o["category"] == gas.UNASSIGNED)
        self.oa_people_note.configure(
            text=(f"{len(people)} people  ·  {un} orders without an STM\nOrange names are not in the DPM list - "
                  "add them in DPM names if they are DPMs.") if self.oa_time else "")

    def _oa_person_picked(self):
        sel = self.oa_people.selection()
        if sel and sel[0] != self.oa_dpm.get():
            self.oa_dpm.set(sel[0])
            self._oa_fill()

    # --- actions
    def _oa_selected_orders(self):
        return list(dict.fromkeys(self.oa_tree.set(i, "order") for i in self.oa_tree.selection()))

    def _oa_open(self, order=None):
        orders = [order] if order else self._oa_selected_orders()
        if orders:
            self._my_open(orders[0])

    def _oa_to_bulk(self):
        """Put the selected orders in Bulk assignment, ready to give them a DPM."""
        orders = self._oa_selected_orders()
        if not orders:
            return
        g = self.assign_grid
        empty = [i for i in g.rows() if not any(g.record(i).values())]
        if len(empty) < len(orders):
            g.add_rows(len(orders) - len(empty) + 5)
            empty = [i for i in g.rows() if not any(g.record(i).values())]
        for iid, order in zip(empty, orders):
            g.set_value(iid, "order", order)
        self.notebook.select(self.assign_grid.master.master.master)
        self.log(f"{len(orders)} order(s) added to Bulk assignment - choose the DPM, then Run update.")

    def _oa_menu(self, e):
        iid = self.oa_tree.identify_row(e.y)
        if not iid:
            return
        if iid not in self.oa_tree.selection():
            self.oa_tree.selection_set(iid)
        n = len(self.oa_tree.selection())
        m = tk.Menu(self, tearoff=0)
        m.add_command(label="Open in Order explorer", command=lambda: self._oa_open(self.oa_tree.set(iid, "order")))
        m.add_command(label=f"Assign a DPM (Bulk assignment) - {n} order(s)", command=self._oa_to_bulk)
        m.add_separator()
        cell = self.oa_tree_copy._cell_at(e)
        self.oa_tree_copy.add_items(m, *(cell or (iid, None)))
        m.tk_popup(e.x_root, e.y_root)

    def _oa_export(self):
        rows = self._oa_rows()
        if not rows:
            return
        """The orders shown (current filters) as an Excel dashboard: headline numbers, charts and the data."""
        import gold_assign_dashboard as gad
        path = Path(self._ctx()["workbook"]).parent / f"Order_assignment_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx"
        source = (f"{gas.file_name(self._oa_link())}, refreshed on SharePoint "
                  f"{datetime.fromtimestamp(self.oa_time):%d-%b-%Y %H:%M}")
        parts = [p for p in (dict(self.OA_SHOWS).get(self.oa_show.get()) if self.oa_show.get() != "all" else "",
                             self.oa_region.get() if self.oa_region.get() != "All regions" else "",
                             self.oa_dpm.get().strip(),
                             f"search “{self.oa_search.get().strip()}”" if self.oa_search.get().strip() else "")
                 if p]
        filt = ("Filter: " + ", ".join(parts)) if parts else "All orders"
        self.oa_info.configure(text="Building the dashboard...")
        self.log(f"Order Assignment: building the Excel dashboard ({len(rows):,} orders)...")

        def work():  # a few seconds for 10,000 orders: off the UI thread
            try:
                gad.export(path, rows, source, filt)
            except Exception as e:
                self.ui(lambda e=e: (self._oa_fill(), messagebox.showerror(
                    APP_NAME, f"The dashboard could not be saved:\n\n{e}")))
                return
            self.ui(lambda: finished())

        def finished():
            self._oa_fill()
            self.log(f"Order Assignment dashboard saved: {path}")
            if messagebox.askyesno(APP_NAME, f"Dashboard with {len(rows):,} orders saved to\n{path}\n\nOpen it now?"):
                os.startfile(path)
        threading.Thread(target=work, daemon=True).start()

    def _tile_head(self, parent, title, icon):
        """A tile's title row (icon + bold title); returns the row so buttons can go on its right."""
        h = tk.Frame(parent, background=CARD)
        h.pack(fill="x", pady=(0, 6))
        tk.Label(h, text=icon, bg=CARD, fg=ORANGE, font=(gui.ICONS, 11)).pack(side="left", padx=(0, 8))
        tk.Label(h, text=title, bg=CARD, fg=INK, font=(FONT, 11, "bold")).pack(side="left")
        return h

    def _build_overview_tab(self, page):
        # Two views of the same orders: Overview (counters + details) and Tracker (a sheet to share)
        switch = tk.Frame(page, background=BG)
        switch.pack(fill="x", padx=6, pady=(0, 8))
        self.ov_view = tk.StringVar(value="overview")
        for value, label in (("overview", "Overview"), ("tracker", "Tracker")):
            ttk.Radiobutton(switch, text=label, value=value, variable=self.ov_view, style="Seg.TRadiobutton",
                            command=self._ov_switch).pack(side="left", padx=(0, 4))
        tk.Label(switch, text="Tracker: the same orders as a tracker sheet for management or customers",
                 bg=BG, fg=MUTED, font=(FONT, 8)).pack(side="left", padx=10)
        f = self.ov_main = tk.Frame(page, background=BG)
        f.pack(fill="both", expand=True)
        self.ov_tracker = tk.Frame(page, background=BG)  # built at the end (it shares the order box)

        # Colourful counters for all the orders read (same look as My Orders)
        strip = tk.Frame(f, background=BG)
        strip.pack(fill="x", padx=6, pady=(2, 8))
        self.ov_summary = {}
        for col, (key, caption, icon, palette) in enumerate((
                ("all", "Orders", "", "v_blue"), ("Green", "On track", "", "v_green"),
                ("Amber", "Due soon", "", "v_orange"), ("Red", "Late", "", "v_rose"),
                ("lead", "Lead time at risk", "", "v_purple"))):
            card = gui.StatCard(strip, caption, icon, palette, sub_lines=1, big_size=15,
                                command=lambda k=key: self._ov_pick(k))
            card.grid(row=0, column=col, sticky="ew", padx=(0 if col == 0 else 12, 0))
            strip.columnconfigure(col, weight=1, uniform="ovsum")
            self.ov_summary[key] = card
        self._ov_summary_set({})

        lower = tk.Frame(f, background=BG)
        lower.pack(fill="both", expand=True)
        lower.rowconfigure(0, weight=1)
        lower.columnconfigure(0, weight=1, uniform="ovlower")   # orders (a narrow list)
        lower.columnconfigure(1, weight=3, uniform="ovlower")   # the selected order's details

        # Orders: type them, read them, pick one
        lt = self._tile(lower)
        lt.grid(row=0, column=0, sticky="nsew")
        left = lt.body
        head = self._tile_head(left, "Orders", "")
        self.ov_export_btn = self._btn(head, "⤓  Excel", self._ov_export, "Small.TButton")
        self.ov_export_btn.pack(side="right")
        self._btn(head, "↻  Refresh", lambda: self._ov_run(refresh=True), "Small.TButton").pack(side="right",
                                                                                              padx=6)
        bar = tk.Frame(left, background=CARD)
        bar.pack(fill="x", pady=(0, 4))
        self.ov_input = tk.StringVar()
        ent = ttk.Entry(bar, textvariable=self.ov_input, width=16, font=(FONT, 10))
        ent.pack(side="left", fill="x", expand=True)
        ent.bind("<Return>", lambda e: self._ov_run())
        self._btn(bar, "Show", self._ov_run, "Accent.TButton").pack(side="left", padx=(6, 0))
        hint = tk.Label(left, text=f"Orders separated by spaces or commas.  Red = date passed, Amber = "
                                   f"{go.AMBER_DAYS} days or less.", bg=CARD,
                        fg=FAINT, font=(FONT, 8), anchor="w", justify="left")
        hint.pack(fill="x", pady=(0, 8))
        hint.bind("<Configure>", lambda e: hint.configure(wraplength=max(e.width - 4, 100)))
        lf = tk.Frame(left, background=CARD)
        lf.pack(fill="both", expand=True)
        self.ov_list = ttk.Treeview(lf, columns=("order", "status", "days", "lead"), show="headings",
                                    selectmode="browse")
        self.ov_list_copy = gui.TableCopy(self.ov_list, self.log, menu=True)  # Excel-like copying
        scale = self.winfo_fpixels("1i") / 96
        for c, w, t in [("order", 90, "Order"), ("status", 60, "Status"), ("days", 80, "Days left"),
                        ("lead", 60, "Lead time")]:
            self.ov_list.heading(c, text=t, anchor="w")
            self.ov_list.column(c, width=int(w * scale), minwidth=int(50 * scale), anchor="w", stretch=True)
        for k, (fg, bg) in self.RAG_UI.items():
            self.ov_list.tag_configure(k, foreground=fg, background=bg)
        self.ov_list.tag_configure("err", foreground=RED)
        self.ov_list.pack(fill="both", expand=True)
        self.ov_list.bind("<<TreeviewSelect>>", lambda e: self._ov_show_selected())

        # The selected order
        rt = self._tile(lower)
        rt.grid(row=0, column=1, sticky="nsew", padx=(12, 0))
        right = rt.body
        self._tile_head(right, "Order details", "")
        self.ov_title = tk.Label(right, text="Type orders on the left and press Show", bg=CARD, fg=INK,
                                 font=(FONT, 13, "bold"), anchor="w")
        self.ov_title.pack(fill="x")
        # Cards: order value, delivery status, 120-day lead time
        cards = tk.Frame(right, background=CARD)
        cards.pack(fill="x", pady=(8, 10))
        self.ov_cards = {}
        for col, (key, caption, icon) in enumerate((("value", "Order total value", ""),
                                                    ("delivery", "Delivery status", ""),
                                                    ("lead", "120-day lead time", ""))):
            card = gui.StatCard(cards, caption, icon, "sun" if key == "value" else "grey", sub_lines=2,
                                big_size=14, bg=CARD)
            card.grid(row=0, column=col, sticky="ew", padx=(0 if col == 0 else 10, 0))
            cards.columnconfigure(col, weight=1, uniform="ovcards")
            self.ov_cards[key] = card
        self.ov_banner = tk.Label(right, text="", anchor="w", font=(FONT, 11, "bold"), background=CARD)
        tf = tk.Frame(right, background=CARD)
        tf.pack(fill="both", expand=True)
        self.ov_text = tk.Text(tf, wrap="word", font=(FONT, 10), relief="flat", background=CARD, foreground=INK,
                               padx=4, pady=4, tabs=("230",), spacing1=2, spacing3=2, highlightthickness=0)
        sb = ttk.Scrollbar(tf, command=self.ov_text.yview)
        self.ov_text.configure(yscrollcommand=sb.set, state="disabled")
        sb.pack(side="right", fill="y")
        self.ov_text.pack(fill="both", expand=True)
        self.ov_text.tag_configure("section", font=(FONT, 11, "bold"), foreground=ORANGE, spacing1=10, spacing3=4)
        self.ov_text.tag_configure("label", font=(FONT, 10, "bold"), foreground=MUTED)
        self.ov_text.tag_configure("value", font=(FONT, 10), foreground=INK, lmargin2=234)
        self.ov_results = {}   # order -> overview dict
        self.ov_cache = {}     # order -> (time read, overview dict)
        self.ov_errors = {}    # order -> error text
        self.ov_export_btn.state(["disabled"])
        self._build_tracker(self.ov_tracker)

    OV_CACHE_MIN = 10        # an order read in the last 10 minutes is shown again instantly
    OV_PARALLEL_FROM = 3     # from this many orders, read with extra GOLD sessions side by side
    OV_SESSIONS = 3          # GOLD sessions in total (the main one + 2 extra)

    def _ov_run(self, refresh=False):
        if self.busy:
            return
        orders = []
        for o in re.split(r"[\s,;]+", self.ov_input.get().strip()):
            if o and o not in orders:
                orders.append(o)
        bad = [o for o in orders if not re.fullmatch(r"\d{2}-\d{5,8}", o)]
        if not orders or bad:
            messagebox.showwarning(APP_NAME, "Type GOLD order numbers like 10-1764901"
                                   + (f"\n\nNot an order number: {', '.join(bad)}" if bad else ""))
            return
        self.ov_results, self.ov_errors = {}, {}
        self.ov_order_list = orders  # in the order typed (the tracker keeps it)
        self._ov_summary_set({})
        self.ov_list.delete(*self.ov_list.get_children())
        self._tr_fill()
        now = time.time()
        todo = []
        for o in orders:
            self.ov_list.insert("", "end", iid=o, values=(o, "...", "waiting", ""))
            cached = self.ov_cache.get(o)
            if cached and not refresh and now - cached[0] < self.OV_CACHE_MIN * 60:
                self.ov_results[o] = cached[1]
                self._ov_listed(o, cached[1])
            else:
                todo.append(o)
        if not todo:
            self.log(f"Overview: {len(orders)} order(s) shown from the last read "
                     f"(less than {self.OV_CACHE_MIN} min old) - Refresh reads GOLD again")
            self._ov_show_selected()
            self._refresh_action_state()
            self._tr_fill()
            return
        w = self.worker
        lock = threading.Lock()

        def read_one(driver, o):
            self.ui(lambda: self.ov_list.item(o, values=(o, "...", "reading GOLD...", "")))
            ov = go.collect(driver, Wait(driver, 60), o)
            with lock:
                self.ov_results[o] = ov
                self.ov_cache[o] = (time.time(), ov)
            self.ui(lambda: self._ov_listed(o, ov))

        def failed(o, e):
            err = str(e).splitlines()[0] if str(e) else type(e).__name__
            with lock:
                self.ov_errors[o] = err
            self.ui(lambda: self.ov_list.item(o, values=(o, "error", err, ""), tags=("err",)))
            self.log(f"Overview {o}: FAILED - {err}")

        def job():
            w.ensure_browser()
            queue_ = list(todo)
            helpers = []
            if len(todo) >= self.OV_PARALLEL_FROM and self.settings.get("parallel_sessions", True):
                helpers = w.helper_drivers(min(self.OV_SESSIONS - 1, len(todo) - 1))

            def take():
                with lock:
                    return queue_.pop(0) if queue_ else None

            def main_lane():  # the app's own session: reconnects and retries once
                while (o := take()) is not None:
                    try:
                        try:
                            w.ensure_browser()
                            read_one(w.driver, o)
                        except Exception as e:
                            if not w.session_problem(e):
                                raise
                            self.log(f"Connection problem on {o} - reconnecting and retrying...")
                            w.wait_for_network()
                            w.restart = True
                            w.ensure_browser()
                            read_one(w.driver, o)
                    except Exception as e:
                        failed(o, e)

            def helper_lane(d):  # an extra session: on a session problem, hand the order back
                while (o := take()) is not None:
                    try:
                        read_one(d, o)
                    except Exception as e:
                        if w._healthy(d) and not looks_like_connection_error(e):
                            failed(o, e)
                            continue
                        with lock:
                            queue_.insert(0, o)
                        self.log(f"An extra GOLD session stopped - its orders go to the others")
                        w.drop_helper(d)
                        return

            lanes = [threading.Thread(target=helper_lane, args=(d,)) for d in helpers]
            for t in lanes:
                t.start()
            main_lane()
            for t in lanes:
                t.join()
            main_lane()  # anything a stopped extra session handed back
            return len(self.ov_results), len(helpers) + 1

        def done(res):
            n, sessions = res
            self.log(f"Overview ready for {n} of {len(orders)} order(s)"
                     + (f" ({sessions} GOLD sessions side by side)" if sessions > 1 else ""))
            if not self.ov_list.selection():
                first = next((o for o in orders if o in self.ov_results), None)
                if first:
                    self.ov_list.selection_set(first)
            self._refresh_action_state()
            self.after(100, self._tr_fill)  # after the run (the Excel button is enabled once idle)

        w.submit(f"Reading {len(todo)} order(s)", job, done)

    def _ov_listed(self, o, ov):
        days = "" if ov["days_left"] is None else ov["rag_text"]
        self.ov_list.item(o, values=(o, ov["rag"], days, ov["lt_rag"]), tags=(ov["rag"],))
        self._ov_summary_set({k: v["rag"] for k, v in self.ov_results.items()})
        self._tr_fill()
        if not self.ov_list.selection():
            self.ov_list.selection_set(o)

    def _ov_show_selected(self):
        sel = self.ov_list.selection()
        if not sel:
            return
        o = sel[0]
        t = self.ov_text
        t.configure(state="normal")
        t.delete("1.0", "end")
        if o not in self.ov_results:
            self.ov_title.configure(text=o)
            for key in self.ov_cards:
                self._ov_card(key, "-", "", None)
            self._ov_card("value", "Not available", self.ov_errors.get(o, "Not read yet"), None)
            t.configure(state="disabled")
            return
        ov = self.ov_results[o]
        self.ov_title.configure(text=f"{o}   ·   {ov['customer']}")
        cur = ov.get("currency", "")
        if ov.get("total_value") is not None:
            term = f" x {ov['term_months']} months" if ov.get("term_months") else ""
            self._ov_card("value", go.money(ov["total_value"], cur),
                          f"New monthly {go.money(ov['mrc_new'], cur)}{term} + one-off {go.money(ov['one_off'], cur)}",
                          None)
        else:
            self._ov_card("value", "-", "No pricing totals found", None)
        which = "Revised TDD" if ov["revised_tdd"] else "Initial TDD"
        when = go.fmt(ov["key_date"]) if ov["key_date"] else ""
        self._ov_card("delivery", f"● {ov['rag']}  ·  {ov['rag_text']}", f"{which} {when}" if when else "", ov["rag"])
        self._ov_card("lead", f"● {ov['lt_rag']}", ov["lt_text"] + (
            f"\nDeadline {go.fmt(ov['lt_deadline'])} (signed {go.fmt(ov['signature_date'])} + "
            f"{go.LEAD_TIME_DAYS} days)" if ov["lt_deadline"] else ""), ov["lt_rag"])
        prev = None
        for section, label, key in go.LAYOUT:
            if section != prev:
                t.insert("end", section + "\n", "section")
                prev = section
            t.insert("end", label, "label")
            t.insert("end", "\t" + (go.display_value(ov, label, key) or "-") + "\n", "value")
        if ov["rtdd_history"]:
            t.insert("end", "Revised Target Date history\n", "section")
            for title, d, closed in ov["rtdd_history"]:
                t.insert("end", d or "-", "label")
                t.insert("end", f"\t{title}" + (f"  (closed {closed})" if closed else "  (current)") + "\n", "value")
        t.configure(state="disabled")

    def _ov_card(self, key, big_text, sub_text, rag):
        palette = "sun" if key == "value" else gui.RAG_PALETTE.get(rag, "grey")
        self.ov_cards[key].set(big_text.replace("● ", ""), sub_text, palette)

    def _ov_summary_set(self, rags):
        """rags: order -> RAG of every order read; fills the counters above the list."""
        n = len(rags)
        counts = {k: sum(1 for r in rags.values() if r == k) for k in ("Green", "Amber", "Red")}
        self.ov_summary["all"].set(str(n), "orders read" if n != 1 else "order read")
        for key, text in (("Green", "> {d} days left"), ("Amber", "{d} days or less"), ("Red", "date passed")):
            pct = f"{counts[key] * 100 // n}% · " if n else ""
            self.ov_summary[key].set(str(counts[key]), pct + text.format(d=go.AMBER_DAYS))
        leads = sum(1 for o in rags if (self.ov_results.get(o) or {}).get("lt_rag") == "Red")
        self.ov_summary["lead"].set(str(leads), "past the 120-day deadline")

    def _ov_pick(self, key):
        """A counter was clicked: select the first order it counts."""
        for o in self.ov_list.get_children():
            ov = self.ov_results.get(o) or {}
            if key == "all" or ov.get("rag") == key or (key == "lead" and ov.get("lt_rag") == "Red"):
                self.ov_list.selection_set(o)
                self.ov_list.see(o)
                return

    def _ov_export(self):
        order_list = [o for o in self.ov_list.get_children() if o in self.ov_results]
        if not order_list:
            return
        ctx = self._ctx()
        name = (f"Overview_{order_list[0]}" if len(order_list) == 1
                else f"GOLD_overview_{datetime.now():%Y-%m-%d_%H%M}")
        path = Path(ctx["workbook"]).parent / f"{name}.xlsx"
        try:
            go.export([self.ov_results[o] for o in order_list], path)
        except PermissionError:
            messagebox.showerror(APP_NAME, f"{path.name} is open in Excel. Close it and try again.")
            return
        self.log(f"Overview exported: {path}")
        if messagebox.askyesno(APP_NAME, f"Saved to\n{path}\n\nOpen it now?"):
            os.startfile(path)

    # ----- Overview > Tracker: the orders read as a tracker sheet for management or customers
    def _ov_switch(self):
        if self.ov_view.get() == "tracker":
            self.ov_main.pack_forget()
            self.ov_tracker.pack(fill="both", expand=True)
            self._tr_fill()
        else:
            self.ov_tracker.pack_forget()
            self.ov_main.pack(fill="both", expand=True)

    def _build_tracker(self, f):
        self.tr_comments = gtr.load_comments(TRACKER_COMMENTS)
        self.ov_order_list = []
        tt = self._tile(f, fill="both", expand=True)
        body = tt.body
        head = self._tile_head(body, "Order tracker", "")
        self.tr_export_btn = self._btn(head, "⤓  Excel tracker", self._tr_export, "Accent.TButton")
        self.tr_export_btn.pack(side="right")
        self.tr_audience = tk.StringVar(value=self.settings.get("tracker_view", "management"))
        seg = tk.Frame(head, background=CARD)
        seg.pack(side="right", padx=12)
        tk.Label(seg, text="For", bg=CARD, fg=MUTED, font=(FONT, 9)).pack(side="left", padx=(0, 6))
        for value, label in gtr.VIEWS.items():
            ttk.Radiobutton(seg, text=label, value=value, variable=self.tr_audience, style="SegSm.TRadiobutton",
                            command=self._tr_audience_changed).pack(side="left", padx=(0, 4))

        bar = tk.Frame(body, background=CARD)
        bar.pack(fill="x", pady=(2, 4))
        tk.Label(bar, text="Orders", bg=CARD, fg=MUTED, font=(FONT, 9)).pack(side="left", padx=(0, 6))
        ent = ttk.Entry(bar, textvariable=self.ov_input, width=50, font=(FONT, 10))  # the same orders as Overview
        ent.pack(side="left", fill="x", expand=True)
        ent.bind("<Return>", lambda e: self._ov_run())
        self._btn(bar, "Show", self._ov_run, "Accent.TButton").pack(side="left", padx=(6, 0))
        self._btn(bar, "↻  Refresh", lambda: self._ov_run(refresh=True), "Small.TButton").pack(
            side="left", padx=(6, 0))
        tk.Label(bar, text="Prepared for", bg=CARD, fg=MUTED, font=(FONT, 9)).pack(side="left", padx=(16, 6))
        self.tr_for = tk.StringVar()
        ttk.Entry(bar, textvariable=self.tr_for, width=22).pack(side="left")
        self.tr_hint = tk.Label(body, text="", bg=CARD, fg=FAINT, font=(FONT, 8), anchor="w")
        self.tr_hint.pack(fill="x", pady=(0, 8))

        tf = tk.Frame(body, background=CARD)
        tf.pack(fill="both", expand=True)
        self.tr_tree = ttk.Treeview(tf, show="headings", selectmode="extended")
        self.tr_tree_copy = gui.TableCopy(self.tr_tree, self.log, menu=True)
        for word, (fill, fnt) in gtr.RAG_FILLS.items():
            self.tr_tree.tag_configure(word, foreground="#" + fnt)
        self.tr_tree.tag_configure("odd", background=gui.ZEBRA)
        ys = ttk.Scrollbar(tf, orient="vertical", command=self.tr_tree.yview)
        xs = ttk.Scrollbar(tf, orient="horizontal", command=self.tr_tree.xview)
        self.tr_tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        ys.pack(side="right", fill="y")
        xs.pack(side="bottom", fill="x")
        self.tr_tree.pack(side="left", fill="both", expand=True)
        self.tr_tree.bind("<Double-1>", self._tr_edit)
        self._tr_editor = None
        self._tr_fill()

    def _tr_audience_changed(self):
        self.settings["tracker_view"] = self.tr_audience.get()
        save_settings(self.settings)
        self._tr_fill()

    def _tr_orders(self):
        return [self.ov_results[o] for o in self.ov_order_list if o in self.ov_results]

    def _tr_fill(self):
        if not hasattr(self, "tr_tree"):
            return
        view = self.tr_audience.get()
        cols = gtr.columns(view)
        t = self.tr_tree
        t.delete(*t.get_children())
        scale = self.winfo_fpixels("1i") / 96
        t.configure(columns=[c[0] for c in cols])
        for key, head, width, _, _ in cols:
            t.heading(key, text=head, anchor="w")
            t.column(key, width=int(width * 7.5 * scale), minwidth=int(60 * scale), anchor="w", stretch=False)
        orders = self._tr_orders()
        status_at = [c[0] for c in cols].index("status")
        for i, (o, values) in enumerate(zip(orders, gtr.rows(orders, view, self.tr_comments))):
            shown = [go.fmt(v) if isinstance(v, (date, datetime)) else ("" if v is None else v) for v in values]
            tags = [values[status_at]] + (["odd"] if i % 2 else [])
            t.insert("", "end", iid=o["order"], values=shown, tags=tags)
        pending = [o for o in self.ov_order_list if o not in self.ov_results]
        self.tr_hint.configure(text=(
            f"{len(orders)} order(s)" + (f"  ·  {len(pending)} still being read" if pending and self.busy else "")
            + "  ·  double-click the Comments / next steps cell to type  ·  "
            + ("Customer view: only what a customer should see" if view == "customer"
               else "Management view: everything, incl. STM, value and the latest GOLD note"))
            if self.ov_order_list else "Type the orders above and press Show - they are read from GOLD "
                                       "(read only), then shown here as a tracker.")
        self.tr_export_btn.state(["!disabled"] if orders and not self.busy else ["disabled"])

    def _tr_edit(self, e):
        """Double-click the Comments cell: type a comment / next step (kept for next time)."""
        t = self.tr_tree
        iid, col = t.identify_row(e.y), t.identify_column(e.x)
        if not iid or not col:
            return
        key = t["columns"][int(col[1:]) - 1]
        if key != "comment":
            return
        box = t.bbox(iid, key)
        if not box:
            return
        x, y, w, h = box
        var = tk.StringVar(value=self.tr_comments.get(iid, ""))
        ed = ttk.Entry(t, textvariable=var, font=(FONT, 10))
        ed.place(x=x, y=y, width=max(w, 320), height=h)
        ed.focus_set()
        ed.select_range(0, "end")

        def done(save=True):
            if save:
                text = " ".join(var.get().split())
                self.tr_comments[iid] = text
                try:
                    gtr.save_comments(TRACKER_COMMENTS, self.tr_comments)
                except OSError as ex:
                    self.log(f"Could not save the comment: {ex}")
                t.set(iid, "comment", text)
            ed.destroy()
        ed.bind("<Return>", lambda ev: done())
        ed.bind("<Escape>", lambda ev: done(False))
        ed.bind("<FocusOut>", lambda ev: done())

    def _tr_export(self):
        orders = self._tr_orders()
        if not orders:
            return
        view = self.tr_audience.get()
        who = self.tr_for.get().strip()
        customers = sorted({o.get("customer", "") for o in orders if o.get("customer")})
        name = re.sub(r"[^\w\- ]", "", who or (customers[0] if view == "customer" and len(customers) == 1 else ""))
        path = (Path(self._ctx()["workbook"]).parent /
                f"Order_tracker_{(name.strip().replace(' ', '_') + '_') if name.strip() else ''}"
                f"{gtr.VIEWS[view]}_{datetime.now():%Y-%m-%d_%H%M}.xlsx")
        try:
            gtr.export(path, orders, view, self.tr_comments, prepared_for=who)
        except PermissionError:
            messagebox.showerror(APP_NAME, f"{path.name} is open in Excel. Close it and try again.")
            return
        self.log(f"Tracker saved ({gtr.VIEWS[view]} view): {path}")
        if messagebox.askyesno(APP_NAME, f"Tracker with {len(orders)} order(s) saved to\n{path}\n\nOpen it now?"):
            os.startfile(path)

    # ----- shared pieces of the bulk tabs (Revised Target Date, Bulk Order Assignment)
    def _grid_toolbar(self, f, title, hint, grid_attr):
        tb = tk.Frame(f, background=CARD)
        tb.pack(fill="x", pady=(6, 4))
        words = tk.Frame(tb, background=CARD)
        words.pack(side="left", fill="x", expand=True)
        ttk.Label(words, text=title, style="Head.TLabel").pack(anchor="w")
        ttk.Label(words, text=hint, style="Muted.TLabel", font=(FONT, 8)).pack(anchor="w")
        grid = lambda: getattr(self, grid_attr)
        # self._btn: disabled while GOLD Express is busy (changing the rows mid-run would mix up the results)
        self._btn(tb, "Clear all", lambda: grid().clear()).pack(side="right")
        self._btn(tb, "+ 10 rows", lambda: grid().add_rows(10)).pack(side="right", padx=4)
        self._btn(tb, "Paste from Excel", lambda: self._grid_paste(grid())).pack(side="right")
        self._btn(tb, "Load Excel/CSV...", lambda: self._grid_load(grid())).pack(side="right", padx=4)
        words.pack_forget()  # packed again after the buttons: they keep their size, the hint gets the rest
        words.pack(side="left", fill="x", expand=True)

    def _grid_paste(self, grid):
        n = grid.paste()
        if not n:
            messagebox.showinfo(APP_NAME, "The clipboard is empty. Copy the cells in Excel first.")
        else:
            self.log(f"Pasted {n} row(s)")

    def _grid_load(self, grid):
        path = filedialog.askopenfilename(
            title="Load orders", filetypes=[("Excel / CSV", "*.xlsx *.xlsm *.csv *.txt"), ("All files", "*.*")])
        if not path:
            return
        try:
            n = grid.load(path)
        except Exception as e:
            messagebox.showerror(APP_NAME, f"Could not read the file:\n{e}")
            return
        self.log(f"Loaded {n} row(s) from {Path(path).name}")

    def _bulk_bar(self, f, dry_var, run_cmd, check_cmd):
        rb = tk.Frame(f, background=CARD, pady=8)
        rb.pack(fill="x")
        ttk.Checkbutton(rb, text="Test only (no final save)", variable=dry_var).pack(side="left")
        stop = ttk.Button(rb, text="Stop", command=self._bulk_stop)
        stop.pack(side="right")
        stop.state(["disabled"])
        self._btn(rb, "Run update", run_cmd, "Accent.TButton").pack(side="right", padx=6)
        self._btn(rb, "Check", check_cmd).pack(side="right")
        # Run status (shown while a run is going, in the same row): clock, which order, progress
        info = tk.Frame(rb, background=CARD)
        clock = tk.Label(info, text="", bg=ORANGE_SOFT, fg=ORANGE_DARK, font=("Consolas", 11, "bold"), padx=10,
                         pady=2)
        clock.pack(side="left")
        text = tk.Label(info, text="", bg=CARD, fg=INK, font=(FONT, 9), anchor="w")
        text.pack(side="left", padx=10)
        bar = gui.Progress(info, height=6)
        bar.pack(side="left", fill="x", expand=True, padx=(4, 0))
        stop.run_info = (info, clock, text, bar, rb)
        return stop

    def _bulk_stop(self):
        self.bulk_stop_flag.set()
        self.log("Stopping after the current order...")

    def _bulk_run(self, grid, stop_btn, items, label, action, columns, sheet, header_line, dry):
        """Run action(driver, wait, item) for each (iid, item, row_values) and report per row.
        Writes <sheet>_<date>.xlsx with columns + Result + Details."""
        ctx = self._ctx()
        self.bulk_stop_flag.clear()
        stop_btn.state(["!disabled"])
        w = self.worker
        # A running clock so the user sees it is working: total time, current order and its time
        info, clock, text, bar, rb = stop_btn.run_info
        run = {"start": time.time(), "n": len(items), "k": 0, "iid": None, "order": "", "t": time.time(),
               "on": True, "ok": 0}
        info.pack(side="left", fill="x", expand=True, padx=(16, 12))
        bar.set(0)

        def mmss(s):
            s = int(s)
            return f"{s // 3600}:{s // 60 % 60:02d}:{s % 60:02d}" if s >= 3600 else f"{s // 60:02d}:{s % 60:02d}"

        def tick():
            if not run["on"]:
                return
            now = time.time()
            clock.configure(text=f"⏱ {mmss(now - run['start'])}")
            if run["iid"]:
                secs = int(now - run["t"])
                text.configure(text=f"{run['k']}/{run['n']}  ·  {run['order']}  ·  {secs} s")
                grid.set_status(run["iid"], f"Working...  {secs} s", "run")
            self.after(1000, tick)
        tick()

        def set_status(iid, text, tag):
            self.ui(lambda: grid.set_status(iid, text, tag))

        def job():
            w.ensure_browser()
            results = []
            for k, (iid, item, row) in enumerate(items, 1):
                if self.bulk_stop_flag.is_set():
                    set_status(iid, "Skipped (stopped)", "bad")
                    results.append(row + ["SKIPPED", "stopped"])
                    continue
                self.ui(lambda iid=iid, k=k, order=row[0]: (
                    run.update(iid=iid, k=k, order=order, t=time.time()), bar.set((k - 1) / max(run["n"], 1))))
                set_status(iid, "Working...", "run")
                try:
                    w.ensure_browser()  # reconnects if the connection dropped (saves are not retried)
                    t0 = time.time()
                    msg = action(w.driver, w.wait, item)
                    msg = f"{msg}  ({int(time.time() - t0)} s)"
                    self.ui(lambda: run.update(iid=None))  # stop the row's own clock before its result shows
                    set_status(iid, f"✓ {msg}", "ok")
                    if not dry:  # saved in GOLD: a re-run skips this row (no duplicate milestone / save)
                        self.ui(lambda iid=iid: grid.mark_done(iid))
                    results.append(row + ["OK", msg])
                    self.log(f"{row[0]}: {msg}")
                except Exception as e:
                    err = str(e).splitlines()[0] if str(e) else type(e).__name__
                    self.ui(lambda: run.update(iid=None))
                    set_status(iid, f"✗ {err}", "bad")
                    results.append(row + ["FAILED", err])
                    self.log(f"{row[0]}: FAILED - {err}")
                    try:  # back to a clean search page for the next order
                        oo.open_orders_tab(w.driver, w.wait, discard_changes=True)
                    except Exception:
                        pass
            stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
            path = Path(ctx["workbook"]).parent / f"{sheet.replace(' ', '_')}_{stamp}.xlsx"
            try:  # the GOLD work is done: a file problem (open in Excel, OneDrive) must not hide the results
                gx.write_sheet(path, sheet, [header_line + ("  |  TEST ONLY - nothing saved" if dry else ""),
                                             f"Run {datetime.now():%d-%b-%Y %H:%M}"],
                               [("Results", columns + ["Result", "Details"], results)])
            except Exception as e:
                self.log(f"Could not save the results file: {e}")
                path = None
            return path, results

        def finish(summary):
            run["on"] = False
            bar.set(1)
            clock.configure(text=f"⏱ {mmss(time.time() - run['start'])}")
            text.configure(text=summary)

        def watch_failure():  # the whole run stopped with an error: stop the clock too
            if run["on"] and not self.busy:
                finish("Stopped - see the message")
            elif run["on"]:
                self.after(1000, watch_failure)
        self.after(1500, watch_failure)

        def done(res):
            path, results = res
            ok = sum(1 for r in results if r[len(columns)] == "OK")
            finish(f"{ok} of {len(results)} {'tested' if dry else 'done'}")
            self.log(f"{label}: {ok} of {len(results)} OK. Results: {path or 'not saved'}")
            if not path:
                messagebox.showinfo(APP_NAME, f"{ok} of {len(results)} order(s) {'tested' if dry else 'done'}."
                                              "\n\nThe results file could not be saved - the Status column "
                                              "shows each order's result.")
            elif messagebox.askyesno(APP_NAME, f"{ok} of {len(results)} order(s) "
                                               f"{'tested' if dry else 'done'}.\n\nResults saved to\n{path}"
                                               "\n\nOpen it now?"):
                os.startfile(path)

        w.submit(f"{label}: {len(items)} order(s)", job, done)

    # ----- Revised Target Date tab
    def _build_rtdd_tab(self, f):
        top = self._tile(f, fill="x").body
        ttk.Label(top, text="Adds a 'Revised Target Date' milestone (Milestones > Create new) to each order "
                            "in the table, with its Target Date and reason.", style="Muted.TLabel",
                  wraplength=900, justify="left").pack(anchor="w")
        opts = tk.Frame(top, background=CARD, pady=10)
        opts.pack(fill="x")
        ttk.Label(opts, text="Reason type").grid(row=0, column=0, sticky="w")
        self.rtdd_kind = tk.StringVar(value="Customer")
        kinds = tk.Frame(opts, background=CARD)
        kinds.grid(row=0, column=1, columnspan=3, sticky="w", padx=8)
        for kind, label in [("Customer", "Customer reason"), ("Orange", "Orange Business Services reason")]:
            ttk.Radiobutton(kinds, text=label, value=kind, variable=self.rtdd_kind,
                            style="Seg.TRadiobutton", command=self._rtdd_kind_changed).pack(side="left", padx=(0, 4))
        ttk.Label(opts, text="For empty cells use", style="Muted.TLabel").grid(row=1, column=0, sticky="w", pady=(8, 0))
        dflt = tk.Frame(opts, background=CARD)
        dflt.grid(row=1, column=1, columnspan=3, sticky="w", padx=8, pady=(8, 0))
        ttk.Label(dflt, text="Date").pack(side="left")
        self.rtdd_default = tk.StringVar()
        ttk.Entry(dflt, textvariable=self.rtdd_default, width=14).pack(side="left", padx=(4, 16))
        ttk.Label(dflt, text="Reason").pack(side="left")
        self.rtdd_reason = tk.StringVar(value="")
        self.rtdd_reason_cb = ttk.Combobox(dflt, textvariable=self.rtdd_reason, state="readonly", width=62)
        self.rtdd_reason_cb.pack(side="left", padx=4)

        def norm_date(v):
            try:
                return gm.gold_date(gm.parse_date(v))
            except ValueError:
                return str(v)

        f = self._tile(f, fill="both", expand=True).body  # the table and its buttons in their own tile
        self._grid_toolbar(f, "Orders to update",
                           "click a cell to type  ·  Ctrl+V pastes rows from Excel  ·  Delete clears rows",
                           "rtdd_grid")
        self.rtdd_grid = EditGrid(f, [
            {"key": "order", "title": "Gold order", "width": 130},
            {"key": "date", "title": "Date", "width": 120, "normalize": norm_date},
            {"key": "reason", "title": "Reason", "width": 430, "stretch": True,
             "choices": lambda: gm.REASONS[self.rtdd_kind.get()]},
        ])
        self.rtdd_grid.pack(fill="both", expand=True)
        self.rtdd_dry = tk.BooleanVar(value=False)
        self.rtdd_stop_btn = self._bulk_bar(f, self.rtdd_dry, self._rtdd_run, self._rtdd_check)
        self._rtdd_kind_changed()

    def _rtdd_kind_changed(self):
        reasons = [""] + gm.REASONS[self.rtdd_kind.get()]
        self.rtdd_reason_cb.configure(values=reasons)
        if self.rtdd_reason.get() not in reasons:
            self.rtdd_reason.set("")

    def _rtdd_check(self):
        """Validate the table; mark Ready or the problem. Returns ([(iid, item, row)], n_bad) or None."""
        g = self.rtdd_grid
        g.close_editor(True)
        kind = self.rtdd_kind.get()
        default_date = self.rtdd_default.get().strip()
        if default_date:
            try:
                gm.parse_date(default_date)
            except ValueError as e:
                messagebox.showwarning(APP_NAME, f"Default date: {e}")
                return None
        items, seen, bad = [], set(), 0
        g.skipped = 0
        for iid, r in g.filled():
            if g.is_done(iid, r):  # already saved in GOLD in this session: never add it twice
                g.skipped += 1
                g.set_status(iid, "✓ Already done - skipped (change the row to run it again)", "ok")
                seen.add(r["order"])
                continue
            try:
                order = r["order"]
                if not re.fullmatch(r"\d{2}-\d{5,8}", order):
                    raise ValueError("not a GOLD order number (like 10-1764901)")
                if order in seen:
                    raise ValueError("this order is already in the list")
                seen.add(order)
                if not (r["date"] or default_date):
                    raise ValueError("no date")
                dt = gm.parse_date(r["date"] or default_date)
                reason = gm.match_reason(kind, r["reason"]) if r["reason"] else self.rtdd_reason.get()
                if r["reason"] and reason != r["reason"]:
                    g.set_value(iid, "reason", reason)  # GOLD's exact wording
                items.append((iid, (order, dt, reason), [order, gm.gold_date(dt), reason]))
                g.set_status(iid, "Ready")
            except ValueError as e:
                bad += 1
                g.set_status(iid, f"✗ {e}", "bad")
        self.log(f"Revised Target Date check: {len(items)} ready" + (f", {bad} with problems" if bad else ""))
        return items, bad

    def _rtdd_run(self):
        if self.busy:
            return
        checked = self._rtdd_check()
        if not checked:
            return
        items, bad = checked
        if not items:
            messagebox.showinfo(APP_NAME, "Every row was already done in GOLD (✓) - nothing to run again."
                                if self.rtdd_grid.skipped else
                                "Fill in at least one Gold order and Date, then try again.")
            return
        if bad and not messagebox.askyesno(APP_NAME, "Some rows have problems (shown in red) and will be "
                                                     "skipped.\nContinue with the others?"):
            return
        kind, dry = self.rtdd_kind.get(), self.rtdd_dry.get()
        title = gm.RTDD_TITLES[kind]
        if not dry and not messagebox.askyesno(
                "Confirm GOLD update", f"Add the milestone\n\n    {title}\n\nto {len(items)} order(s) in GOLD, "
                                       "each with its Date (and Reason, where given)?\n\nThis saves in GOLD."):
            return
        self._bulk_run(self.rtdd_grid, self.rtdd_stop_btn, items, "Revised Target Date",
                       lambda d, w, it: gm.add_rtdd(d, w, it[0], kind, it[1], it[2], dry_run=dry),
                       ["Gold order", "Date", "Reason"], "RTDD update", title, dry)

    # ----- Bulk Order Assignment tab
    def _dpm_teams(self):
        """The DPM teams in the list (DPM Egypt, DPM India, DPM Brazil...), then Other (names without a team)."""
        teams = sorted({gn.team_of(e) for e in self.dpm_names} - {gn.NO_TEAM} | set(gn.TEAMS.values()))
        return teams + [gn.NO_TEAM]

    def _dpm_choices(self):
        """The names offered in Bulk assignment: those of the teams ticked there."""
        ticked = {t for t, v in getattr(self, "assign_team_vars", {}).items() if v.get()}
        return [gn.full_name(e) for e in self.dpm_names
                if not hasattr(self, "assign_team_vars") or gn.team_of(e) in ticked]

    def _dpm_normalize(self, value):
        """Snap a typed / pasted name to the list's spelling - only when it IS that name, ignoring spaces,
        hyphens, capitals and order ('Ahmed El-Bourgy', 'ELBOURGY Ahmed'). A partial name is never snapped
        to someone else (pick it from the suggestions instead)."""
        e = gn.lookup(self.dpm_names, value)
        if e:
            return gn.full_name(e)
        same = [x for x in self.dpm_names if gn.same_person_name(gn.full_name(x), value)]
        if len(same) == 1:
            return gn.full_name(same[0])
        return " ".join(str(value).split())

    def _build_assign_tab(self, f):
        top = self._tile(f, fill="x").body
        ttk.Label(top, text="Makes the chosen DPM the primary STM of each order (Teams > Manage Teams > STM), "
                            "then saves the teams. Pick names from the list - manage it in the DPM Names tab.",
                  style="Muted.TLabel", wraplength=900, justify="left").pack(anchor="w")
        teams_row = tk.Frame(top, background=CARD, pady=(8))
        teams_row.pack(fill="x")
        ttk.Label(teams_row, text="DPM teams", style="Muted.TLabel").pack(side="left", padx=(0, 8))
        saved = self.settings.get("assign_teams")
        self.assign_team_vars = {}
        self.assign_teams_row = teams_row
        for team in self._dpm_teams():
            self._assign_team_box(team, saved is None or team in saved)
        self.assign_team_info = ttk.Label(teams_row, text="", style="Muted.TLabel")
        self.assign_team_info.pack(side="right")
        opts = tk.Frame(top, background=CARD, pady=10)
        opts.pack(fill="x")
        ttk.Label(opts, text="For empty cells use", style="Muted.TLabel").pack(side="left")
        self.assign_default = tk.StringVar()
        self.assign_default_cb = ttk.Combobox(opts, textvariable=self.assign_default, width=34,
                                              values=[""] + self._dpm_choices())
        self.assign_default_cb.pack(side="left", padx=8)

        def filter_default(event):
            if event.keysym in ("Return", "Tab", "Escape", "Up", "Down"):
                return
            typed = self.assign_default.get().strip()
            self.assign_default_cb.configure(values=[c for c in self._dpm_choices() if gn.matches(c, typed)]
                                             if typed else [""] + self._dpm_choices())
        self.assign_default_cb.bind("<KeyRelease>", filter_default)
        self.assign_default_cb.bind("<FocusOut>", lambda e: self.assign_default.get().strip() and
                                    self.assign_default.set(self._dpm_normalize(self.assign_default.get())))

        f = self._tile(f, fill="both", expand=True).body  # the table and its buttons in their own tile
        self._grid_toolbar(f, "Orders to assign",
                           "click a cell to type  ·  in the DPM cell type any part of a name (e.g. elbourgy)  ·  "
                           "Ctrl+V pastes from Excel", "assign_grid")
        self.assign_grid = EditGrid(f, [
            {"key": "order", "title": "Gold order", "width": 140},
            {"key": "dpm", "title": "STM (DPM)", "width": 300, "choices": self._dpm_choices,
             "typeahead": True, "match": gn.matches, "normalize": self._dpm_normalize},
        ])
        self._assign_teams_changed(save=False)
        self.assign_grid.pack(fill="both", expand=True)
        self.assign_dry = tk.BooleanVar(value=False)
        self.assign_stop_btn = self._bulk_bar(f, self.assign_dry, self._assign_run, self._assign_check)

    def _assign_team_box(self, team, on=True):
        var = tk.BooleanVar(value=on)
        self.assign_team_vars[team] = var
        ttk.Checkbutton(self.assign_teams_row, text=team, variable=var, style="Card.TCheckbutton",
                        command=self._assign_teams_changed).pack(side="left", padx=(0, 10))

    def _assign_teams_changed(self, save=True):
        """The DPM team tick-boxes changed: offer only those names (and remember the choice)."""
        for team in self._dpm_teams():  # a team that appeared since (e.g. after an Org Chart import)
            if team not in self.assign_team_vars:
                self._assign_team_box(team)
        names = self._dpm_choices()
        self.assign_default_cb.configure(values=[""] + names)
        ticked = [t for t, v in self.assign_team_vars.items() if v.get()]
        counts = {t: sum(gn.team_of(e) == t for e in self.dpm_names) for t in self.assign_team_vars}
        self.assign_team_info.configure(
            text=f"{len(names)} names offered  ·  " + ", ".join(f"{t} {counts[t]}" for t in ticked if counts[t]))
        if save:
            self.settings["assign_teams"] = ticked
            save_settings(self.settings)

    def _assign_check(self):
        g = self.assign_grid
        g.close_editor(True)
        default = self.assign_default.get().strip()
        items, seen, bad = [], set(), 0
        g.skipped = 0
        for iid, r in g.filled():
            if g.is_done(iid, r):  # already saved in GOLD in this session
                g.skipped += 1
                g.set_status(iid, "✓ Already done - skipped (change the row to run it again)", "ok")
                seen.add(r["order"])
                continue
            try:
                order = r["order"]
                if not re.fullmatch(r"\d{2}-\d{5,8}", order):
                    raise ValueError("not a GOLD order number (like 10-1764901)")
                if order in seen:
                    raise ValueError("this order is already in the list")
                seen.add(order)
                name = r["dpm"] or default
                if not name:
                    raise ValueError("choose the STM (DPM)")
                entry = gn.lookup(self.dpm_names, name)
                first, last = (entry["first"], entry["last"]) if entry else gn.split_name(name)
                if not last:
                    raise ValueError("the STM needs a first and a last name")
                user = (entry or {}).get("user") or None
                items.append((iid, (order, first, last, user), [order, f"{first} {last}"]))
                g.set_status(iid, "Ready" if entry else "Ready (name is not in the DPM list)")
            except ValueError as e:
                bad += 1
                g.set_status(iid, f"✗ {e}", "bad")
        self.log(f"Bulk Order Assignment check: {len(items)} ready" + (f", {bad} with problems" if bad else ""))
        return items, bad

    def _assign_run(self):
        if self.busy:
            return
        items, bad = self._assign_check()
        if not items:
            messagebox.showinfo(APP_NAME, "Every row was already done in GOLD (✓) - nothing to run again."
                                if self.assign_grid.skipped else
                                "Fill in at least one Gold order with its STM (DPM), then try again.")
            return
        if bad and not messagebox.askyesno(APP_NAME, "Some rows have problems (shown in red) and will be "
                                                     "skipped.\nContinue with the others?"):
            return
        dry = self.assign_dry.get()
        people = sorted({row[1] for _, _, row in items})
        if not dry and not messagebox.askyesno(
                "Confirm GOLD update",
                f"Make {', '.join(people[:3])}{' and others' if len(people) > 3 else ''} the primary STM "
                f"of {len(items)} order(s) and save the teams in GOLD?\n\n"
                "Where STM is 'Same as project', that is unticked so the order gets its own STM.\n\n"
                "This saves in GOLD."):
            return
        self._bulk_run(self.assign_grid, self.assign_stop_btn, items, "Bulk Order Assignment",
                       lambda d, w, it: gt.assign_stm(d, w, it[0], it[1], it[2], dry_run=dry, user=it[3]),
                       ["Gold order", "STM (DPM)"], "STM assignment", "Primary STM assignment", dry)

    # ----- DPM Names tab
    def _build_names_tab(self, f):
        top = self._tile(f, fill="x").body
        ttk.Label(top, text="The DPMs (STMs) used by Order Assignment and Bulk assignment, with their names as they appear in "
                            "GOLD. Add, edit or remove names; the list is saved on this computer.",
                  style="Muted.TLabel", wraplength=900, justify="left").pack(anchor="w")
        tb = tk.Frame(top, background=CARD)
        tb.pack(fill="x", pady=(10, 0))
        self._btn(tb, "+ Add name", self._names_add, "Accent.TButton").pack(side="left")
        self._btn(tb, "Edit", self._names_edit).pack(side="left", padx=6)
        self._btn(tb, "Delete", self._names_delete).pack(side="left")
        self._btn(tb, "Check names in GOLD", self._names_check).pack(side="left", padx=(18, 0))
        self._btn(tb, "Export to Excel", self._names_export).pack(side="right")
        self._btn(tb, "Import Excel/CSV...", self._names_import).pack(side="right", padx=6)
        self._btn(tb, "👥  Import DPM teams from Org Chart", self._names_import_oc).pack(side="right")
        fb = tk.Frame(top, background=CARD)
        fb.pack(fill="x", pady=(10, 0))
        ttk.Label(fb, text="Find").pack(side="left", padx=(0, 4))
        self.names_filter = tk.StringVar()
        fe = ttk.Entry(fb, textvariable=self.names_filter, width=22)
        fe.pack(side="left")
        fe.bind("<KeyRelease>", lambda e: self._names_fill())
        ttk.Label(fb, text="Team").pack(side="left", padx=(18, 6))
        self.names_team = tk.StringVar(value="All")
        self.names_team_row = tk.Frame(fb, background=CARD)
        self.names_team_row.pack(side="left")

        f = self._tile(f, fill="both", expand=True).body
        lf = tk.Frame(f, background=CARD)
        lf.pack(fill="both", expand=True)
        cols = ("name", "team", "first", "last", "user", "check")
        self.names_tree = ttk.Treeview(lf, columns=cols, show="headings", selectmode="extended")
        self.names_tree_copy = gui.TableCopy(self.names_tree, self.log, menu=True)  # Excel-like copying
        for c, w, t in [("name", 220, "Name"), ("team", 110, "Team"), ("first", 150, "GOLD first name"),
                        ("last", 150, "GOLD last name"), ("user", 110, "GOLD user"), ("check", 280, "Checked in GOLD")]:
            self.names_tree.heading(c, text=t, anchor="w")
            self.names_tree.column(c, width=w, anchor="w", stretch=c == "check")
        self.names_tree.tag_configure("ok", foreground=GREEN)
        self.names_tree.tag_configure("bad", foreground=RED)
        self.names_tree.tag_configure("odd", background=gui.ZEBRA)
        sb = ttk.Scrollbar(lf, command=self.names_tree.yview)
        self.names_tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.names_tree.pack(fill="both", expand=True)
        self.names_tree.bind("<Double-1>", lambda e: self._names_edit())
        self.names_tree.bind("<Delete>", lambda e: self._names_delete())
        self.names_count = ttk.Label(f, text="", style="Muted.TLabel")
        self.names_count.pack(anchor="w", pady=(6, 0))
        self._names_fill()

    def _names_saved(self):
        self.dpm_names = gn.save(self.dpm_names)
        self._names_fill()
        if hasattr(self, "assign_team_vars"):
            self._assign_teams_changed(save=False)
        if hasattr(self, "oa_orders"):
            self._oa_reclassify()  # Order Assignment: a name added / given a team moves its orders

    def _names_team_buttons(self):
        """The Team filter: All, then each team in the list (DPM Egypt, DPM India, DPM Brazil, Other)."""
        row = self.names_team_row
        teams = ["All"] + [t for t in self._dpm_teams() if any(gn.team_of(e) == t for e in self.dpm_names)]
        if getattr(self, "_names_teams_shown", None) == teams:
            return
        self._names_teams_shown = teams
        for w in row.winfo_children():
            w.destroy()
        if self.names_team.get() not in teams:
            self.names_team.set("All")
        for t in teams:
            ttk.Radiobutton(row, text=t, value=t, variable=self.names_team, style="SegSm.TRadiobutton",
                            command=self._names_fill).pack(side="left", padx=(0, 4))

    def _names_fill(self):
        t = self.names_tree
        t.delete(*t.get_children())
        self._names_team_buttons()
        q = self.names_filter.get().strip() if hasattr(self, "names_filter") else ""
        team = self.names_team.get()
        shown = 0
        for i, e in enumerate(self.dpm_names):
            name = gn.full_name(e)
            if q and not gn.matches(name, q) and q.lower() not in e.get("user", "").lower():
                continue
            if team != "All" and gn.team_of(e) != team:
                continue
            status = e.get("check", "")
            tags = ["ok" if status.startswith("✓") else "bad" if status.startswith("✗") else ""]
            if shown % 2:
                tags.append("odd")
            t.insert("", "end", iid=str(i), values=(name, e.get("team", ""), e["first"], e["last"],
                                                    e.get("user", ""), status), tags=[x for x in tags if x])
            shown += 1
        self.names_count.configure(
            text=f"{len(self.dpm_names)} names" + (f"  ·  {shown} shown" if q or team != "All" else "")
            if self.dpm_names else "No names yet - click “Import DPM teams from Org Chart” to add the DPM teams "
                                   "(Egypt, India, Brazil), or add names one by one.")

    def _names_import_oc(self):
        """Org Chart > Search by Organization > Department 'APM': APM Europe -> DPM Egypt, APM APAC -> DPM India,
        APM AME -> DPM Brazil. Adds the people (or sets the team of names already in the list). Read-only."""
        if self.busy:
            return

        def job():
            def read(oc):
                depts = [x for x in oc.departments("APM") if x["name"] in gn.TEAMS]
                if not depts:
                    raise RuntimeError("No APM departments found in the Org Chart.")
                out = []
                for x in depts:
                    self.log(f"Org Chart: reading {x['name']}...")
                    out += [gn.from_org_chart(r, gn.TEAMS[x["name"]]) for r in oc.department_people([x["id"]])]
                return out
            return self._org(read)

        def done(people):
            before = len(self.dpm_names)
            added, updated = gn.merge_team(self.dpm_names, people)
            self._names_saved()
            counts = ", ".join(f"{t}: {sum(p['team'] == t for p in people)}" for t in gn.TEAMS.values())
            self.log(f"DPM teams from the Org Chart ({counts}): {added} added, {updated} updated")
            self._message(APP_NAME, "DPM teams imported",
                          f"{len(people)} people in the Org Chart ({counts}).\n\n{added} added to the list, "
                          f"{updated} names already there now have their team ({before} names before, "
                          f"{len(self.dpm_names)} now).\n\nTip: 'Check names in GOLD' makes sure the new "
                          "names are spelled as in GOLD.")

        self.worker.submit("Reading DPM teams from the Org Chart", job, done, retry=True)

    def _names_dialog(self, title, entry=None):
        """Small form for one name. Returns {'first','last'} or None."""
        win = tk.Toplevel(self, background=CARD)
        win.title(title)
        win.transient(self)
        win.resizable(False, False)
        body = ttk.Frame(win, style="Card.TFrame", padding=(20, 16))
        body.pack(fill="both", expand=True)
        ttk.Label(body, text="Type the name exactly as it appears in GOLD.", style="CardMuted.TLabel").grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))
        first = tk.StringVar(value=(entry or {}).get("first", ""))
        last = tk.StringVar(value=(entry or {}).get("last", ""))
        user = tk.StringVar(value=(entry or {}).get("user", ""))
        team = tk.StringVar(value=(entry or {}).get("team", ""))
        for r, (label, var) in enumerate((("First name", first), ("Last name", last),
                                          ("GOLD user ID", user), ("Team", team)), 1):
            ttk.Label(body, text=label, style="Card.TLabel").grid(row=r, column=0, sticky="w", pady=4)
            if var is team:
                e = ttk.Combobox(body, textvariable=var, width=32,
                                 values=[""] + [t for t in self._dpm_teams() if t != gn.NO_TEAM])
            else:
                e = ttk.Entry(body, textvariable=var, width=34)
            e.grid(row=r, column=1, padx=(10, 0), pady=4)
            if r == 1:
                e.focus_set()
        result = {}

        def ok(*_):
            f, l = " ".join(first.get().split()), " ".join(last.get().split())
            if not f or not l:
                messagebox.showwarning(APP_NAME, "Fill in both the first and the last name.", parent=win)
                return
            dup = gn.lookup(self.dpm_names, f"{f} {l}")
            if dup is not None and dup is not entry:
                messagebox.showwarning(APP_NAME, f"{f} {l} is already in the list.", parent=win)
                return
            result.update(first=f, last=l, user=user.get().strip(), team=" ".join(team.get().split()))
            win.destroy()

        ttk.Label(body, text="GOLD user ID is optional - only needed when two GOLD users share the name "
                             "(e.g. asmith2). Team: e.g. DPM Egypt (used to filter names in Bulk assignment and Order Assignment).",
                  style="CardMuted.TLabel", wraplength=360).grid(row=5, column=0, columnspan=2, sticky="w",
                                                                 pady=(4, 0))
        btns = ttk.Frame(body, style="Card.TFrame")
        btns.grid(row=6, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(btns, text="Cancel", command=win.destroy).pack(side="right")
        ttk.Button(btns, text="Save", style="Accent.TButton", command=ok).pack(side="right", padx=6)
        win.bind("<Return>", ok)
        win.bind("<Escape>", lambda e: win.destroy())
        win.update_idletasks()
        win.geometry(f"+{self.winfo_rootx() + 200}+{self.winfo_rooty() + 160}")
        win.grab_set()
        self.wait_window(win)
        return result or None

    def _names_add(self):
        new = self._names_dialog("Add DPM name")
        for k in ("user", "team"):
            if new and not new.get(k):
                new.pop(k, None)
        if new:
            self.dpm_names.append(new)
            self._names_saved()
            self.log(f"DPM added: {new['first']} {new['last']}")

    def _names_edit(self):
        sel = self.names_tree.selection()
        if not sel:
            messagebox.showinfo(APP_NAME, "Select a name to edit (or double-click it).")
            return
        entry = self.dpm_names[int(sel[0])]
        new = self._names_dialog("Edit DPM name", entry)
        if new:
            old = gn.full_name(entry)
            entry.update(new)
            for k in ("user", "team"):
                if not entry.get(k):
                    entry.pop(k, None)
            entry.pop("check", None)
            self._names_saved()
            self.log(f"DPM edited: {old} -> {gn.full_name(entry)}")

    def _names_delete(self):
        sel = self.names_tree.selection()
        if not sel:
            return
        names = [gn.full_name(self.dpm_names[int(i)]) for i in sel]
        if not messagebox.askyesno(APP_NAME, f"Remove {len(names)} name(s) from the list?\n\n" +
                                   "\n".join(names[:10]) + ("\n..." if len(names) > 10 else "")):
            return
        for i in sorted((int(i) for i in sel), reverse=True):
            del self.dpm_names[i]
        self._names_saved()
        self.log(f"Removed {len(names)} DPM name(s)")

    def _names_import(self):
        path = filedialog.askopenfilename(title="Import DPM names", filetypes=[
            ("Excel / CSV", "*.xlsx *.xlsm *.csv *.txt"), ("All files", "*.*")])
        if not path:
            return
        try:
            rows = []
            if path.lower().endswith((".xlsx", ".xlsm")):
                from openpyxl import load_workbook
                ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
                rows = [[str(c or "").strip() for c in r[:2]] for r in ws.iter_rows(values_only=True) if r]
            else:
                import csv
                text = Path(path).read_text(encoding="utf-8-sig")
                rows = [r[:2] for r in csv.reader(text.splitlines())]
        except Exception as e:
            messagebox.showerror(APP_NAME, f"Could not read the file:\n{e}")
            return
        added = 0
        for r in rows:
            r = [c for c in r if c and str(c).strip()]
            if not r or re.search(r"\bname\b", " ".join(r), re.I):
                continue  # empty or header row
            first, last = (r[0], r[1]) if len(r) >= 2 else gn.split_name(r[0])
            if last and not gn.lookup(self.dpm_names, f"{first} {last}"):
                self.dpm_names.append({"first": first.strip(), "last": last.strip()})
                added += 1
        self._names_saved()
        self.log(f"Imported {added} new DPM name(s) from {Path(path).name}")
        messagebox.showinfo(APP_NAME, f"{added} new name(s) added.")

    def _names_export(self):
        ctx = self._ctx()
        path = Path(ctx["workbook"]).parent / "DPM_names.xlsx"
        try:
            gx.write_sheet(path, "DPM names", [f"{len(self.dpm_names)} names", f"Exported {datetime.now():%d-%b-%Y %H:%M}"],
                           [("DPM names", ["Name", "Team", "GOLD first name", "GOLD last name", "GOLD user", "Email",
                                           "Checked"],
                             [[gn.full_name(e), e.get("team", ""), e["first"], e["last"], e.get("user", ""),
                               e.get("email", ""), e.get("check", "")]
                              for e in self.dpm_names])])
        except PermissionError:
            messagebox.showerror(APP_NAME, f"{path.name} is open in Excel. Close it and try again.")
            return
        if messagebox.askyesno(APP_NAME, f"Saved to\n{path}\n\nOpen it now?"):
            os.startfile(path)

    def _names_check(self):
        """Look every name up in GOLD (Administration > users, read-only) and store GOLD's exact
        first/last name and user id, so Bulk Order Assignment always finds them."""
        if self.busy:
            return
        sel = self.names_tree.selection()
        entries = [self.dpm_names[int(i)] for i in sel] if sel else list(self.dpm_names)
        if not messagebox.askyesno(APP_NAME, f"Check {len(entries)} name(s) in GOLD?\n\n"
                                             "This only reads GOLD (about 5 seconds per name)."):
            return
        w = self.worker

        def job():
            w.ensure_browser()
            gt.open_user_list(w.driver, w.wait)
            ok = 0
            for n, e in enumerate(entries, 1):
                name = gn.full_name(e)
                self.ui(lambda n=n, name=name: self._set_pill(f"Checking {n}/{len(entries)}: {name}...",
                                                              ORANGE, "#2A2A2A"))
                try:
                    user, gfirst, glast = gt.check_name(w.driver, w.wait, e["first"], e["last"], e.get("user"))
                    e.update(first=gfirst, last=glast, user=user, check=f"✓ found ({datetime.now():%d-%b-%Y})")
                    ok += 1
                except Exception as ex:
                    msg = str(ex).splitlines()[0] if str(ex) else type(ex).__name__
                    e.update(check=f"✗ {msg.replace(' - check the name in the DPM Names tab', '')}")
                    if not w.session_problem(ex):
                        continue
                    w.ensure_browser()
                    gt.open_user_list(w.driver, w.wait)
            return ok

        def done(ok):
            self._names_saved()
            self.log(f"Checked {len(entries)} name(s) in GOLD: {ok} found, {len(entries) - ok} not found")

        w.submit(f"Checking {len(entries)} name(s) in GOLD", job, done)

    # ----- Org Chart tab
    OC_COLUMNS = (("name", "Name", 190), ("title", "Job title", 220), ("dept", "Department", 200),
                  ("country", "Country", 110), ("team", "GOLD team", 150), ("email", "Email", 230))

    def _org(self, fn):
        """Worker thread: run fn(org_chart) in the Org Chart tab kept open next to GOLD in the hidden
        Firefox (opened on first use, about 10 s), then switch back to GOLD."""
        w = self.worker
        w.ensure_browser()
        d = w.driver
        oc = getattr(self, "_oc", None)
        if oc is None or oc.d is not d or oc.handle not in d.window_handles:
            self.log("Opening the Org Chart...")
            oc = goc.OrgChart(d, cache_file=str(OC_CACHE), sign_in_wait=60, log=self.log)
            try:
                oc.__enter__()
            except Exception:
                oc.__exit__(None, None, None)
                raise
            self._oc = oc
        try:
            return fn(oc)
        finally:
            back = oc.back if oc.back in d.window_handles else d.window_handles[0]
            d.switch_to.window(back)

    def _build_org_tab(self, f):
        bar = self._tile(f, pad=6, fill="x").body
        tk.Label(bar, text="", bg=CARD, fg=ORANGE, font=(gui.ICONS, 13)).pack(side="left", padx=(4, 8))
        self.oc_query = tk.StringVar()
        ent = ttk.Entry(bar, textvariable=self.oc_query, width=34, font=(FONT, 10))
        ent.pack(side="left")
        ent.bind("<Return>", lambda e: self._oc_search())
        self._btn(bar, "Search", self._oc_search, "Accent.TButton").pack(side="left", padx=(8, 0))
        self.oc_all = tk.BooleanVar(value=False)
        ttk.Checkbutton(bar, text="all members", variable=self.oc_all, style="Card.TCheckbutton").pack(side="right")
        self.oc_team_btn = self._btn(bar, "👥  Order team", self._oc_order_team)
        self.oc_team_btn.pack(side="right", padx=(0, 8))

        panes = ttk.Panedwindow(f, orient="horizontal")
        panes.pack(fill="both", expand=True)
        lt = self._tile(panes)
        left = lt.body
        self.oc_title = tk.Label(left, text="", bg=CARD, fg=MUTED, font=(FONT, 9, "bold"), anchor="w")
        self.oc_title.pack(fill="x", pady=(0, 6))
        tf = tk.Frame(left, background=CARD)
        tf.pack(fill="both", expand=True)
        self.oc_tree = ttk.Treeview(tf, columns=[c for c, _, _ in self.OC_COLUMNS], show="headings",
                                    selectmode="browse")
        self.oc_tree_copy = gui.TableCopy(self.oc_tree, self.log, menu=True)  # Excel-like copying
        for c, t, w in self.OC_COLUMNS:
            self.oc_tree.heading(c, text=t, anchor="w")
            self.oc_tree.column(c, width=w, anchor="w", stretch=False)  # full width: scroll sideways
        self.oc_tree.tag_configure("odd", background=gui.ZEBRA)
        self.oc_tree.tag_configure("missing", foreground=MUTED)
        sb = ttk.Scrollbar(tf, orient="vertical", command=self.oc_tree.yview)
        hsb = ttk.Scrollbar(tf, orient="horizontal", command=self.oc_tree.xview)  # wide columns
        self.oc_tree.configure(yscrollcommand=sb.set, xscrollcommand=hsb.set)
        sb.pack(side="right", fill="y")
        hsb.pack(side="bottom", fill="x")
        self.oc_tree.pack(side="left", fill="both", expand=True)
        self.oc_tree.bind("<<TreeviewSelect>>", lambda e: self._oc_show())
        panes.add(lt, weight=3)
        ct = self._tile(panes, pad=14)
        # The person card (details + the "Reports to" chain) scrolls when it is taller than the tile
        cv = self.oc_card_canvas = tk.Canvas(ct.body, background=CARD, highlightthickness=0, borderwidth=0)
        vsb = ttk.Scrollbar(ct.body, orient="vertical", command=cv.yview)

        def set_bar(lo, hi):
            vsb.set(lo, hi)
            if float(lo) <= 0 and float(hi) >= 1:
                vsb.pack_forget()
            elif not vsb.winfo_ismapped():
                vsb.pack(side="right", fill="y", before=cv)
        cv.configure(yscrollcommand=set_bar)
        cv.pack(side="left", fill="both", expand=True)
        self.oc_card = tk.Frame(cv, background=CARD)
        win = cv.create_window(0, 0, window=self.oc_card, anchor="nw")
        cv.bind("<Configure>", lambda e: cv.itemconfigure(win, width=e.width))
        self.oc_card.bind("<Configure>", lambda e: cv.configure(scrollregion=(0, 0, e.width, e.height)))

        def wheel(e):
            w = e.widget if isinstance(e.widget, tk.Misc) else None
            while w is not None and w is not cv:
                w = getattr(w, "master", None)
            if w is cv and cv.yview() != (0.0, 1.0):
                cv.yview_scroll(int(-e.delta / 120), "units")
        cv.bind_all("<MouseWheel>", wheel, add="+")
        panes.add(ct, weight=2)

        def place_sash(e=None):  # give the person card ~430 px, once the tab has its real size
            if panes.winfo_width() > 900:
                panes.sashpos(0, panes.winfo_width() - 430)
                panes.unbind("<Configure>")
        panes.bind("<Configure>", place_sash)
        self.oc_people = []
        self._oc_fill([], "Search for a person, or look up the team of the order you have open.")

    def _oc_search(self, text=None):
        q = (text if text is not None else self.oc_query.get()).strip()
        if not q or self.busy:
            return
        self.oc_query.set(q)
        self.notebook.select(self.org_frame)
        self.worker.submit(f"Searching the Org Chart for {q}",
                           lambda: self._org(lambda oc: oc.query(q)),
                           lambda people: self._oc_fill(people, f"{len(people)} result(s) for “{q}”" if people
                                                        else f"Nobody found for “{q}”. Try the last name only."),
                           retry=True)

    def _oc_lookup_member(self, event=None):
        """Explorer > Teams: double-click a member to look them up in the Org Chart."""
        sel = self.members.curselection()
        if sel:
            self._oc_search(self.members.get(sel[0]).replace("★", "").strip())

    def _oc_order_team(self):
        order = self.order or (self._ask_context_order() if hasattr(self, "ask_input") else "")
        if not order:
            messagebox.showinfo(APP_NAME, "Open an order in the Order explorer first.")
            return
        if self.busy:
            return
        teams = self.teams if self.order == order and self.teams else None
        everyone = self.oc_all.get()

        def job():
            nonlocal teams
            if not teams:
                self._ensure_order(order)
                w = self.worker
                oo.open_tab(w.driver, w.wait, "Teams")
                teams = oo.read_teams(w.driver)
            # By default: each team's primary member, and everyone in small teams (big pools such as
            # Order Entry have ~100 people). "everyone" looks them all up (slower).
            roles = {}
            for team, t in teams.items():
                for m in t["members"]:
                    if everyone or m == t["primary"] or len(t["members"]) <= 5:
                        roles.setdefault(m, []).append(team + (" (primary)" if m == t["primary"] else ""))
            found = self._org(lambda oc: oc.lookup(list(roles)))
            people = []
            for name, team_list in roles.items():
                e = found.get(name) or {}
                p = dict(e.get("person") or {"display_name": name})
                p["chain"] = [e["manager"]] if e.get("manager") else []
                p["team"] = ", ".join(team_list)
                p["note"] = e.get("note", "" if e else "not a person")
                people.append(p)
            people.sort(key=lambda p: ("(primary)" not in p["team"], p["team"]))
            return people

        self.notebook.select(self.org_frame)
        self.worker.submit(f"Looking up the team of {order}", job,
                           lambda people: self._oc_fill(people, f"Team of order {order}: {len(people)} people"),
                           retry=True)

    def _oc_fill(self, people, title):
        self.oc_people = people
        self.oc_title.configure(text=title.upper() if people else "")
        self.oc_tree.delete(*self.oc_tree.get_children())
        has_team = any(p.get("team") for p in people)
        self.oc_tree.configure(displaycolumns=[c for c, _, _ in self.OC_COLUMNS if has_team or c != "team"])
        for i, p in enumerate(people):
            tags = ["odd"] if i % 2 else []
            if not p.get("email"):
                tags.append("missing")
            self.oc_tree.insert("", "end", iid=str(i), tags=tags, values=(
                p.get("display_name", ""), p.get("busi_title") or p.get("note", ""), p.get("dept_desc", ""),
                p.get("country_descr", ""), p.get("team", ""), p.get("email", "")))
        if people:
            self.oc_tree.selection_set("0")
        else:
            self._oc_card_message(title)

    def _oc_card_message(self, text):
        for w in self.oc_card.winfo_children():
            w.destroy()
        tk.Label(self.oc_card, text=text, bg=CARD, fg=MUTED, font=(FONT, 10), wraplength=360,
                 justify="left").pack(anchor="nw", pady=24)

    def _oc_show(self):
        sel = self.oc_tree.selection()
        if not sel:
            return
        p = self.oc_people[int(sel[0])]
        c = self.oc_card
        for w in c.winfo_children():
            w.destroy()
        self.oc_card_canvas.yview_moveto(0)
        tk.Label(c, text=p.get("display_name", ""), bg=CARD, fg=INK, font=(FONT, 14, "bold"),
                 anchor="w").pack(fill="x")
        tk.Label(c, text=p.get("busi_title") or p.get("note", ""), bg=CARD, fg=MUTED, font=(FONT, 10),
                 anchor="w", wraplength=380, justify="left").pack(fill="x", pady=(0, 10))
        rows = [("Email", p.get("email")), ("Phone", p.get("phone")), ("Mobile", p.get("mobile")),
                ("Department", p.get("dept_desc")),
                ("Location", ", ".join(x for x in (p.get("city"), p.get("country_descr")) if x)),
                ("Company", p.get("vendor_name")), ("GOLD team", p.get("team")), ("CUID", p.get("cuid"))]
        grid = tk.Frame(c, background=CARD)
        grid.pack(fill="x")
        for r, (label, value) in enumerate((l, v) for l, v in rows if v):
            tk.Label(grid, text=label, bg=CARD, fg=MUTED, font=(FONT, 9), anchor="nw").grid(
                row=r, column=0, sticky="nw", pady=2)
            v = tk.Entry(grid, font=(FONT, 10), relief="flat", readonlybackground=CARD, fg=INK,
                         highlightthickness=0, borderwidth=0)
            v.insert(0, value)
            v.configure(state="readonly")
            v.grid(row=r, column=1, sticky="ew", padx=(14, 0), pady=2)
        grid.columnconfigure(1, weight=1)

        btns = tk.Frame(c, background=CARD)
        btns.pack(fill="x", pady=(10, 4))
        if p.get("email"):
            ttk.Button(btns, text="✉  Email", command=lambda: os.startfile(f"mailto:{p['email']}")).pack(side="left")
            ttk.Button(btns, text="Copy email", command=lambda: self._oc_copy(p["email"])).pack(side="left", padx=6)
        ttk.Button(btns, text="Copy all", command=lambda: self._oc_copy(self._oc_details(p))).pack(side="left")

        chain = p.get("chain") or []
        if chain:
            tk.Label(c, text="REPORTS TO", bg=CARD, fg=MUTED, font=(FONT, 8, "bold"), anchor="w").pack(
                fill="x", pady=(14, 4))
            for level, m in enumerate(chain):
                row = tk.Frame(c, background=CARD)
                row.pack(fill="x", pady=1)
                tk.Frame(row, background=ORANGE if level == 0 else LINE, width=3).pack(side="left", fill="y",
                                                                                     padx=(level * 14, 10))
                txt = tk.Frame(row, background=CARD)
                txt.pack(side="left", fill="x", expand=True, pady=3)
                name = tk.Label(txt, text=m.get("display_name", ""), bg=CARD, fg=ORANGE_DARK,
                                font=(FONT, 10, "bold" if level == 0 else "normal"), anchor="w", cursor="hand2")
                name.pack(fill="x")
                name.bind("<Button-1>", lambda e, m=m: self._oc_search(m.get("email") or m.get("display_name")))
                tk.Label(txt, text=m.get("busi_title", ""), bg=CARD, fg=MUTED, font=(FONT, 9), anchor="w",
                         wraplength=max(320 - level * 14, 180), justify="left").pack(fill="x")

    @staticmethod
    def _oc_details(p):
        lines = [p.get("display_name", ""), p.get("busi_title", "")]
        for label, key in (("Email", "email"), ("Phone", "phone"), ("Mobile", "mobile"),
                           ("Department", "dept_desc"), ("Country", "country_descr")):
            if p.get(key):
                lines.append(f"{label}: {p[key]}")
        if p.get("chain"):
            m = p["chain"][0]
            lines.append(f"Manager: {m.get('display_name', '')} ({m.get('email', '')})")
        return "\n".join(x for x in lines if x)

    def _oc_copy(self, text):
        self.clipboard_clear()
        self.clipboard_append(text)
        self.log("Copied.")

    # ----- Ask GOLD tab (AI assistant)
    CHAT_BG = CARD
    ASK_INDENT = 40  # answers line up under the assistant's name, right of its avatar
    # (title, description, question) - "{order}" is the order open in the Order explorer
    ASK_SUGGESTIONS = (("Deadlines", "Target dates, milestones, lead time",
                        "What are the deadlines and milestones of {order}?"),
                       ("Pricing", "Monthly and one-off charges, order value",
                        "Show the pricing details of {order}."),
                       ("Site", "Address, site ID and references",
                        "What is the site address of {order}?"),
                       ("People", "STM, owners, technical support",
                        "Who is working on {order}?"))

    def _build_ask_tab(self, f):
        self.ask_frame = f
        scale = self.winfo_fpixels("1i") / 96
        f.columnconfigure(1, weight=1)
        f.rowconfigure(0, weight=1)
        # Saved chats, grouped by order number (left)
        ht = self._tile(f, pad=6)
        ht.grid(row=0, column=0, sticky="nsew", padx=(0, 12))
        ht.body.configure(width=int(250 * scale))
        self._build_ask_history(ht.body)
        ct = self._tile(f, pad=8)
        ct.grid(row=0, column=1, sticky="nsew")
        f = ct.body  # the chat itself (right)
        # Header: what it does | order in context, New chat
        head = tk.Frame(f, background=CARD, padx=12, pady=4)
        head.pack(fill="x")
        bot_png = self._asset_path("bot.png")
        if bot_png:  # the bot itself, floating next to its name (click it!)
            gui.Bot(head, bot_png, size=int(34 * scale), bg=CARD).pack(side="left", padx=(0, 8))
        tk.Label(head, text="The Friendly GOLD bot", bg=CARD, fg=INK, font=(FONT, 11, "bold")).pack(side="left")
        self.ask_hint = tk.Label(head, text="   reads GOLD and the Org Chart live for you - it never changes "
                                            "anything", bg=CARD, fg=MUTED, font=(FONT, 9), anchor="w")
        self.ask_hint.pack(side="left", fill="x", expand=True)
        self._btn(head, "＋  New chat", self._ask_new).pack(side="right")
        self.ask_chip = gui.Pill(head, "", ORANGE_DARK, ORANGE_SOFT, bg=CARD)

        # Composer (bottom): a rounded field with a round send button
        foot = tk.Frame(f, background=CARD, padx=8, pady=4)
        foot.pack(side="bottom", fill="x")
        comp = gui.Tile(foot, radius=14, pad=0, bg=CARD, fill=gui.FIELD, shadow=False, outline=gui.FIELD_LINE)
        comp.pack(fill="x")
        box = comp.body
        d = int(28 * scale)
        self._send_imgs = {s: ImageTk.PhotoImage(gui.rounded_image(d, d, d // 2, c), master=self)
                           for s, c in (("on", ORANGE), ("hover", ORANGE_DARK), ("off", gui.ACCENT_OFF))}
        self.ask_send = tk.Label(box, image=self._send_imgs["on"], text="", compound="center",
                                 fg="white", font=(gui.ICONS, 9), bg=gui.FIELD, cursor="hand2", bd=0)
        self.ask_send.pack(side="right", padx=(6, 0))  # packed first: always visible
        self.ask_send.bind("<Enter>", lambda e: self.busy or self.ask_send.configure(image=self._send_imgs["hover"]))
        self.ask_send.bind("<Leave>", lambda e: self.busy or self.ask_send.configure(image=self._send_imgs["on"]))
        # one line to start with; it grows (up to 5 lines) as the question gets longer
        self.ask_input = tk.Text(box, height=1, wrap="word", font=(FONT, 10), relief="flat", padx=6, pady=3,
                                 background=gui.FIELD, foreground=INK, insertbackground=INK, highlightthickness=0,
                                 borderwidth=0)
        self.ask_input.pack(side="left", fill="both", expand=True)
        self.ask_send.bind("<Button-1>", lambda e: self._ask_send())
        self.ask_placeholder = tk.Label(box, bg=gui.FIELD, fg=gui.PLACEHOLDER, font=(FONT, 10), cursor="xterm")
        self.ask_placeholder.place(x=8, y=3)
        self.ask_placeholder.bind("<Button-1>", lambda e: self.ask_input.focus_set())
        self.ask_input.bind("<Return>", lambda e: (self._ask_send(), "break")[1])
        self.ask_input.bind("<Shift-Return>", lambda e: self.after(1, self._ask_input_changed))
        self.ask_input.bind("<KeyRelease>", lambda e: self._ask_input_changed())
        comp_outline = {True: ORANGE, False: gui.FIELD_LINE}
        for ev, on in (("<FocusIn>", True), ("<FocusOut>", False)):
            self.ask_input.bind(ev, lambda e, on=on: (setattr(comp, "outline", comp_outline[on]),
                                                       setattr(comp, "_size", None), comp._soon()))
        hints = tk.Frame(foot, background=CARD)
        hints.pack(fill="x", pady=(6, 0))
        self.ask_status = tk.Label(hints, text="", bg=CARD, fg=MUTED, font=(FONT, 8))
        self.ask_status.pack(side="left")
        tk.Label(hints, text="Enter to send  ·  Shift+Enter for a new line  ·  right-click a table to copy it",
                 bg=CARD, fg=gui.PLACEHOLDER, font=(FONT, 8)).pack(side="right")

        # Conversation
        body = tk.Frame(f, background=self.CHAT_BG)
        body.pack(fill="both", expand=True)
        self.ask_text = tk.Text(body, wrap="word", relief="flat", background=self.CHAT_BG, foreground=INK,
                                padx=20, pady=10, cursor="arrow", state="disabled", font=(FONT, 10),
                                highlightthickness=0, borderwidth=0, selectbackground=ORANGE_SOFT,
                                selectforeground=INK, inactiveselectbackground=ORANGE_SOFT)
        sb = ttk.Scrollbar(body, orient="vertical", command=self.ask_text.yview)
        self.ask_text.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.ask_text.pack(side="left", fill="both", expand=True)
        t = self.ask_text
        gmd.setup(t, font=FONT, ink=INK, muted=MUTED, accent=ORANGE, soft=ORANGE_SOFT, line=LINE,
                  indent=self.ASK_INDENT)
        t.tag_configure("right", justify="right", spacing1=14, spacing3=2, rmargin=4)
        t.tag_configure("meta", font=(FONT, 8), foreground=gui.PLACEHOLDER)
        t.tag_configure("who", font=(FONT, 10, "bold"), foreground=INK, spacing1=16, spacing3=4)
        t.tag_raise("meta")
        t.tag_configure("center", justify="center")
        t.tag_configure("hero", font=(FONT, 13, "bold"), foreground=INK, justify="center", spacing1=4)
        t.tag_configure("sub", font=(FONT, 10), foreground=MUTED, justify="center", spacing1=4, spacing3=16)
        t.tag_configure("err", foreground=RED, font=(FONT, 10), lmargin1=self.ASK_INDENT + 2,
                        lmargin2=self.ASK_INDENT + 2, spacing1=2, spacing3=6)
        t.tag_configure("action", font=(FONT, 8, "bold"), foreground=MUTED, lmargin1=self.ASK_INDENT + 2,
                        spacing1=4, spacing3=6)
        # Mouse wheel over embedded tables / bubbles scrolls the chat
        t.bind_all("<MouseWheel>", self._ask_wheel, add="+")
        self._ask_size_cards = lambda: None  # the welcome screen's cards follow the chat's width
        t.bind("<Configure>", lambda e: self._ask_size_cards(), add="+")
        self.notebook.bind("<<NotebookTabChanged>>", lambda e: self._ask_tab_shown(), add="+")
        self.ask_order = ""     # order the welcome screen / chip was drawn for
        self.ask_pending = None  # "thinking" row while a question runs
        self._ask_welcome()

    # --- helpers
    def _ask_context_order(self):
        """The order the user is looking at: open in the Order explorer, else typed there."""
        if self.order:
            return self.order
        typed = self.order_var.get().strip()
        return typed if re.fullmatch(r"\d{2}-\d{5,8}", typed) else ""

    def _ask_tab_shown(self):
        if self.notebook.current() is not self.ask_frame:
            return
        order = self._ask_context_order()
        if order != self.ask_order and not (self.chat and self.chat.messages) and not self.busy:
            self._ask_welcome()
        self._ask_refresh_chip()
        self.ask_input.focus_set()

    def _ask_refresh_chip(self):
        order = self._ask_context_order()
        if order:
            self.ask_chip.set(f"Order in view:  {order}")
            self.ask_chip.pack(side="right", padx=(0, 12))
            self.ask_hint.pack_forget()  # packed again last: the pill keeps its size, the hint gets the rest
            self.ask_hint.pack(side="left", fill="x", expand=True)
        else:
            self.ask_chip.pack_forget()
        self._ask_input_changed()

    def _ask_input_changed(self):
        try:  # grow with the question: 1 to 5 lines
            lines = self.ask_input.count("1.0", "end-1c", "displaylines")
            lines = (lines[0] if isinstance(lines, tuple) else lines or 0) + 1
        except tk.TclError:
            lines = 1
        self.ask_input.configure(height=max(1, min(5, lines)))
        if self.ask_input.get("1.0", "end").strip():
            self.ask_placeholder.place_forget()
        else:
            order = self._ask_context_order()
            self.ask_placeholder.configure(text=f"Ask about {order} or any other order..." if order
                                           else "Ask about a GOLD order, e.g. deadlines of 10-1234567")
            self.ask_placeholder.place(x=8, y=3)

    def _ask_wheel(self, e):
        w = e.widget
        while w is not None and w is not self.ask_text:
            w = getattr(w, "master", None)
        if w is self.ask_text and e.widget is not self.ask_text:
            self.ask_text.yview_scroll(int(-e.delta / 120), "units")

    def _ask_write(self, fn):
        t = self.ask_text
        t.configure(state="normal")
        fn(t)
        t.configure(state="disabled")
        t.see("end")

    def _ask_clear(self):
        def clear(t):
            t.delete("1.0", "end")
            for w in t.winfo_children():
                w.destroy()
        self._ask_write(clear)

    def _ask_avatar(self, t):
        if not getattr(self, "_avatar_img", None):
            self._avatar_img = self._asset_image("bot_28.png") or self._asset_image("logo_28.png")
        t.window_create("end", window=tk.Label(t, image=self._avatar_img, bg=self.CHAT_BG, bd=0)
                        if self._avatar_img else tk.Label(t, text="GX", bg=ORANGE, fg="white",
                                                          font=(FONT, 8, "bold"), width=3, pady=2))
        t.insert("end", "   ", ("who",))

    # --- welcome screen with suggestions
    def _ask_welcome(self):
        self._ask_clear()
        order = self.ask_order = self._ask_context_order()

        def write(t):
            t.insert("end", "\n", ("center",))
            bot_png = self._asset_path("bot.png")
            scale = self.winfo_fpixels("1i") / 96
            t.window_create("end", window=gui.Bot(t, bot_png, size=int(118 * scale), bg=self.CHAT_BG)
                            if bot_png else tk.Label(t, text="✦", bg=ORANGE, fg="white", font=(FONT, 18, "bold")))
            t.tag_add("center", "end-2c")
            hour = datetime.now().hour
            hello = "Good morning" if hour < 12 else "Good afternoon" if hour < 18 else "Good evening"
            t.insert("end", f"\n{hello}! I'm the Friendly GOLD bot\n", ("hero",))
            t.insert("end", (f"Ask me anything about order {order} (open in the Order explorer) or any other order."
                             if order else "Ask me anything about a GOLD order - deadlines, pricing, sites, people.")
                     + "\n", ("sub",))
            grid = tk.Frame(t, background=self.CHAT_BG)
            cards = []
            for i, (title, desc, q) in enumerate(self.ASK_SUGGESTIONS):
                question = q.format(order=order or "10-1234567")
                card = gui.ClickCard(grid, title, desc, lambda q=question: self._ask_suggestion(q),
                                     bg=self.CHAT_BG, width=int(200 * scale))
                card.grid(row=i // 2, column=i % 2, padx=6, pady=6, sticky="nsew")
                cards.append(card)

            def size_cards(e=None):  # two cards per row, as wide as the chat allows
                width = t.winfo_width()
                if width < 200 or not grid.winfo_exists():
                    return
                w = min(int(380 * scale), (width - int(70 * scale)) // 2)
                for c in cards:
                    c.configure(width=max(w, int(160 * scale)))
            self._ask_size_cards = size_cards
            self.after(50, size_cards)
            t.window_create("end", window=grid)
            t.tag_add("center", "end-2c")
            t.insert("end", "\n", ("center",))
        self._ask_write(write)
        self.ask_text.see("1.0")
        self._ask_refresh_chip()

    def _ask_suggestion(self, q):
        self.ask_input.delete("1.0", "end")
        self.ask_input.insert("1.0", q)
        if self._ask_context_order():
            self._ask_send()
        else:  # no order in view: let the user put in the order number
            self.ask_input.tag_add("sel", "1.0", "end-1c")
            self.ask_input.focus_set()
            self._ask_input_changed()

    def _ask_new(self):
        if self.busy:
            return
        if self.chat:
            self.chat.messages = []
        self.ask_restore, self.ask_log, self.ask_chat_id = None, [], None
        self._chat_fill_history()
        self._ask_welcome()

    # --- saved chats (Ask GOLD history), one JSON file per chat in <data folder>\chats
    ORDER_RE = re.compile(r"\b\d{2}-\d{6,8}\b")

    def _build_ask_history(self, f):
        self.ask_log, self.ask_chat_id, self.ask_restore = [], None, None
        self.save_chats = tk.BooleanVar(value=self.settings.get("save_chats", True))
        head = self._tile_head(f, "History", "")
        ttk.Checkbutton(f, text="Save chats", variable=self.save_chats, style="Card.TCheckbutton",
                        command=lambda: (self.settings.update(save_chats=self.save_chats.get()),
                                         save_settings(self.settings))).pack(anchor="w", pady=(0, 6))
        self.chat_find = tk.StringVar()
        ent = ttk.Entry(f, textvariable=self.chat_find, font=(FONT, 9))
        ent.pack(fill="x", pady=(0, 6))
        ent.bind("<KeyRelease>", lambda e: self._chat_fill_history())
        tf = tk.Frame(f, background=CARD)
        tf.pack(fill="both", expand=True)
        self.chat_tree = ttk.Treeview(tf, show="tree", selectmode="browse")
        self.chat_tree.column("#0", width=200, stretch=True)
        self.chat_tree.tag_configure("group", font=(FONT, 9, "bold"), foreground=INK)
        self.chat_tree.tag_configure("chat", font=(FONT, 9), foreground=MUTED)
        self.chat_tree.tag_configure("current", foreground=ORANGE_DARK)
        sb = ttk.Scrollbar(tf, orient="vertical", command=self.chat_tree.yview)
        self.chat_tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.chat_tree.pack(side="left", fill="both", expand=True)
        self.chat_tree.bind("<<TreeviewSelect>>", lambda e: self._chat_open_selected())
        self.chat_tree.bind("<Button-3>", self._chat_menu)
        tk.Label(f, text="Grouped by the order numbers in each chat", bg=CARD, fg=FAINT, font=(FONT, 8),
                 anchor="w").pack(fill="x", pady=(4, 0))
        self._chat_fill_history()

    @staticmethod
    def _plain_messages(messages):
        """Claude's conversation as plain JSON (the SDK's content blocks become dicts)."""
        out = []
        for m in messages:
            content = m["content"]
            if not isinstance(content, str):
                content = [b.model_dump(exclude_none=True) if hasattr(b, "model_dump") else b for b in content]
            out.append({"role": m["role"], "content": content})
        return out

    def _chat_save(self):
        if not self.save_chats.get() or not self.ask_log:
            return
        if not self.ask_chat_id:
            self.ask_chat_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        orders = set()
        for e in self.ask_log:
            orders.update(self.ORDER_RE.findall(e["you"] + " " + e["answer"]))
            if e.get("order"):
                orders.add(e["order"])
        data = {"id": self.ask_chat_id, "title": self.ask_log[0]["you"][:80], "orders": sorted(orders),
                "updated": datetime.now().isoformat(timespec="minutes"), "log": self.ask_log,
                "messages": self._plain_messages(self.chat.messages) if self.chat else []}
        try:
            CHATS_DIR.mkdir(parents=True, exist_ok=True)
            (CHATS_DIR / f"{self.ask_chat_id}.json").write_text(json.dumps(data, ensure_ascii=False, default=str),
                                                               encoding="utf-8")
        except OSError as e:
            self.log(f"Could not save the chat: {e}")
        self._chat_fill_history()

    def _chat_list(self):
        chats = []
        for p in CHATS_DIR.glob("*.json") if CHATS_DIR.exists() else []:
            try:
                d = json.loads(p.read_text(encoding="utf-8"))
                chats.append({"id": d["id"], "title": d.get("title", ""), "orders": d.get("orders", []),
                              "updated": d.get("updated", "")})
            except (OSError, ValueError, KeyError):
                continue
        return sorted(chats, key=lambda c: c["updated"], reverse=True)

    def _chat_fill_history(self):
        if not hasattr(self, "chat_tree"):
            return
        q = self.chat_find.get().strip().lower()
        t = self.chat_tree
        open_groups = {i for i in t.get_children() if t.item(i, "open")}
        t.delete(*t.get_children())
        groups = {}
        for c in self._chat_list():
            if q and q not in (c["title"] + " " + " ".join(c["orders"])).lower():
                continue
            for o in c["orders"] or ["Other chats"]:
                groups.setdefault(o, []).append(c)
        # orders with the newest chat first; "Other chats" last
        for o in sorted(groups, key=lambda o: (o == "Other chats", -max(
                datetime.fromisoformat(c["updated"]).timestamp() if c["updated"] else 0 for c in groups[o]))):
            gid = f"g:{o}"
            t.insert("", "end", iid=gid, text=f"{o}   ({len(groups[o])})", tags=("group",),
                     open=gid in open_groups or bool(q) or o == self._ask_context_order())
            for c in groups[o]:
                when = c["updated"][5:16].replace("T", " ") if c["updated"] else ""
                tags = ("chat", "current") if c["id"] == self.ask_chat_id else ("chat",)
                t.insert(gid, "end", iid=f"c:{o}:{c['id']}", text=f"{c['title'][:40]}  ·  {when}", tags=tags)

    def _chat_open_selected(self):
        sel = self.chat_tree.selection()
        if not sel or not sel[0].startswith("c:") or self.busy:
            return
        chat_id = sel[0].rsplit(":", 1)[1]
        if chat_id == self.ask_chat_id:
            return
        try:
            d = json.loads((CHATS_DIR / f"{chat_id}.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return
        self.ask_chat_id, self.ask_log = chat_id, d.get("log", [])
        # continue the conversation: Claude gets the earlier messages back
        if self.chat:
            self.chat.messages = d.get("messages", [])
        else:
            self.ask_restore = d.get("messages", [])
        self._ask_clear()
        for e in self.ask_log:
            self._ask_user_message(e["you"])
            self._ask_answer(e["answer"], e.get("ok", True))
        self.ask_text.see("end")
        self._chat_fill_history()

    def _chat_menu(self, e):
        iid = self.chat_tree.identify_row(e.y)
        if not iid or not iid.startswith("c:"):
            return
        chat_id = iid.rsplit(":", 1)[1]

        def delete():
            try:
                (CHATS_DIR / f"{chat_id}.json").unlink()
            except OSError:
                pass
            if chat_id == self.ask_chat_id:
                self._ask_new()
            self._chat_fill_history()
        m = tk.Menu(self, tearoff=0)
        m.add_command(label="Delete this chat", command=delete)
        m.tk_popup(e.x_root, e.y_root)

    # --- messages
    def _ask_user_message(self, question):
        def write(t):
            stamp = datetime.now().strftime("%H:%M")
            t.mark_set("question", "end-1c")
            t.mark_gravity("question", "left")
            t.insert("end", f"You · {stamp}\n", ("right", "meta"))
            width = max(300, int(t.winfo_width() * 0.62))
            t.window_create("end", window=gui.Bubble(t, question, width, bg=self.CHAT_BG))
            t.tag_add("right", "end-2c")
            t.insert("end", "\n", ("right",))
        self._ask_write(write)

    def _ask_thinking(self, text="Thinking..."):
        """Show (or update) the assistant's 'working' row under the last question."""
        if self.ask_pending is None:
            def write(t):
                t.mark_set("pending", "end-1c")
                t.mark_gravity("pending", "left")
                t.insert("end", "\n", ("who",))
                bot_png = self._asset_path("bot.png")
                if bot_png:  # the bot bobs along while it works
                    t.window_create("end", window=gui.Bot(t, bot_png, size=int(40 * self.winfo_fpixels("1i") / 96),
                                                          bg=self.CHAT_BG, mode="busy", clickable=False))
                    t.insert("end", "  ", ("who",))
                else:
                    self._ask_avatar(t)
                lab = tk.Label(t, bg=self.CHAT_BG, fg=MUTED, font=(FONT, 9, "italic"))
                t.window_create("end", window=lab)
                t.insert("end", "\n", ("who",))
                self.ask_pending = lab
            self._ask_write(write)
            self._ask_pulse(0)
        self.ask_pending.configure(text=text)
        self.ask_pending.base = text

    def _ask_pulse(self, n):
        lab = self.ask_pending
        if lab is None or not lab.winfo_exists():
            return
        base = getattr(lab, "base", "").rstrip(".")
        lab.configure(text=base + "." * (n % 3 + 1))
        self.after(450, self._ask_pulse, n + 1)

    def _ask_end_thinking(self):
        if self.ask_pending is not None:
            self.ask_pending = None
            self._ask_write(lambda t: t.delete("pending", "end-1c"))

    def _ask_answer(self, answer, ok=True):
        self._ask_end_thinking()

        def write(t):
            t.insert("end", "\n", ("who",))
            self._ask_avatar(t)
            t.insert("end", "GOLD bot", ("who",))
            t.insert("end", "   " + datetime.now().strftime("%H:%M") + "\n", ("who", "meta"))
            if not ok:
                t.insert("end", "⚠  " + answer + "\n", ("err",))
                return
            gmd.render(t, answer, max(500, t.winfo_width() - 60))
            tag = f"copy{len(t.tag_names())}"
            t.insert("end", "Copy answer", ("action", tag))
            t.insert("end", "\n", ("action",))
            text = gmd.plain(answer)
            t.tag_bind(tag, "<Button-1>", lambda e: (self.clipboard_clear(), self.clipboard_append(text),
                                                      self.log("Answer copied (tables paste into Excel as columns).")))
            t.tag_bind(tag, "<Enter>", lambda e: (t.configure(cursor="hand2"), t.tag_configure(tag, foreground=ORANGE)))
            t.tag_bind(tag, "<Leave>", lambda e: (t.configure(cursor="arrow"), t.tag_configure(tag, foreground=MUTED)))
        self._ask_write(write)
        # Long answers are read from the top: scroll so the question is at the top of the view.
        self.ask_text.update_idletasks()
        self.ask_text.yview("question")

    def _ask_send(self):
        question = self.ask_input.get("1.0", "end").strip()
        if not question or self.busy:
            return
        self.ask_input.delete("1.0", "end")
        self._ask_input_changed()
        if not (self.chat and self.chat.messages):
            self._ask_clear()
        self._ask_user_message(question)
        self._ask_thinking("Thinking")
        self.ask_send.configure(image=self._send_imgs["off"], cursor="watch")
        chosen = self.ai_key_file.get().strip()
        order = self._ask_context_order()
        context = ""
        if order:
            banner = self.order_label.cget("text") if order == self.order else ""
            context = (f"Order open in the Order explorer: {order}" + (f" ({banner})" if banner else "")
                       + ". Questions that name no order are about it.")

        def progress(text):
            written.clear()
            self.ui(lambda: self._ask_thinking(text.rstrip(".")))

        written = []

        def typing(piece):  # the answer as it is being written: show its latest words straight away
            written.append(piece)
            preview = " ".join("".join(written).replace("*", "").replace("#", "").split())
            self.ui(lambda p=preview: self._ask_thinking("Writing: " + (p[-110:] if len(p) > 110 else p)))

        def driver():
            self.worker.ensure_browser()
            return self.worker.driver

        def connect(fresh=False):
            if fresh:  # the saved key was refused (e.g. replaced by a new one): fetch it again
                gai.forget_cached_key(AI_KEY_CACHE)
            key = gai.get_key(chosen, AI_KEY_CACHE, driver, self.log)
            history = self.chat.messages if self.chat else getattr(self, "ask_restore", None) or []
            self.ask_restore = None
            fast = self.settings.get("bot_speed", "fast") == "fast"  # Settings > GOLD bot answers
            self.chat = gai.Chat(key, gai.GoldReader(self), effort="low" if fast else "medium", fast=fast)
            self.chat.messages = history

        def job():
            try:
                if self.chat is None:
                    progress("Connecting to the AI service")
                    connect()
                try:
                    return True, self.chat.ask(question, progress, context, on_text=typing)
                except Exception as e:
                    if not isinstance(e, self.chat.anthropic.AuthenticationError) or chosen:
                        raise
                    connect(fresh=True)
                    return True, self.chat.ask(question, progress, context, on_text=typing)
            except Exception as e:  # shown in the chat, not as a pop-up
                (HERE / "gold_app_error.log").write_text(traceback.format_exc(), encoding="utf-8")
                return False, gai.friendly_error(e)

        def done(result):
            ok, answer = result
            self.ask_send.configure(image=self._send_imgs["on"], cursor="hand2")
            self._ask_answer(answer, ok)
            self.ask_input.focus_set()
            self.ask_log.append({"you": question, "answer": answer, "ok": ok, "order": order,
                                 "time": datetime.now().isoformat(timespec="minutes")})
            self._chat_save()

        self.worker.submit("Asking GOLD", job, done)

    def _build_text_view(self):
        f = self.text_view = ttk.Frame(self.content)
        self.text = tk.Text(f, wrap="word", font=("Segoe UI", 10), relief="flat",
                            background=CARD, foreground=INK, insertbackground=INK, padx=10, pady=8)
        sb = ttk.Scrollbar(f, command=self.text.yview)
        tab = int(260 * self.winfo_fpixels("1i") / 96)
        self.text.configure(yscrollcommand=sb.set, state="disabled", tabs=(str(tab),), highlightthickness=0,
                            cursor="arrow")
        self.text.tag_configure("section", font=("Segoe UI Semibold", 11), foreground=ORANGE,
                                spacing1=10, spacing3=4)
        self.text.tag_configure("label", font=(FONT, 10), foreground=MUTED, spacing1=2, spacing3=2)
        self.text.tag_configure("value", font=(FONT, 10), foreground=INK, lmargin2=tab + 8)
        self._grid_fits = []  # tables in the view follow its width
        self.text.bind("<Configure>", lambda e: [fit() for fit in self._grid_fits])
        sb.pack(side="right", fill="y")
        self.text.pack(fill="both", expand=True)

    def _build_docs_view(self):
        f = self.docs_view = ttk.Frame(self.content)
        bar = ttk.Frame(f)
        bar.pack(fill="x", pady=(0, 6))
        ttk.Label(bar, text="Select documents (Ctrl/Shift-click for several)").pack(side="left")
        self._btn(bar, "Download all", lambda: self.download(all_docs=True)).pack(side="right")
        self._btn(bar, "Download selected", self.download, "Accent.TButton").pack(side="right", padx=6)
        cols = ("n", "title", "modified", "type", "saved")
        self.docs_tree = ttk.Treeview(f, columns=cols, show="headings", selectmode="extended")
        self.docs_tree_copy = gui.TableCopy(self.docs_tree, self.log, menu=True)  # Excel-like copying
        for c, w, t in [("n", 40, "#"), ("title", 330, "Document"), ("modified", 170, "Last modified"),
                        ("type", 190, "Type"), ("saved", 200, "Downloaded")]:
            self.docs_tree.heading(c, text=t)
            self.docs_tree.column(c, width=w, anchor="w", stretch=c in ("title", "saved"))
        sb = ttk.Scrollbar(f, command=self.docs_tree.yview)
        self.docs_tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.docs_tree.pack(fill="both", expand=True)

    def _build_teams_view(self):
        f = self.teams_view = ttk.Frame(self.content)
        bar = ttk.Frame(f)
        bar.pack(fill="x", pady=(0, 6))
        ttk.Label(bar, text="Find a person").pack(side="left")
        self.find_var = tk.StringVar()
        e = ttk.Entry(bar, textvariable=self.find_var, width=28)
        e.pack(side="left", padx=6)
        e.bind("<KeyRelease>", lambda ev: self._find_person())
        self.find_result = ttk.Label(bar, text="", foreground=MUTED)
        self.find_result.pack(side="left", padx=6)

        panes = ttk.Panedwindow(f, orient="horizontal")
        panes.pack(fill="both", expand=True)
        lf = ttk.Frame(panes)
        self.teams_tree = ttk.Treeview(lf, columns=("team", "primary", "count"), show="headings",
                                       selectmode="extended")
        self.teams_tree_copy = gui.TableCopy(self.teams_tree, self.log, menu=True)  # Excel-like copying
        for c, w, t in [("team", 230, "Team"), ("primary", 200, "Primary member"), ("count", 80, "Members")]:
            self.teams_tree.heading(c, text=t)
            self.teams_tree.column(c, width=w, anchor="w")
        self.teams_tree.pack(fill="both", expand=True)
        self.teams_tree.bind("<<TreeviewSelect>>", lambda e: self._team_selected())
        panes.add(lf, weight=3)
        rf = ttk.Frame(panes, padding=(8, 0, 0, 0))
        self.members_title = ttk.Label(rf, text="Members", style="Head.TLabel")
        self.members_title.pack(anchor="w")
        self.members = tk.Listbox(rf, font=("Segoe UI", 10), activestyle="none", background=CARD, foreground=INK,
                                  selectbackground=ORANGE, selectforeground="white", borderwidth=0,
                                  highlightthickness=0)
        self.members.pack(fill="both", expand=True, pady=(4, 0))
        self.members.bind("<Double-Button-1>", self._oc_lookup_member)
        ttk.Label(rf, text="Double-click a member to see their email, job and manager (Org Chart)",
                  style="Muted.TLabel", font=(FONT, 8)).pack(anchor="w", pady=(4, 0))
        panes.add(rf, weight=2)

    def _show(self, view):
        for v in (self.text_view, self.docs_view, self.teams_view):
            v.pack_forget()
        view.pack(fill="both", expand=True)

    def _refresh_action_state(self):
        if hasattr(self, "ov_export_btn"):
            self.ov_export_btn.state(["!disabled"] if getattr(self, "ov_results", None) and not self.busy
                                     else ["disabled"])
        on = self.current_tab is not None
        for b in (self.b_shot, self.b_xls):
            b.state(["!disabled"] if on else ["disabled"])
        self.b_all.state(["!disabled"] if self.order else ["disabled"])

    # ----- folders / settings
    def _store_settings(self):
        self.settings.update(screenshots=self.shot_dir.get(), exports=self.export_dir.get(),
                             full_page=self.full_page.get(), auto_shot=self.auto_shot.get(),
                             per_order_folder=self.per_order.get(), parallel_sessions=self.parallel.get(),
                             ai_key_file=self.ai_key_file.get(), notify_on=self.notify_on.get(),
                             notify_orders=self.notify_orders.get(), notify_every=self.notify_every.get(),
                             notify_digest=self.notify_digest.get(), notify_updates=self.notify_updates.get(),
                             notify_tasks=self.notify_tasks.get())
        save_settings(self.settings)

    def _browse(self, var):
        d = filedialog.askdirectory(initialdir=var.get() or str(HERE), title="Choose folder")
        if d:
            var.set(d)
            self._store_settings()

    def _browse_key(self):
        start = Path(self.ai_key_file.get()).parent if self.ai_key_file.get() else Path.home()
        p = filedialog.askopenfilename(initialdir=str(start), title="Choose the AI key file",
                                       filetypes=[("Key file", "*.key"), ("All files", "*.*")])
        if p:
            self.ai_key_file.set(p)
            self.chat = None  # the next question uses the new key
            self._store_settings()

    def _open_folder(self, path):
        Path(path).mkdir(parents=True, exist_ok=True)
        os.startfile(path)

    def _ctx(self):
        """Snapshot of the window's settings, taken on the UI thread for a worker job."""
        self._store_settings()
        shots = Path(self.shot_dir.get())
        exports = Path(self.export_dir.get())
        if self.per_order.get() and self.order:
            shots = shots / self.order
        return {"order": self.order,
                "shots": shots, "full_page": self.full_page.get(),
                "auto_shot": self.auto_shot.get(),
                "workbook": exports / f"{self.order}.xlsx",
                "docs_dest": exports / (self.order or "") / "documents"}

    # ----- actions
    def open_order(self):
        order = self.order_var.get().strip()
        if not re.fullmatch(r"\d+-\d+", order):
            messagebox.showwarning(APP_NAME, "Type an order number like 10-1771012.")
            return
        if self.busy:
            return
        self._store_settings()
        w = self.worker

        def job():
            w.ensure_browser()
            self.log(f"Searching order {order}...")
            oo.open_orders_tab(w.driver, w.wait, discard_changes=True)
            oo.untick_my_orders(w.driver)
            row = oo.search_order(w.driver, w.wait, order)
            oo.open_result(w.driver, w.wait, row, order)
            banner = gx.extract_tab(w.driver)["banner"]
            return [t for t, _ in oo.list_tabs(w.driver)], banner

        def done(result):
            titles, banner = result
            self.order = order
            self.tab_titles = titles
            self.tab_list.delete(0, "end")
            for t in titles:
                self.tab_list.insert("end", "  " + t)
            self.order_label.configure(text=banner.replace("  |  ", "   ·   "))
            self.log(f"Opened {order}: {banner}")
            self.current_tab = None
            self._select_tab_by_name("Information")

        w.submit(f"Opening {order}", job, done, retry=True)

    def _ensure_order(self, order, tab=None):
        """Worker thread: healthy GOLD session with this order (and tab) open. After a
        reconnect the new session starts on GOLD's home page, so reopen the order."""
        w = self.worker
        w.ensure_browser()
        if not oo._on_order_page(w.driver, order):
            self.log(f"Reopening {order}...")
            oo.open_orders_tab(w.driver, w.wait, discard_changes=True)
            oo.untick_my_orders(w.driver)
            oo.open_result(w.driver, w.wait, oo.search_order(w.driver, w.wait, order), order)
            if tab:
                oo.open_tab(w.driver, w.wait, tab)

    def _select_tab_by_name(self, name):
        for i, t in enumerate(self.tab_titles):
            if t.startswith(name):
                self.tab_list.selection_clear(0, "end")
                self.tab_list.selection_set(i)
                self._tab_clicked()
                return

    def _tab_clicked(self):
        sel = self.tab_list.curselection()
        if not sel or not self.order:
            return
        title = self.tab_titles[sel[0]]
        ctx = self._ctx()
        if self.busy:
            return  # a job is running; ignore clicks
        w = self.worker
        order = self.order

        def job():
            self._ensure_order(order)
            name = oo.open_tab(w.driver, w.wait, title)
            if name == "Documents":
                payload = oo.read_documents(w.driver)
            elif name == "Teams":
                payload = oo.read_teams(w.driver)
            else:
                payload = gx.extract_tab(w.driver)
                if not payload["fields"] and not payload["grids"]:
                    payload["text"] = oo.content_text(w.driver)
            shot =self._take_screenshot(name, ctx) if ctx["auto_shot"] else None
            return name, payload, shot

        def done(result):
            name, payload, shot = result
            self.current_tab = name
            self.tab_title.configure(text=f"{name}   —   {self.order}")
            if name == "Documents":
                self._fill_docs(payload)
                self._show(self.docs_view)
            elif name == "Teams":
                self._fill_teams(payload)
                self._show(self.teams_view)
            else:
                self._fill_text(payload)
                self._show(self.text_view)
            if shot:
                self.log(f"Screenshot saved: {shot}")
            self._refresh_action_state()

        w.submit(f"Opening {title}", job, done, retry=True)

    def _take_screenshot(self, tab, ctx):
        """Runs on the worker thread. Saves <Tab>_<n>.png with the next free n."""
        folder = ctx["shots"]
        folder.mkdir(parents=True, exist_ok=True)
        base = safe_name(tab)
        n = 1
        while (folder / f"{base}_{n}.png").exists():
            n += 1
        path = folder / f"{base}_{n}.png"
        d = self.worker.driver
        if ctx["full_page"] and hasattr(d, "get_full_page_screenshot_as_file"):
            d.get_full_page_screenshot_as_file(str(path))
        else:
            d.save_screenshot(str(path))
        return path

    def screenshot(self):
        tab, ctx, order = self.current_tab, self._ctx(), self.order
        self.worker.submit("Taking screenshot",
                           lambda: (self._ensure_order(order, tab), self._take_screenshot(tab, ctx))[1],
                           lambda p: self.log(f"Screenshot saved: {p}"), retry=True)

    def _export_current(self, tab, ctx, docs):
        """Runs on the worker thread: extract the tab shown in GOLD and write its sheet."""
        d = self.worker.driver
        data = gx.extract_tab(d)
        header = gx.header_lines(ctx["order"], data["banner"])
        if tab == "Teams":
            teams = oo.read_teams(d)
            blocks = [("Teams", ["Team", "Primary member", "Members"],
                       [[t, v["primary"] or "", len(v["members"])] for t, v in teams.items()]),
                      ("Members", ["Team", "Member", "Primary"],
                       [[t, m, "Yes" if m == v["primary"] else ""]
                        for t, v in teams.items() for m in v["members"]])]
        elif tab == "Documents":
            saved = {doc["key"]: doc.get("saved", "") for doc in docs}
            blocks = [("Documents", ["#", "Document", "Last modified", "Type", "SRF relating",
                                     "Downloaded to"],
                       [[i, doc["title"], doc["modified"], doc["type"], doc["srf"],
                         saved.get(doc["key"], "")]
                        for i, doc in enumerate(oo.read_documents(d), 1)])]
        else:
            blocks = gx.tab_blocks(data)
        if not blocks:
            blocks = [("Content", ["Text"], [[line] for line in oo.content_text(d).splitlines()])]
        try:
            return gx.write_sheet(ctx["workbook"], tab, header, blocks)
        except PermissionError:
            raise RuntimeError(f"{ctx['workbook'].name} is open in Excel. Close it and try again.")

    def export_tab(self):
        tab, ctx, docs, order = self.current_tab, self._ctx(), list(self.docs), self.order
        self.worker.submit(f"Exporting {tab}",
                           lambda: (self._ensure_order(order, tab), self._export_current(tab, ctx, docs))[1],
                           lambda p: self._exported(p, [tab]), retry=True)

    def export_all(self):
        titles, ctx, docs, order = list(self.tab_titles), self._ctx(), list(self.docs), self.order
        w = self.worker

        def job():
            self._ensure_order(order)
            done_tabs = []
            for t in titles:
                name = oo.open_tab(w.driver, w.wait, t)
                self.log(f"  exporting {name}...")
                try:
                    self._export_current(name, ctx, docs)
                    done_tabs.append(name)
                except RuntimeError:
                    raise  # Excel file open: stop
                except Exception as e:
                    self.log(f"  {name}: skipped ({e})")
            if self.current_tab:  # go back to the tab the user was on
                oo.open_tab(w.driver, w.wait, self.current_tab)
            return ctx["workbook"], done_tabs

        w.submit("Exporting all tabs", job, lambda r: self._exported(*r), retry=True)

    def _exported(self, path, tabs):
        self.log(f"Excel updated: {path}  (sheets: {', '.join(tabs)})")
        if messagebox.askyesno(APP_NAME, f"Saved to\n{path}\n\nOpen it now?"):
            os.startfile(path)

    def download(self, all_docs=False):
        items = self.docs_tree.get_children() if all_docs else self.docs_tree.selection()
        if not items:
            messagebox.showinfo(APP_NAME, "Select one or more documents first.")
            return
        picks = [self.docs[int(i)] for i in items]
        dest = self._ctx()["docs_dest"]
        w = self.worker
        order = self.order

        def job():
            results = []
            self._ensure_order(order, "Documents")
            for doc in picks:
                self.log(f"Downloading {doc['title']}...")
                try:
                    saved, _ = oo.download_document(w.driver, w.wait, doc, dest)
                except Exception as e:
                    saved = []
                    self.log(f"  failed: {str(e).splitlines()[0] if str(e) else type(e).__name__}")
                if saved:
                    doc["saved"] = saved[0]
                    self.log(f"  saved {saved[0]}")
                results.append(doc)
                self._ensure_order(order, "Documents")  # reconnect if the connection dropped
                oo.open_tab(w.driver, w.wait, "Documents")
            return results

        def done(_):
            self._fill_docs(self.docs, keep_saved=True)
            ok = sum(1 for d in picks if d.get("saved"))
            self.log(f"{ok} of {len(picks)} documents downloaded to {dest}")

        w.submit(f"Downloading {len(picks)} document(s)", job, done)

    # ----- fill views
    def _fill_text(self, data):
        """Show a GOLD tab: its fields as label / value pairs by section, then each of its tables as a real
        table (headers, zebra rows, copy cells / rows / the whole table)."""
        t = self.text
        t.configure(state="normal")
        t.delete("1.0", "end")
        for w in t.winfo_children():
            w.destroy()
        self._grid_fits = []
        if data.get("text"):
            t.insert("end", data["text"])
        section = None
        for sec, field, value in data["fields"]:
            if sec != section:
                section = sec
                if sec:
                    t.insert("end", ("\n" if t.index("end-1c") != "1.0" else "") + sec + "\n", "section")
            if field:
                t.insert("end", f"  {field}", "label")
                t.insert("end", f"\t{value}\n", "value")
            else:
                t.insert("end", f"  {value}\n", "value")
        for g in data["grids"]:
            t.insert("end", "\n" if t.index("end-1c") != "1.0" else "")
            t.window_create("end", window=self._grid_table(t, g))
            t.insert("end", "\n")
        t.configure(state="disabled")

    def _grid_table(self, parent, g):
        """One GOLD table as a neat, copyable table."""
        scale = self.winfo_fpixels("1i") / 96
        box = tk.Frame(parent, background=CARD, highlightthickness=1, highlightbackground=LINE, padx=10, pady=8)
        head = tk.Frame(box, background=CARD)
        head.pack(fill="x", pady=(0, 6))
        tk.Label(head, text=g.get("title") or "Table", bg=CARD, fg=INK, font=(FONT, 10, "bold")).pack(side="left")
        tk.Label(head, text=f"  {len(g['rows'])} rows", bg=CARD, fg=MUTED, font=(FONT, 9)).pack(side="left")
        cols = [c or f"Column {i + 1}" for i, c in enumerate(g["columns"])]
        ids = [f"c{i}" for i in range(len(cols))]
        tf = tk.Frame(box, background=CARD)
        tf.pack(fill="both", expand=True)
        tree = ttk.Treeview(tf, columns=ids, show="headings", selectmode="browse",
                            height=max(1, min(len(g["rows"]), 14)))
        copy = gui.TableCopy(tree, self.log)
        ttk.Button(head, text="Copy table", style="Small.TButton", command=copy.copy_table).pack(side="right")
        from tkinter import font as tkfont
        f_body, f_head = tkfont.Font(family=FONT, size=10), tkfont.Font(family=FONT, size=9, weight="bold")
        for i, (cid, title) in enumerate(zip(ids, cols)):
            longest = max([f_head.measure(title)] + [f_body.measure(str(r[i])) for r in g["rows"] if i < len(r)])
            tree.heading(cid, text=title, anchor="w")
            tree.column(cid, width=min(int(360 * scale), longest + int(24 * scale)), minwidth=int(50 * scale),
                        anchor="w", stretch=True)
        # milestone-like tables: done in green, open and past the target date in red
        target = next((i for i, c in enumerate(cols) if c.lower().startswith("target date")), None)
        completed = next((i for i, c in enumerate(cols) if c.lower().startswith("completed")), None)
        today = datetime.now().date()
        tree.tag_configure("odd", background=gui.ZEBRA)
        tree.tag_configure("done", foreground=gui.GREEN_FG)
        tree.tag_configure("late", foreground=gui.RED_FG)
        for n, r in enumerate(g["rows"]):
            tags = ["odd"] if n % 2 else []
            if completed is not None and completed < len(r) and str(r[completed]).strip():
                tags.append("done")
            elif target is not None and target < len(r):
                try:
                    if gm.parse_date(str(r[target])) < today:
                        tags.append("late")
                except (ValueError, TypeError):
                    pass
            tree.insert("", "end", values=[str(v) for v in r] + [""] * (len(cols) - len(r)), tags=tags)
        xs = ttk.Scrollbar(tf, orient="horizontal", command=tree.xview)
        tree.configure(xscrollcommand=xs.set)
        if len(g["rows"]) > 14:
            ys = ttk.Scrollbar(tf, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=ys.set)
            ys.pack(side="right", fill="y")
        tree.pack(fill="both", expand=True)
        natural = [tree.column(c, "width") for c in ids]
        least = int(80 * scale)

        def fit(e=None):  # the table takes the page's width; columns share it; scroll sideways only if needed
            if not box.winfo_exists():
                return
            w = max(self.text.winfo_width() - int(40 * scale), int(300 * scale))
            room = w - int(30 * scale)
            widths = natural if sum(natural) <= room else [max(least, int(n * room / sum(natural))) for n in natural]
            for c, cw in zip(ids, widths):
                tree.column(c, width=cw)
            if sum(widths) > room:
                xs.pack(side="bottom", fill="x", before=tree)
            else:
                xs.pack_forget()
            box.update_idletasks()
            h = head.winfo_reqheight() + tf.winfo_reqheight() + 2 * int(box.cget("pady")) + 8
            box.configure(width=w, height=h)
        box.pack_propagate(False)
        self._grid_fits.append(fit)
        fit()
        return box

    def _fill_docs(self, docs, keep_saved=False):
        if not keep_saved:
            old = {d["key"]: d.get("saved") for d in self.docs}
            for d in docs:
                if old.get(d["key"]):
                    d["saved"] = old[d["key"]]
        self.docs = docs
        self.docs_tree.delete(*self.docs_tree.get_children())
        for i, d in enumerate(docs):
            saved = Path(d["saved"]).name if d.get("saved") else ""
            self.docs_tree.insert("", "end", iid=str(i),
                                  values=(i + 1, d["title"], d["modified"], d["type"], saved))

    def _fill_teams(self, teams):
        self.teams = teams
        self.teams_tree.delete(*self.teams_tree.get_children())
        for name, t in teams.items():
            self.teams_tree.insert("", "end", iid=name,
                                   values=(name, t["primary"] or "—", len(t["members"])))
        self.members.delete(0, "end")
        self.members_title.configure(text="Members")
        self._find_person()

    def _team_selected(self):
        sel = self.teams_tree.selection()
        if not sel:
            return
        t = self.teams[sel[0]]
        self.members_title.configure(text=f"{sel[0]} — {len(t['members'])} members")
        self.members.delete(0, "end")
        for m in t["members"]:
            self.members.insert("end", ("★ " if m == t["primary"] else "    ") + m)

    def _find_person(self):
        q = self.find_var.get().strip().lower()
        if not q:
            self.find_result.configure(text="")
            return
        hits = [(name, m) for name, t in self.teams.items() for m in t["members"] if q in m.lower()]
        teams_hit = sorted({name for name, _ in hits})
        people = sorted({m for _, m in hits})
        self.find_result.configure(
            text=(f"{', '.join(people[:3])}{'…' if len(people) > 3 else ''} → in "
                  f"{len(teams_hit)} team(s): {', '.join(teams_hit)}") if hits else "No match")
        self.teams_tree.selection_remove(self.teams_tree.selection())
        for name in teams_hit:
            self.teams_tree.selection_add(name)

    def _on_close(self):
        self._store_settings()
        self.worker.close()
        self.destroy()


def selftest(order):
    """GOLD Express.exe --selftest 10-1771012 : open Firefox + GOLD, open the order,
    read its tabs, export one sheet, and write a report. For troubleshooting."""
    report = HERE / "selftest_report.txt"
    lines = []
    d = None
    try:
        d = start_firefox(headless=True)
        lines.append("Firefox started (background)")
        w = Wait(d, 60)
        d.get(GOLD_URL)
        if not wait_for_gold(d):
            raise RuntimeError("GOLD did not load (SSO?)")
        lines.append("GOLD loaded")
        oo.open_orders_tab(d, w)
        oo.untick_my_orders(d)
        oo.open_result(d, w, oo.search_order(d, w, order), order)
        tabs = [t for t, _ in oo.list_tabs(d)]
        lines.append(f"Order {order} opened, tabs: {', '.join(tabs)}")
        oo.open_tab(d, w, "Milestones")
        data = gx.extract_tab(d)
        path = gx.write_sheet(HERE / "output" / f"selftest_{order}.xlsx", "Milestones",
                              gx.header_lines(order, data["banner"]), gx.tab_blocks(data))
        lines.append(f"Excel written: {path}")
        lines.append("RESULT: OK")
    except Exception:
        lines.append("RESULT: FAILED\n" + traceback.format_exc())
    finally:
        if d is not None:
            d.quit()
        report.write_text("\n".join(lines), encoding="utf-8")


def ask_selftest(sharepoint=False):
    """GOLD Express.exe --ask-selftest [--sharepoint] : check the AI key and the AI service (no GOLD).
    --sharepoint does what a colleague's first question does: a brand-new hidden Firefox gets the key
    from the team SharePoint folder (the synced folder and the saved key are not used)."""
    import shutil
    import tempfile
    lines = []
    d = prof = None
    try:
        if sharepoint:
            prof = tempfile.mkdtemp(prefix="gx-ask-")
            d = start_firefox(headless=True, profile_dir=prof)
            key = gai.fetch_key_from_sharepoint(d)
            lines.append("Key: fetched from the team SharePoint folder with a new Firefox profile")
        else:
            def driver():
                nonlocal d
                d = start_firefox(headless=True, profile_dir=tempfile.mkdtemp(prefix="gx-ask-"))
                return d
            key = gai.get_key(load_settings().get("ai_key_file", ""), AI_KEY_CACHE, driver, lines.append)
            lines.append("Key: found")
        chat = gai.Chat(key, reader=None)
        answer = chat.ask("This is a connection test. Reply with just: OK")
        lines.append(f"Answer: {answer}")
        lines.append("RESULT: OK")
    except Exception as e:
        lines.append(f"RESULT: FAILED - {gai.friendly_error(e)}\n" + traceback.format_exc())
    finally:
        if d is not None:
            d.quit()
        if prof:
            shutil.rmtree(prof, ignore_errors=True)
    (HERE / "ask_selftest_report.txt").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    import sys
    HERE.mkdir(parents=True, exist_ok=True)
    if len(sys.argv) >= 3 and sys.argv[1] == "--selftest":
        selftest(sys.argv[2])
    elif len(sys.argv) >= 2 and sys.argv[1] == "--ui-selftest":
        # Open the window without GOLD, show every page, save a screenshot, close (for troubleshooting).
        App._prewarm = lambda self: None
        App._check_updates = lambda self, quiet=True: None
        App._my_load = lambda self, refresh=False: None  # showing My Orders must not start GOLD here
        App._oa_shown = lambda self: None  # nor Order Assignment fetch its sheet
        report = HERE / "ui_selftest_report.txt"
        try:
            app = App()

            def tour(i=0):
                if i < len(app.notebook.pages):
                    app.notebook.select(i)
                    app.after(300, tour, i + 1)
                    return
                app.notebook.select(0)
                app.update()
                try:
                    from PIL import ImageGrab
                    x, y = app.winfo_rootx(), app.winfo_rooty()
                    ImageGrab.grab((x, y, x + app.winfo_width(), y + app.winfo_height())).save(
                        HERE / "ui_selftest.png")
                except Exception:
                    pass
                report.write_text(f"RESULT: OK - {len(app.notebook.pages)} pages shown", encoding="utf-8")
                app._on_close()
            app.after(800, tour)
            app.mainloop()
        except Exception:
            report.write_text("RESULT: FAILED\n" + traceback.format_exc(), encoding="utf-8")
    elif len(sys.argv) >= 2 and sys.argv[1] == "--ask-selftest":
        ask_selftest(sharepoint="--sharepoint" in sys.argv)
    else:
        # One GOLD Express at a time: two copies can't share the Firefox profile.
        import ctypes
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        if "--after-pid" in sys.argv:  # a restart (e.g. light / dark mode): wait for the old window to close
            try:
                pid = int(sys.argv[sys.argv.index("--after-pid") + 1])
                handle = kernel32.OpenProcess(0x00100000, False, pid)  # SYNCHRONIZE
                if handle:
                    kernel32.WaitForSingleObject(handle, 30_000)
                    kernel32.CloseHandle(handle)
            except (ValueError, IndexError):
                pass
        mutex = kernel32.CreateMutexW(None, False, "Local\\GOLDExpressSingleInstance")
        if ctypes.get_last_error() == 183:  # ERROR_ALREADY_EXISTS
            r = tk.Tk(); r.withdraw()
            messagebox.showinfo(APP_NAME, f"{APP_NAME} is already open.\n\n"
                                          "Look for it on the taskbar.")
            sys.exit(0)
        App().mainloop()

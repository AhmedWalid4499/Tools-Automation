"""Turn the current GOLD tab into structured data and write neat Excel workbooks.

extract_tab(driver) -> {"grids": [ {"columns": [...], "rows": [[...], ...]} ],
                        "fields": [ (section, field, value), ... ]}
  grids  = list-style tables with a header row (Milestones, Notes, Contacts, Pricing...)
  fields = label/value pairs from form-style pages (Information, Billing, Service build...)
"""
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Rows of every visible, innermost table of the content area (not the menus).
# Dropdowns contribute their selected text, checkboxes "Yes"/"No".
LEAF_TABLES_JS = r"""
const skip = el => el.closest('#LeftNavbarDiv') || el.closest('#headMenu');
const visible = el => !!(el.offsetWidth || el.offsetHeight || el.getClientRects().length);
function cellText(td) {
  const clone = td.cloneNode(true);
  const origSel = td.querySelectorAll('select'), cloneSel = clone.querySelectorAll('select');
  cloneSel.forEach((s, i) => { const o = origSel[i];
    s.replaceWith(document.createTextNode(o.selectedIndex >= 0 ? o.options[o.selectedIndex].text : '')); });
  const origIn = td.querySelectorAll('input,textarea'), cloneIn = clone.querySelectorAll('input,textarea');
  cloneIn.forEach((e, i) => { const o = origIn[i]; let v = '';
    if (o.type === 'checkbox' || o.type === 'radio') v = o.checked ? 'Yes' : 'No';
    else if (o.type !== 'hidden' && o.type !== 'button' && o.type !== 'image') v = o.value;
    e.replaceWith(document.createTextNode(v)); });
  clone.querySelectorAll('script,style').forEach(e => e.remove());
  return clone.textContent.replace(/ /g, ' ').replace(/\s+/g, ' ').trim();
}
// Every table, keeping only its own rows that don't wrap another table
// (so a grid whose rows embed detail tables is still read as a grid).
const out = [];
for (const t of document.querySelectorAll('table')) {
  if (skip(t) || !visible(t)) continue;
  const rows = [];
  for (const tr of t.rows) {
    if (!visible(tr) || tr.querySelector('table')) continue;
    const cells = [...tr.cells].map(cellText);
    if (cells.some(c => c)) rows.push(cells);
  }
  if (rows.length) out.push(rows);
}
return out;
"""

HEADER_SKIP = re.compile(r"^(Previous|Next|\d+)$")
NOISE = re.compile(r"\s*Mobile:\s*$")  # contacts show an empty "Mobile:" label


def _compact(row):
    return [c for c in row if c]


def _existing_vs_new(rows):
    """Service build layout: a header with 'Existing X' and 'New X', then per item
    a label and its existing / new values. Returns a grid or None."""
    head = rows[0]
    if len(_compact(head)) != 2:  # just "Existing X" and "New X" (not a charges grid)
        return None
    ex = next((i for i, c in enumerate(head) if c.startswith("Existing ")), None)
    new = next((i for i, c in enumerate(head) if c.startswith("New ")), None)
    if ex is None or new is None:
        return None
    split = (ex + new) / 2 + 1  # values sit just right of their header cell
    out, section = [], ""
    for row in rows[1:]:
        cells = [(i, c) for i, c in enumerate(row) if c]
        if len(cells) == 1:
            section = cells[0][1]
            continue
        label_i, label = cells[0]
        existing = " ".join(c for i, c in cells[1:] if i <= split)
        new_val = " ".join(c for i, c in cells[1:] if i > split)
        out.append([section, label.lstrip("* "), existing, new_val])
    return {"title": f"{head[ex]}  ->  {head[new]}",
            "columns": ["Section", "Item", "Existing", "New"], "rows": out}


def _header_index(rows):
    """Index of a header row (3+ short labels) followed by rows of about the same width."""
    for h, header in enumerate(rows[:3]):
        width = len(header)
        if len(_compact(header)) < 3 or any(len(c) > 40 for c in header):
            continue
        if sum(1 for c in header if c.endswith(":")) >= 2:
            continue  # "Label: value Label: value" form rows, not a list header
        body = rows[h + 1:]
        same = sum(1 for r in body if abs(len(r) - width) <= 1)
        if body and same >= max(1, len(body) * 0.5):
            return h
    return None


def _grid(rows, h):
    """Header = rows[h]; earlier rows form the title. Rows of a different width are
    detail lines of the row above (added as a 'Details' column). Empty columns dropped."""
    title = " ".join(" ".join(_compact(r)) for r in rows[:h])
    header = rows[h]
    width = len(header)
    body, details = [], []
    for r in rows[h + 1:]:
        if abs(len(r) - width) <= 1:
            body.append((r + [""] * width)[:width])
            details.append([])
        elif body:
            details[-1].append(" ".join(_compact(r)))
    body = [[NOISE.sub("", c) for c in r] for r in body]
    keep = [i for i in range(width) if header[i] or any(r[i] for r in body)]
    cols = [header[i] or f"Column {n + 1}" for n, i in enumerate(keep)]
    rows_out = [[r[i] for i in keep] for r in body]
    if any(details):
        cols.append("Details")
        rows_out = [r + [" | ".join(d)] for r, d in zip(rows_out, details)]
    return {"title": title, "columns": cols, "rows": [r for r in rows_out if any(r)]}


def _fields(rows, section):
    """Pair labels with values: 'Label:' followed by its value, label/value cells, or a
    row of headings followed by a (differently laid out) row of their values."""
    pairs = []
    if len(rows) == 2 and len(_compact(rows[0])) == len(_compact(rows[1])) == 1:
        return [(section, _compact(rows[0])[0].lstrip("* ").rstrip(":"), _compact(rows[1])[0])], section
    skip_next = False
    for n, row in enumerate(rows):
        if skip_next:
            skip_next = False
            continue
        cells = _compact(row)
        # A lone label row followed by a lone value row ("Quote Owner" / "swapnil b.").
        nxt = rows[n + 1] if n + 1 < len(rows) else None
        if (len(cells) == 1 and nxt and len(_compact(nxt)) == 1 and len(cells[0]) < 40
                and not cells[0].endswith(":") and (cells[0][0].isupper() or cells[0][0] == "*")
                and n + 2 >= len(rows)):
            pairs.append((section, cells[0].lstrip("* "), _compact(nxt)[0]))
            skip_next = True
            continue
        nxt = rows[n + 1] if n + 1 < len(rows) else None
        if (nxt and len(row) != len(nxt) and len(cells) >= 2
                and len(_compact(nxt)) == len(cells) and all(len(c) <= 40 for c in cells)):
            for label, value in zip(cells, _compact(nxt)):
                pairs.append((section, label.lstrip("* ").rstrip(":"), value))
            skip_next = True
            continue
        if len(cells) == 1:
            c = cells[0]
            if c.endswith(":"):
                pairs.append((section, c.rstrip(": *").strip(), ""))
            elif len(c) < 60:
                section = c.lstrip("* ").strip()  # a lone short cell is usually a heading
            else:
                pairs.append((section, "", c))
            continue
        i = 0
        while i < len(cells):
            label = cells[i]
            value = cells[i + 1] if i + 1 < len(cells) else ""
            if value.endswith(":") and not label.endswith(":"):
                pairs.append((section, "", label))  # orphan value before next label
                i += 1
                continue
            pairs.append((section, label.rstrip(":").lstrip("* ").strip(), value))
            i += 2
    return pairs, section


YES_NO_ITEM = re.compile(r"^(Yes|No) (.+)$")


def extract_tab(driver, tab=""):
    tables = driver.execute_script(LEAF_TABLES_JS)
    # Everything up to the order banner ("Customer: X | Site: Y | Order Status: Z") is
    # GOLD's top bar; the banner itself goes into the sheet header.
    banner = ""
    for i, rows in enumerate(tables):
        if rows[0] and rows[0][0].startswith("Customer:"):
            banner = "  |  ".join(" ".join(_compact(r)) for r in rows)
            tables = tables[i + 1:]
            break

    grids, fields, section, note_blocks = [], [], "", []
    for rows in tables:
        rows = [r for r in rows if not all(HEADER_SKIP.match(c) or not c for c in r)]
        if not rows:
            continue
        if _compact(rows[0]) == ["Note:"]:  # Notes tab: full text of one note
            note_blocks.append([" ".join(_compact(r)) for r in rows[1:]])
            continue
        evn = _existing_vs_new(rows)
        if evn:
            grids.append(evn)
            continue
        h = _header_index(rows)
        if h is not None:
            g = _grid(rows, h)
            if g["rows"]:
                grids.append(g)
        else:
            new, section = _fields(rows, section)
            fields += new

    if note_blocks and grids:  # attach each note's full text to its row, in order
        g = max(grids, key=lambda g: len(g["rows"]))
        if len(g["rows"]) == len(note_blocks):
            g["columns"] += ["Full note", "Visibility / RFB status"]
            for r, nb in zip(g["rows"], note_blocks):
                r += [nb[0] if nb else "", " ".join(nb[1:])]

    # Checkbox lists (Companion services): "Yes Service name" -> rows of (item, selected).
    items = [YES_NO_ITEM.match(v) for f in fields for v in f[1:] if v]
    if fields and sum(1 for m in items if m) >= len(items) * 0.8:
        rows = [[m.group(2), m.group(1)] for m in items if m]
        grids.insert(0, {"title": fields[0][0], "columns": ["Item", "Selected"], "rows": rows})
        fields = [f for f in fields if not any(v and YES_NO_ITEM.match(v) for v in f[1:])]

    return {"banner": banner, "grids": grids, "fields": fields}


# ---------------- Excel ----------------

ORANGE = "FF7900"
HEADER_FILL = PatternFill("solid", fgColor=ORANGE)
TITLE_FONT = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet_name(wb, name):
    name = re.sub(r"[\[\]:*?/\\]", "", name)[:31] or "Sheet"
    if name in wb.sheetnames:
        del wb[name]  # re-export replaces the old sheet
    return name


def _table_name(wb, base):
    base = re.sub(r"\W", "_", base) or "T"
    existing = {t for ws in wb.worksheets for t in ws.tables}
    n = 1
    while f"{base}_{n}" in existing:
        n += 1
    return f"{base}_{n}"


def _write_block(wb, ws, top, title, columns, rows, table_base):
    """Title line, orange header, rows, as a filterable Excel table. Returns next free row."""
    ws.cell(row=top, column=1, value=title).font = Font(bold=True, size=12)
    hdr = top + 1
    for c, name in enumerate(columns, 1):
        cell = ws.cell(row=hdr, column=c, value=name)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, hdr + 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=_typed(val))
            cell.alignment = WRAP
            cell.border = Border(bottom=THIN)
    last = hdr + max(len(rows), 1)
    if rows:
        ref = f"A{hdr}:{get_column_letter(len(columns))}{last}"
        # Excel tables need unique header names.
        seen = {}
        for c in range(1, len(columns) + 1):
            v = str(ws.cell(row=hdr, column=c).value)
            seen[v] = seen.get(v, 0) + 1
            if seen[v] > 1:
                ws.cell(row=hdr, column=c, value=f"{v} ({seen[v]})")
        t = Table(displayName=_table_name(wb, table_base), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
    return last + 3


NUM = re.compile(r"^-?[\d,]+\.\d+$|^-?\d+$")


def _typed(val):
    """Numbers as numbers (so Excel can sum them); keep IDs with leading zeros as text."""
    if isinstance(val, str) and NUM.match(val) and not (val.startswith("0") and len(val) > 1
                                                         and "." not in val):
        try:
            return float(val.replace(",", "")) if "." in val else int(val.replace(",", ""))
        except ValueError:
            pass
    return val


def _autosize(ws, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                longest = max(len(line) for line in str(cell.value).splitlines() or [""])
                widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), longest)
    for col, w in widths.items():
        ws.column_dimensions[col].width = min(max(w + 2, 8), max_width)


def open_workbook(path):
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def write_sheet(path, sheet, header_lines, blocks):
    """Add/replace one sheet. blocks = [(title, columns, rows), ...]."""
    wb = open_workbook(path)
    ws = wb.create_sheet(_sheet_name(wb, sheet))
    ws.cell(row=1, column=1, value=sheet).font = TITLE_FONT
    for i, line in enumerate(header_lines, 2):
        ws.cell(row=i, column=1, value=line).font = Font(italic=True, color="595959")
    row = len(header_lines) + 3
    for title, columns, rows in blocks:
        row = _write_block(wb, ws, row, title, columns, rows, sheet)
    _autosize(ws)
    ws.sheet_view.showGridLines = False
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def tab_blocks(data):
    blocks = []
    if data["fields"]:
        blocks.append(("Details", ["Section", "Field", "Value"],
                       [list(f) for f in data["fields"]]))
    for i, g in enumerate(data["grids"], 1):
        title = g.get("title") or ("Table" if len(data["grids"]) == 1 else f"Table {i}")
        blocks.append((title, g["columns"], g["rows"]))
    return blocks


def header_lines(order, banner):
    return [f"Order: {order}    {banner}".strip(),
            f"Exported {datetime.now():%d-%b-%Y %H:%M}"]

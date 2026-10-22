"""An Excel-like editable table for tkinter: type in cells, Tab/Enter to move, Ctrl+V pastes
rows copied from Excel, Delete clears rows, load .xlsx/.csv, and a Status column the app fills.

columns: [{"key", "title", "width", "stretch"?, "choices"?: callable -> [str],
           "normalize"?: callable(str) -> str}]   (a Status column is added automatically)
"""
import csv
import re
import tkinter as tk
from pathlib import Path
from tkinter import ttk

ORDER_RE = re.compile(r"\s*\d{2}-\d+")


class EditGrid(tk.Frame):
    def __init__(self, parent, columns, on_change=None, start_rows=15, **kw):
        super().__init__(parent, background="white", highlightthickness=0, **kw)  # sits in a rounded tile
        self.cols = columns
        self.keys = [c["key"] for c in columns]
        self.all_keys = ["n"] + self.keys + ["status"]
        self.on_change = on_change
        self._editor = None
        ttk.Style(self).configure("Grid.Treeview", rowheight=28)
        self.tree = ttk.Treeview(self, columns=self.all_keys, show="headings",
                                 selectmode="extended", style="Grid.Treeview")
        self.tree.heading("n", text="#", anchor="w")
        self.tree.column("n", width=40, anchor="w", stretch=False)
        for c in columns:
            self.tree.heading(c["key"], text=c["title"], anchor="w")
            self.tree.column(c["key"], width=c.get("width", 140), anchor="w",
                             stretch=c.get("stretch", False))
        self.tree.heading("status", text="Status", anchor="w")
        self.tree.column("status", width=360, anchor="w", stretch=True)
        self.tree.tag_configure("odd", background="#FAFAFB")
        self.tree.tag_configure("ok", foreground="#1E7B34")
        self.tree.tag_configure("bad", foreground="#C62828")
        self.tree.tag_configure("run", foreground="#FF7900")
        sb = ttk.Scrollbar(self, command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<ButtonRelease-1>", self._click)
        self.tree.bind("<Return>", lambda e: self._edit_selected())
        self.tree.bind("<Control-v>", lambda e: (self.paste(), "break")[1])
        self.tree.bind("<Control-V>", lambda e: (self.paste(), "break")[1])
        self.tree.bind("<Delete>", lambda e: self.delete_selected())
        self.add_rows(start_rows)

    # ----- rows / values
    def rows(self):
        return list(self.tree.get_children())

    def get(self, iid, key):
        return self.tree.item(iid, "values")[self.all_keys.index(key)]

    def record(self, iid):
        vals = self.tree.item(iid, "values")
        return {k: str(vals[self.all_keys.index(k)]).strip() for k in self.keys}

    def filled(self):
        """[(iid, {key: value})] for rows with anything typed in them."""
        return [(i, r) for i in self.rows() if any((r := self.record(i)).values())]

    def _blank(self):
        return [""] * len(self.all_keys)

    def _renumber(self):
        for n, iid in enumerate(self.rows(), 1):
            vals = list(self.tree.item(iid, "values"))
            vals[0] = n
            tags = [t for t in self.tree.item(iid, "tags") if t != "odd"]
            if n % 2 == 0:
                tags.append("odd")
            self.tree.item(iid, values=vals, tags=tags)

    def add_rows(self, n):
        for _ in range(n):
            self.tree.insert("", "end", values=self._blank())
        self._renumber()

    def clear(self, start_rows=15):
        self.close_editor(False)
        self.tree.delete(*self.rows())
        self.add_rows(start_rows)

    def delete_selected(self):
        for iid in self.tree.selection():
            self.tree.item(iid, values=self._blank(), tags=())
        self._renumber()

    def set_value(self, iid, key, value):
        vals = list(self.tree.item(iid, "values"))
        vals[self.all_keys.index(key)] = value
        vals[-1] = ""  # edited: the old status no longer applies
        tags = [t for t in self.tree.item(iid, "tags") if t == "odd"]
        self.tree.item(iid, values=vals, tags=tags)

    def set_status(self, iid, text, tag=None):
        vals = list(self.tree.item(iid, "values"))
        vals[-1] = text
        tags = [t for t in self.tree.item(iid, "tags") if t == "odd"] + ([tag] if tag else [])
        self.tree.item(iid, values=vals, tags=tags)
        self.tree.see(iid)

    def fill(self, rows, start_iid=None, start_key=None):
        """Write rows (lists in column order) into the grid from start_iid / start_key down."""
        ids = self.rows()
        i = ids.index(start_iid) if start_iid in ids else 0
        offset = self.keys.index(start_key) if start_key in self.keys else 0
        for r in rows:
            while i >= len(ids):
                self.add_rows(10)
                ids = self.rows()
            vals = list(self.tree.item(ids[i], "values"))
            for k, v in zip(self.keys[offset:], list(r)):
                col = next(c for c in self.cols if c["key"] == k)
                v = "" if v is None else str(v).strip() if not hasattr(v, "year") else v
                norm = col.get("normalize")
                if norm and v != "":
                    v = norm(v)
                vals[self.all_keys.index(k)] = v
            vals[-1] = ""
            self.tree.item(ids[i], values=vals, tags=())
            i += 1
        if i >= len(ids) - 2:
            self.add_rows(10)
        self._renumber()
        self._changed()

    def _changed(self):
        if self.on_change:
            self.on_change()

    # ----- paste / load
    def paste(self, text=None, start_iid=None, start_key=None):
        try:
            text = text if text is not None else self.clipboard_get()
        except tk.TclError:
            return 0
        rows = [line.split("\t") for line in text.replace("\r", "").split("\n") if line.strip()]
        if rows and (start_key in (None, self.keys[0])) and not ORDER_RE.match(rows[0][0]):
            rows = rows[1:]  # a header row copied along
        if start_iid is None:
            sel = self.tree.selection()
            start_iid = sel[0] if sel else next(
                (i for i in self.rows() if not any(self.record(i).values())), None)
        self.fill(rows, start_iid, start_key)
        return len(rows)

    def load(self, path):
        """Read the first len(columns) columns of an .xlsx/.xlsm/.csv/.txt; rows must start
        with an order number. Returns the number of rows loaded."""
        width = len(self.keys)
        rows = []
        if str(path).lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
            for row in ws.iter_rows(values_only=True):
                if row and row[0] is not None and ORDER_RE.match(str(row[0])):
                    rows.append(list(row[:width]))
        else:
            text = Path(path).read_text(encoding="utf-8-sig")
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
            for row in csv.reader(text.splitlines(), dialect):
                if row and ORDER_RE.match(row[0]):
                    rows.append(row[:width])
        self.tree.delete(*self.rows())
        self.add_rows(max(15, len(rows) + 5))
        self.fill(rows, self.rows()[0])
        return len(rows)

    # ----- cell editing
    def _click(self, event):
        if self.tree.identify("region", event.x, event.y) != "cell":
            return
        iid = self.tree.identify_row(event.y)
        key = self.all_keys[int(self.tree.identify_column(event.x)[1:]) - 1]
        if iid and key in self.keys and len(self.tree.selection()) <= 1:
            self.edit(iid, key)

    def _edit_selected(self):
        sel = self.tree.selection()
        if sel:
            self.edit(sel[0], self.keys[0])

    def edit(self, iid, key):
        self.close_editor(True)
        self.tree.see(iid)
        self.update_idletasks()
        box = self.tree.bbox(iid, key)
        if not box:
            return
        x, y, w, h = box
        col = next(c for c in self.cols if c["key"] == key)
        var = tk.StringVar(value=self.get(iid, key))
        if col.get("choices") and col.get("typeahead"):
            # Type to filter a long list (e.g. names); the value snaps to the list's spelling.
            choices = col["choices"]()
            ed = ttk.Combobox(self.tree, textvariable=var, state="normal", values=choices)

            def filter_list(event):
                if event.keysym in ("Return", "Tab", "Escape", "Up", "Down"):
                    return
                typed = var.get().strip().lower()
                ed.configure(values=[c for c in choices if typed in c.lower()] if typed else choices)
            ed.bind("<KeyRelease>", filter_list)
            ed.bind("<<ComboboxSelected>>", lambda e: self.close_editor(True))
            col.setdefault("normalize", lambda v, cs=choices: next(
                (c for c in col["choices"]() if c.lower() == " ".join(v.split()).lower()), v))
        elif col.get("choices"):
            ed = ttk.Combobox(self.tree, textvariable=var, state="readonly", values=[""] + col["choices"]())
            ed.bind("<<ComboboxSelected>>", lambda e: self.close_editor(True))
            self.after(50, lambda: ed.winfo_exists() and ed.event_generate("<Button-1>"))
        else:
            ed = ttk.Entry(self.tree, textvariable=var)
            ed.select_range(0, "end")
        if not col.get("choices") or col.get("typeahead"):
            ed.bind("<Control-v>", self._entry_paste)
        if not col.get("choices"):
            ed.bind("<Down>", lambda e: self._move(1, 0))
            ed.bind("<Up>", lambda e: self._move(-1, 0))
        ed.place(x=x, y=y, width=w, height=h)
        ed.focus_set()
        self._editor = (ed, iid, key, var)
        ed.bind("<Return>", lambda e: self._move(1, 0))
        ed.bind("<Tab>", lambda e: (self._move(0, 1), "break")[1])
        ed.bind("<Shift-Tab>", lambda e: (self._move(0, -1), "break")[1])
        ed.bind("<Escape>", lambda e: self.close_editor(False))
        ed.bind("<FocusOut>", lambda e: self.after(50, self._focus_out, ed))

    def _focus_out(self, ed):
        if self._editor and self._editor[0] is ed:
            focus = self.focus_get()
            if focus is ed or (focus is not None and str(focus).startswith(str(ed))):
                return
            self.close_editor(True)

    def _entry_paste(self, event):
        try:
            clip = self.clipboard_get()
        except tk.TclError:
            return None
        if "\t" in clip or clip.strip().count("\n"):  # several cells from Excel
            _, iid, key, _ = self._editor
            self.close_editor(False)
            self.paste(clip, iid, key)
            return "break"
        return None

    def close_editor(self, commit=True):
        if not self._editor:
            return
        ed, iid, key, var = self._editor
        self._editor = None
        if commit and self.tree.exists(iid):
            value = var.get().strip()
            col = next(c for c in self.cols if c["key"] == key)
            if value and col.get("normalize"):
                value = col["normalize"](value)
            if value != self.get(iid, key):
                self.set_value(iid, key, value)
                self._changed()
        ed.destroy()

    def _move(self, drow, dcol):
        if not self._editor:
            return "break"
        _, iid, key, _ = self._editor
        self.close_editor(True)
        ids = self.rows()
        ci, ri = self.keys.index(key) + dcol, ids.index(iid)
        if ci >= len(self.keys):
            ci, ri = 0, ri + 1
        elif ci < 0:
            ci, ri = len(self.keys) - 1, ri - 1
        ri += drow
        if ri >= len(ids):
            self.add_rows(10)
            ids = self.rows()
        if 0 <= ri < len(ids):
            self.tree.selection_set(ids[ri])
            self.edit(ids[ri], self.keys[ci])
        return "break"
